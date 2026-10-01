"""
ИИ-генератор BPMN по тексту регламента (ПАО «Интер РАО», трек «Архитектор BPMN-диаграмм»).

Конвейер:
    регламент (RU) ──► LLM (Ollama / OpenAI-совместимый API) ──► Python-код для DIAGRAM
                  └──► [fail-safe] семантический эмулятор ────────┘
                                   │
                        безопасная песочница (AST-проверка)
                                   │
                    heal_graph → layout → BPMN XML + аудит

Публичный API:
    generate_bpmn_from_text(regulation_text) -> (bpmn_xml, audit_data, error)
    execute_generated_code(code_str)         -> (bpmn_xml, audit_data, error)
    assistant_chat(...)                      -> диалог сайдбара (аналитика / правка / реверс)
    process_facts(audit, tobe_delta)         -> общие цифры SLA/циклов/bus-factor/To-Be
    canvas_copilot_reply(...)                -> локальный ответ плавающего копайлота (без XML)
    generate_process_passport(xml, audit, text) -> Markdown «Паспорт процесса»
    inspect_task_details(task_name, task_role, process_context) -> операционная карточка
    build_diagram_catalog(xml, audit, text) -> метаданные узлов для клика по холсту
    build_canvas_copilot(xml, audit, text) -> ответы плавающего ассистента на холсте
    optimize_process_to_be(text, audit) -> целевой регламент To-Be и дельта SLA
    generate_raci_matrix(steps, roles) -> матрица ответственности R/A/C/I
    export_docx_passport(xml, audit, text) -> официальный регламент Microsoft Word
    PROMPT_TEMPLATE / build_prompt(text)

Модуль не падает при недоступности моделей: при любом сбое сети, таймауте или
некорректном коде включается встроенный эмулятор на основе семантического
сопоставления шагов регламента (роли, условия, параллельность, декомпозиция).
"""

from __future__ import annotations

import ast
import builtins
from collections import defaultdict
import io
import json
import os
import re
import signal
import threading
import time
import xml.etree.ElementTree as _ET
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Sequence, Set, Tuple

from bpmn_framework import DEFAULT_HOURS, GATEWAY_KINDS, WORK_KINDS, BPMNDiagramBuilder
from validate_bpmn import xsd_errors_xml

ROOT_PROCESS_ID = "Process_Root"
ROOT_START_TASK_ID = "Event_RootStart"
ROOT_END_TASK_ID = "Event_RootEnd"

OLLAMA_PREFERRED_MODELS = ("qwen2.5-coder", "llama3")
OPENAI_DEFAULT_BASE = "https://api.openai.com/v1"

# --------------------------------------------------------------------------- #
# Системный промпт
# --------------------------------------------------------------------------- #
PROMPT_TEMPLATE = """Ты — Senior Enterprise Business Architect ПАО «Интер РАО» (Дирекция бизнес-архитектуры).
Твоя задача — превратить текст регламента в исполняемый Python-код построения BPMN 2.0 диаграммы
через готовый объект DIAGRAM. Объект DIAGRAM и константы ROOT_PROCESS_ID, ROOT_START_TASK_ID,
ROOT_END_TASK_ID уже определены — НЕ импортируй ничего и НЕ создавай их заново.

API объекта DIAGRAM (строго эти сигнатуры):
- pool_id, lane_ids = DIAGRAM.add_pool(ROOT_PROCESS_ID, ["Роль 1", "Роль 2", ...])   # список дорожек
- task_id = DIAGRAM.add_task(name, parent_id)              # обычная задача
- task_id = DIAGRAM.add_user_task(name, parent_id)         # действие человека
- task_id = DIAGRAM.add_script_task(name, parent_id)       # автоматическое действие системы
- sub_id = DIAGRAM.create_subprocess(name, parent_id)      # подпроцесс; внутрь добавляй задачи с parent_id=sub_id
- gw_id = DIAGRAM.add_exclusive_gateway(name, parent_id)   # XOR: name — вопрос («Допуск выдан?»)
- gw_id = DIAGRAM.add_parallel_gateway(name, parent_id)    # AND: разветвление/слияние параллельных веток
- gw_id = DIAGRAM.add_inclusive_gateway(name, parent_id)   # OR
- group_id = DIAGRAM.add_group(name, parent_id)
- DIAGRAM.add_link(source_id, target_id, condition_name="")  # поток; у веток шлюза ОБЯЗАТЕЛЬНА подпись условия
- DIAGRAM.set_sla(task_id, hours)                           # трудозатраты шага В ЧАСАХ. ОБЯЗАТЕЛЬНО, если срок есть в регламенте.
                                                             #   «30 минут» → 0.5; «2 часа» → 2; «5 рабочих дней» → 5 × 8 = 40; «3 календарных дня» → 72

ЖЁСТКИЕ ТРЕБОВАНИЯ:
1. РОЛЕВАЯ МОДЕЛЬ. Выдели всех участников регламента и создай для них дорожки одним вызовом add_pool.
   parent_id для узлов — это id нужной дорожки (элемент lane_ids). Каждый шаг размещай в дорожке исполнителя.
2. ДЕКОМПОЗИЦИЯ (анти-«метро Токио»). Если последовательность шагов ОДНОГО подразделения превышает 3
   действия подряд — обязательно сгруппируй их в DIAGRAM.create_subprocess с понятным названием этапа.
   Задачи внутри подпроцесса соединяй между собой add_link; parent_id = id подпроцесса.
3. УСЛОВИЯ. Каждое «если/иначе» — exclusive-шлюз; обе исходящие ветки подписаны (например «Допуск выдан»,
   «Замечания»). Возвраты на доработку моделируй обратной связью на предыдущий шаг.
4. ПАРАЛЛЕЛЬНОСТЬ. Шаги «одновременно/параллельно» — через parallel-шлюзы (разветвление и слияние).
5. ЦЕЛОСТНОСТЬ. Все ветки обязаны начинаться от ROOT_START_TASK_ID и заканчиваться на ROOT_END_TASK_ID:
   первый шаг связывай DIAGRAM.add_link(ROOT_START_TASK_ID, ...), все финальные ветки —
   DIAGRAM.add_link(..., ROOT_END_TASK_ID). Тупиков и «висящих» узлов быть не должно.
6. Имена задач — короткие, в форме «глагол + объект» (до 70 символов), на русском языке.
7. ПОСЛЕДОВАТЕЛЬНОЕ СОГЛАСОВАНИЕ. Если в тексте описано последовательное согласование несколькими лицами
   («любой вносит замечания — возврат на доработку»), создавай цепочку:
   Задача согласования 1 → Exclusive Gateway (Замечания → возврат, Согласовано → Задача согласования 2) → и т.д.
   НИКОГДА не делай 2 выхода из обычной задачи без шлюза.
8. АТОМАРНОСТЬ ЗАДАЧ: Каждая задача в BPMN должна принадлежать ровно одному исполнителю. Если в одном пункте
   регламента описаны действия нескольких участников (например: «Инициатор готовит заявку, а согласующий проверяет её»)
   или последовательные этапы со связками «после чего / затем» — ОБЯЗАТЕЛЬНО разбивай их на отдельные последовательные
   задачи (userTask / scriptTask) в дорожках соответствующих ролей. Однородные действия одного исполнителя
   («проверить и подписать акт») оставляй одной задачей.
9. ФОРМАТ ОТВЕТА: верни ТОЛЬКО исполняемый Python-код для объекта DIAGRAM. Никаких markdown-тегов
   (без ```), пояснений, импортов, циклов while и обращений к файлам/сети.

ПРИМЕР ОТВЕТА:
pool_id, lanes = DIAGRAM.add_pool(ROOT_PROCESS_ID, ["Инициатор", "Согласующий"])
a = DIAGRAM.add_user_task("Подготовить заявку", lanes[0])
b = DIAGRAM.add_user_task("Проверить заявку", lanes[1])
g = DIAGRAM.add_exclusive_gateway("Заявка корректна?", lanes[1])
c = DIAGRAM.add_user_task("Исполнить заявку", lanes[1])
DIAGRAM.add_link(ROOT_START_TASK_ID, a)
DIAGRAM.add_link(a, b)
DIAGRAM.add_link(b, g)
DIAGRAM.add_link(g, c, "Заявка корректна")
DIAGRAM.add_link(g, a, "Замечания")
DIAGRAM.add_link(c, ROOT_END_TASK_ID)

ТЕКСТ РЕГЛАМЕНТА:
<<REGULATION>>
"""


def build_prompt(regulation_text: str) -> str:
    return PROMPT_TEMPLATE.replace("<<REGULATION>>", regulation_text.strip())


REPAIR_TEMPLATE = """

ТВОЙ ПРЕДЫДУЩИЙ ОТВЕТ:
<<CODE>>

В НЁМ НАЙДЕНЫ ОШИБКИ СТРУКТУРЫ ПРОЦЕССА:
<<ISSUES>>

Исправь их и верни ПОЛНЫЙ исправленный код целиком (только Python-код для DIAGRAM, без пояснений).
Помни: ветвление только через шлюз, у каждой задачи ровно один выход, каждый узел имеет вход и выход."""


def build_repair_prompt(regulation_text: str, code: str, issues: List[str]) -> str:
    """Промпт второй попытки: исходная задача + прошлый код + найденные ошибки."""
    return build_prompt(regulation_text) + (
        REPAIR_TEMPLATE.replace("<<CODE>>", code.strip()).replace("<<ISSUES>>", "\n".join(f"- {i}" for i in issues))
    )


# --------------------------------------------------------------------------- #
# Песочница исполнения сгенерированного кода
# --------------------------------------------------------------------------- #
_ALLOWED_BUILTINS = (
    "len range enumerate zip list dict tuple set str int float bool min max sum sorted reversed abs round "
    "any all map filter isinstance repr print Exception KeyError ValueError TypeError IndexError AttributeError"
).split()
_FORBIDDEN_NAMES = {
    "eval", "exec", "compile", "open", "__import__", "globals", "locals", "vars", "getattr", "setattr",
    "delattr", "input", "breakpoint", "exit", "quit", "help", "memoryview", "classmethod", "staticmethod",
}
_MAX_RANGE = 2000


class UnsafeCodeError(ValueError):
    """Сгенерированный код нарушает правила песочницы."""


class _DiagramExecTimeout(BaseException):
    """Таймаут exec. BaseException, чтобы except Exception внутри кода его не съел."""


def _strip_markdown(code: str) -> str:
    fenced = re.findall(r"```(?:python|py)?\s*\n(.*?)```", code, flags=re.S | re.I)
    if fenced:
        code = "\n".join(fenced)
    else:
        code = code.replace("```python", "").replace("```", "")
    lines = []
    for line in code.splitlines():
        stripped = line.strip()
        if re.match(r"^(import\s+\S+|from\s+\S+\s+import\s+.+)$", stripped):
            continue  # модель любит добавлять импорты — они не нужны и запрещены
        if re.match(r"^DIAGRAM\s*=", stripped):
            continue  # DIAGRAM уже предоставлен песочницей
        lines.append(line)
    return "\n".join(lines).strip() + "\n"


def _validate_ast(tree: ast.AST) -> None:
    for node in ast.walk(tree):
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            raise UnsafeCodeError("импорты запрещены")
        if isinstance(node, (ast.While, ast.AsyncFor, ast.AsyncFunctionDef, ast.Await, ast.Global, ast.Nonlocal)):
            raise UnsafeCodeError(f"конструкция {type(node).__name__} запрещена")
        if isinstance(node, ast.Attribute) and node.attr.startswith("_"):
            raise UnsafeCodeError(f"доступ к атрибуту {node.attr!r} запрещён")
        if isinstance(node, ast.Name) and (node.id in _FORBIDDEN_NAMES or node.id.startswith("__")):
            raise UnsafeCodeError(f"имя {node.id!r} запрещено")
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "range":
            for arg in node.args:
                if isinstance(arg, ast.Constant) and isinstance(arg.value, int) and abs(arg.value) > _MAX_RANGE:
                    raise UnsafeCodeError("слишком большой диапазон range")


def _structure_issues(diagram: BPMNDiagramBuilder) -> List[str]:
    """Логические ошибки графа ДО самовосстановления: то, что heal_graph скрыл бы «метро»."""
    issues: List[str] = []
    outgoing: Dict[str, List[Any]] = {}
    incoming: Dict[str, int] = {}
    for link in diagram.links:
        outgoing.setdefault(link.source_id, []).append(link)
        incoming[link.target_id] = incoming.get(link.target_id, 0) + 1
    for node in diagram.nodes.values():
        outs = outgoing.get(node.id, [])
        targets = {l.target_id for l in outs}  # дубли в одну цель (после переадресации движком) — не ветвление
        if node.kind in WORK_KINDS | {"subProcess"} and len(targets) > 1:
            issues.append(f"У задачи «{node.name}» {len(targets)} выхода — ветвление без шлюза")
        if node.kind in GATEWAY_KINDS and len(outs) <= 1 and incoming.get(node.id, 0) <= 1:
            issues.append(f"Шлюз «{node.name}» ничего не разветвляет и не сливает")
        if node.kind in ("exclusiveGateway", "inclusiveGateway") and len(outs) > 1:
            if any(not l.condition_name.strip() for l in outs):
                issues.append(f"У шлюза «{node.name}» есть ветка без подписи условия")
    return issues


def _quality_report(
    structure_issues: List[str],
    audit: Dict[str, Any],
    diagram: Optional[BPMNDiagramBuilder] = None,
) -> Dict[str, Any]:
    """Отчёт по графу после heal. Журнал лечения («без входа», «Тупик») отказ не вызывает.

    Критично только то, что осталось в графе: нет старта или конца, тупик, узел без входа.
    Битый XML отсекается отдельно, проверкой XSD.
    """
    issues = list(structure_issues)
    issues += [f"Связь пропущена: {s['reason']}" for s in audit.get("skipped_links", [])]
    critical: List[str] = []
    nodes = list(diagram.nodes.values()) if diagram is not None else []
    if diagram is not None:
        if not any(n.kind == "startEvent" for n in nodes):
            critical.append("В графе нет стартового события")
        if not any(n.kind == "endEvent" for n in nodes):
            critical.append("В графе нет конечного события")
    for item in audit.get("dead_ends") or []:
        msg = item.get("message") if isinstance(item, dict) else str(item)
        if msg:
            critical.append(str(msg))
    for oid in audit.get("orphans_without_incoming") or []:
        name = str(oid)
        if diagram is not None and oid in diagram.nodes:
            name = diagram.nodes[oid].name or name
        critical.append(f"Узел «{name}» без входящего потока")
    return {"issues": issues, "critical": critical, "ok": not critical}


def _sla_stems(text: str) -> set:
    return {w[:6] for w in re.findall(r"[А-Яа-яЁёA-Za-z]{4,}", (text or "").lower())}


def _sla_title_score(left: str, right: str) -> int:
    a, b = _sla_stems(left), _sla_stems(right)
    if not a or not b:
        return 0
    return len(a & b)


def _sla_looks_default(node: Any) -> bool:
    """True, если трудозатраты не заданы или совпадают с дефолтом вида узла (2 ч для userTask)."""
    if getattr(node, "sla_hours", None) is None:
        return True
    expected = DEFAULT_HOURS.get(getattr(node, "kind", ""), None)
    if expected is None:
        return False
    try:
        return abs(float(node.sla_hours) - float(expected)) < 1e-6
    except (TypeError, ValueError):
        return True


def _enrich_sla_from_regulation(diagram: BPMNDiagramBuilder, regulation_text: str) -> int:
    """Подставляет часы из текста регламента узлам без явного set_sla (или с дефолтом 2 ч)."""
    if not regulation_text or not str(regulation_text).strip():
        return 0
    try:
        parsed = parse_regulation(normalize_regulation(regulation_text))
    except Exception:  # noqa: BLE001
        return 0
    steps = [s for s in parsed.steps if s.hours]
    work = [n for n in diagram.nodes.values() if n.kind in WORK_KINDS]
    if not steps or not work:
        return 0
    applied = 0
    used_nodes: set = set()

    def _free(node: Any) -> bool:
        return node.id not in used_nodes and _sla_looks_default(node)

    for step in steps:
        title = step.title or ""
        best, score = None, 0
        for node in work:
            if not _free(node):
                continue
            sc = _sla_title_score(title, node.name or "")
            if sc > score:
                best, score = node, sc
        if best is not None and score >= 1:
            best.sla_hours = float(step.hours)
            used_nodes.add(best.id)
            applied += 1

    return applied


def _run_diagram_exec(code: Any, namespace: Dict[str, Any]) -> None:
    """Исполняет код DIAGRAM. На главном потоке обрывает зависший цикл по таймеру."""
    timeout = float(os.getenv("DIAGRAM_EXEC_TIMEOUT", "12"))

    def _handle(signum: int, frame: Any) -> None:
        raise _DiagramExecTimeout(f"дольше {timeout:g} с")

    on_main = threading.current_thread() is threading.main_thread()
    if timeout <= 0 or not on_main or not hasattr(signal, "setitimer"):
        exec(code, namespace)  # noqa: S102 — песочница: AST-фильтр + whitelist builtins
        return
    old = signal.signal(signal.SIGALRM, _handle)
    signal.setitimer(signal.ITIMER_REAL, timeout)
    try:
        exec(code, namespace)  # noqa: S102 — песочница: AST-фильтр + whitelist builtins
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, old)


def execute_generated_code(
    code_str: str,
    process_name: str = "Бизнес-процесс ПАО «Интер РАО»",
    sla_target_hours: Optional[float] = None,
    regulation_text: str = "",
) -> Tuple[str, Dict[str, Any], str]:
    """Безопасно исполняет код для DIAGRAM и возвращает (bpmn_xml, audit_data, error).

    Любые исключения перехватываются: при ошибке возвращается ("", {}, "описание").
    """
    try:
        code = _code_for_sandbox(_strip_markdown(code_str or ""))
        if len(code.strip()) < 10:
            return "", {}, "Пустой код: модель не вернула инструкций для DIAGRAM."
        try:
            tree = ast.parse(code, filename="<generated>")
        except SyntaxError as exc:
            return "", {}, f"Синтаксическая ошибка в сгенерированном коде (строка {exc.lineno}): {exc.msg}"
        try:
            _validate_ast(tree)
        except UnsafeCodeError as exc:
            return "", {}, f"Код отклонён песочницей: {exc}."

        diagram = BPMNDiagramBuilder(
            process_name=process_name,
            process_id=ROOT_PROCESS_ID,
            root_start_id=ROOT_START_TASK_ID,
            root_end_id=ROOT_END_TASK_ID,
            sla_target_hours=sla_target_hours,
        )
        safe_builtins = {name: getattr(builtins, name) for name in _ALLOWED_BUILTINS}
        namespace: Dict[str, Any] = {
            "__builtins__": safe_builtins,
            "DIAGRAM": diagram,
            "ROOT_PROCESS_ID": ROOT_PROCESS_ID,
            "ROOT_START_TASK_ID": ROOT_START_TASK_ID,
            "ROOT_END_TASK_ID": ROOT_END_TASK_ID,
        }
        try:
            _run_diagram_exec(compile(tree, "<generated>", "exec"), namespace)
        except _DiagramExecTimeout as exc:
            return "", {}, f"Исполнение кода прервано по таймауту: {exc}"

        work_nodes = [n for n in diagram.nodes.values() if n.kind not in ("startEvent", "endEvent")]
        if len(work_nodes) < 2:
            return "", {}, _UNRECOGNIZED_PROCESS

        sla_enriched = _enrich_sla_from_regulation(diagram, regulation_text)
        diagram.heal_graph()
        _complete_return_edges(diagram, regulation_text)
        structure_issues = _structure_issues(diagram)
        xml = diagram.to_bpmn_xml(ROOT_PROCESS_ID, ROOT_START_TASK_ID, ROOT_END_TASK_ID)
        audit = diagram.analyze_bottlenecks()
        audit["quality"] = _quality_report(structure_issues, audit, diagram)
        expected_returns = _explicit_return_count(regulation_text)
        actual_returns = len(audit.get("rework_loops") or [])
        if expected_returns > actual_returns:
            return "", {}, _return_diagram_not_ready(expected_returns, actual_returns)
        if sla_enriched:
            audit["sla_enriched_nodes"] = sla_enriched
        errors = xsd_errors_xml(xml)  # официальная XSD BPMN 2.0: невалидный файл не отдаём
        if errors:
            return "", {}, "XML не прошёл проверку по XSD BPMN 2.0: " + "; ".join(errors[:3])
        audit["xsd_valid"] = errors is not None
        audit.setdefault("artifacts", [])
        audit.setdefault("it_systems", [])
        return xml, audit, ""
    except Exception as exc:  # noqa: BLE001 — песочница обязана не падать
        return "", {}, f"Ошибка исполнения кода: {type(exc).__name__}: {exc}"


# --------------------------------------------------------------------------- #
# HTTP-клиенты (Ollama / OpenAI-совместимый API)
# --------------------------------------------------------------------------- #
def _http_json(
    url: str,
    payload: Optional[Dict[str, Any]] = None,
    headers: Optional[Dict[str, str]] = None,
    timeout: float = 5.0,
) -> Dict[str, Any]:
    try:
        import requests  # type: ignore

        if payload is None:
            resp = requests.get(url, headers=headers, timeout=timeout)
        else:
            resp = requests.post(url, json=payload, headers=headers, timeout=timeout)
        resp.raise_for_status()
        return resp.json()
    except ImportError:  # requests не установлен — стандартная библиотека
        import urllib.request

        data = json.dumps(payload).encode("utf-8") if payload is not None else None
        req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json", **(headers or {})})
        with urllib.request.urlopen(req, timeout=timeout) as resp:  # noqa: S310
            return json.loads(resp.read().decode("utf-8"))


def _ollama_base() -> str:
    return os.getenv("OLLAMA_URL", "http://localhost:11434").rstrip("/")


def _ollama_models() -> List[str]:
    data = _http_json(f"{_ollama_base()}/api/tags", timeout=1.5)
    return [m.get("name", "") for m in data.get("models", []) if m.get("name")]


def _pick_ollama_model(models: List[str]) -> Optional[str]:
    prefs = [os.getenv("OLLAMA_MODEL", "")] + list(OLLAMA_PREFERRED_MODELS)
    for pref in filter(None, prefs):
        for name in models:
            if name.lower().startswith(pref.lower()):
                return name
    return None


def _call_ollama(prompt: str) -> Tuple[str, str]:
    model = _pick_ollama_model(_ollama_models())
    if not model:
        raise RuntimeError("в Ollama нет моделей qwen2.5-coder / llama3")
    # num_ctx: промпт ~1800 токенов + ответ до 3500 не помещаются в дефолтные 4096 Ollama.
    # Таймаут 300 с: 7B-модель на MacBook M2 генерирует ~5 ток/с, схема — 3–4 минуты.
    options = {"temperature": 0.1, "num_predict": 3500, "num_ctx": int(os.getenv("OLLAMA_NUM_CTX", "8192"))}
    data = _http_json(
        f"{_ollama_base()}/api/generate",
        {"model": model, "prompt": prompt, "stream": False, "keep_alive": "30m", "options": options},
        timeout=float(os.getenv("OLLAMA_TIMEOUT", "300")),
    )
    return f"ollama:{model}", str(data.get("response", ""))


def _call_openai(prompt: str, prefix: str = "OPENAI") -> Tuple[str, str]:
    """OpenAI-совместимый chat/completions. prefix: OPENAI_* — основной провайдер, FALLBACK_* — запасной."""
    key = os.getenv(f"{prefix}_API_KEY")
    if not key:
        raise RuntimeError(f"{prefix}_API_KEY не задан")
    base = os.getenv(f"{prefix}_BASE_URL", OPENAI_DEFAULT_BASE).rstrip("/")
    model = os.getenv(f"{prefix}_MODEL", "gpt-4o-mini")
    system, _, regulation = prompt.partition("ТЕКСТ РЕГЛАМЕНТА:")
    payload: Dict[str, Any] = {
        "model": model,
        "temperature": 0.1,
        "messages": [
            {"role": "system", "content": system.strip()},
            {"role": "user", "content": "ТЕКСТ РЕГЛАМЕНТА:\n" + regulation.strip()},
        ],
    }
    # Рассуждающие модели (gpt-oss на Groq): low — меньше «мыслей», быстрее и в пределах лимита токенов/мин.
    effort = os.getenv(f"{prefix}_REASONING_EFFORT", "")
    if effort == "none" and "openrouter.ai" in base:
        payload["reasoning"] = {"enabled": False}  # OpenRouter: без «размышлений» qwen3 отвечает в ~2 раза быстрее
    elif effort:
        payload["reasoning_effort"] = effort
    for attempt in range(2):
        try:
            data = _http_json(
                f"{base}/chat/completions",
                payload,
                headers={"Authorization": f"Bearer {key}"},
                timeout=float(os.getenv("OPENAI_TIMEOUT", "90")),
            )
            break
        except Exception as exc:  # noqa: BLE001
            wait = _rate_limit_wait(exc)
            if attempt or wait is None:
                raise
            time.sleep(wait)  # бесплатный тариф Groq: 8K токенов/мин — ждём, сколько просит провайдер
    label = "openai" if prefix == "OPENAI" else prefix.lower()
    return f"{label}:{model}", str(data["choices"][0]["message"]["content"])


# --------------------------------------------------------------------------- #
# Семантический эмулятор (fail-safe)
# --------------------------------------------------------------------------- #
ROLE_PATTERNS: List[Tuple[str, str]] = [
    (r"диспетчерск\w+ служб\w*|диспетчер\w*", "Диспетчер"),
    (r"начальник\w* смены", "Начальник смены"),
    (r"начальник\w* служб\w* подстанц\w*|служб\w* подстанц\w*", "Служба подстанций"),
    (r"ремонтн\w+ бригад\w*|бригад[аыуеой]\b", "Ремонтная бригада"),
    (r"служб\w* безопасности|\bСБ\b", "Служба безопасности"),
    (r"тендерн\w+ комитет\w*|закупочн\w+ комисси\w*", "Тендерный комитет"),
    (r"закупочн\w+ служб\w*|департамент\w* закупок|отдел\w* закупок", "Закупочная служба"),
    (r"технический департамент\w*|технический заказчик\w*|инициатор\w* закупки", "Технический департамент"),
    (r"юридическ\w+ служб\w*|юридическ\w+ департамент\w*|юрист\w*", "Юридическая служба"),
    (r"финансов\w+ департамент\w*|казначейств\w*", "Финансовый департамент"),
    (r"заявител\w+|потребител\w+", "Заявитель"),
    (r"центр\w* обслуживания клиентов|клиентск\w+ служб\w*", "Центр обслуживания клиентов"),
    (r"отдел\w* технологического присоединения", "Отдел технологического присоединения"),
    (r"эксплуатационн\w+ служб\w*|служб\w* эксплуатации", "Эксплуатационная служба"),
    (r"эколог\w*|служб\w* экологии|отдел\w* экологии", "Служба экологии"),
    (r"охран\w* труда|специалист\w* по охране труда|инженер\w* по охране труда", "Служба охраны труда"),
    (r"\bРВБ\b|ремонтно-восстановительн\w+ бригад\w*", "РВБ"),
    (r"\bОМТО\b|отдел\w* материально-техническ\w+|служб\w* МТО", "ОМТО"),
    (r"главн\w+ инженер\w*", "Главный инженер"),
    (r"бухгалтер\w*", "Бухгалтерия"),
    (r"служб\w* (?:информационных технологий|ИТ)\b|ИТ-служб\w*|\bИТ-отдел\w*", "Служба ИТ"),
    (r"руководител\w+|директор\w*", "Руководитель"),
    (r"согласующ\w*", "Согласующий"),
    (r"\bинициатор\w*", "Инициатор"),
]
_ROLE_RE = [(re.compile(p, re.I), name) for p, name in ROLE_PATTERNS]

VERB_MAP = {
    "принимает": "принять", "подаёт": "подать", "подает": "подать", "выдаёт": "выдать", "выдает": "выдать",
    "вносит": "внести", "проводит": "провести", "готовит": "подготовить", "производит": "произвести",
    "оформляет": "оформить", "направляет": "направить", "определяет": "определить", "проверяет": "проверить",
    "выполняет": "выполнить", "составляет": "составить", "подписывает": "подписать", "заключает": "заключить",
    "передаёт": "передать", "передает": "передать", "допускает": "допустить", "выводит": "вывести",
    "восстанавливает": "восстановить", "измеряет": "измерить", "рассчитывает": "рассчитать",
    "регистрирует": "зарегистрировать", "формирует": "сформировать", "фиксирует": "зафиксировать",
    "рассматривает": "рассмотреть", "рассматривают": "рассмотреть",
    "уведомляет": "уведомить", "закрывает": "закрыть", "обосновывает": "обосновать",
    "уточняет": "уточнить", "разрабатывает": "разработать", "оценивает": "оценить", "подключает": "подключить",
    "выезжает": "выехать", "выезжают": "выехать",
    "организует": "организовать", "устраняет": "устранить", "выбирает": "выбрать", "готовят": "подготовить",
    "создаёт": "создать", "создает": "создать", "заносит": "занести", "платит": "оплатить", "берёт": "взять",
    "отвечает": "ответить", "решает": "решить", "принимают": "принять", "отправит": "отправить",
    "вводит": "ввести", "выносит": "вынести", "запросит": "запросить", "получит": "получить",
}
_VERB_SUFFIXES = (("ирует", "ировать"), ("ует", "овать"), ("ает", "ать"), ("яет", "ять"))
_SYSTEM_RE = re.compile(
    r"автоматическ|в (?:оперативном )?журнал|в (?:информационной |автоматизированной |электронной )?систем|"
    r"рассчитывает|электронн\w+ площадк|в реестр",
    re.I,
)
_DUR_RE = re.compile(
    r"\(?\s*(?:(?:в течение|не более|до|срок[:\s]*)\s*)?(\d+(?:[.,]\d+)?)\s*"
    r"(рабоч\w*\s+(?:дн\w*|день)|календарн\w*\s+(?:дн\w*|день)|сут\w*|дн\w*|день|час\w*|ч\b|мин\w*)\s*\)?",
    re.I,
)
_REF_RE = re.compile(r"(?:п(?:ункт\w*|\.|п\.)?|шаг\w*)\s*(\d+)", re.I)
_END_KW = re.compile(r"завершить|завершается|завершение процесса|прекрат|закрыть\s+(?:заявку|процесс|закупку)", re.I)
_BACK_KW = re.compile(r"верну|возврат|возвраща|доработ|повтор|заново|перенос\s+срок", re.I)
_EXPLICIT_RETURN_RE = re.compile(
    r"верну\w*|возврат\w*|возвраща\w*|на\s+доработк\w*|\bповторно\b|при\s+замечаниях",
    re.I,
)
_NAZAD_RETURN_RE = re.compile(r"(?:если|иначе|при)\b[^.]{0,80}\bназад\b", re.I)
_CASE_RE = re.compile(r"\b(?:если|в случае(?:\s+если)?)\b", re.I)
_RETURN_RE = re.compile(
    r"(?:возвраща\w+|верну\w+|направляется\s+на\s+доработ|на\s+доработк|перенос\s+срок)",
    re.I,
)
_INITIATOR_RE = re.compile(r"инициатор\w*", re.I)
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.;!?])\s+(?=[А-ЯЁA-Z«\"])")
_NUMBERED_LINE_RE = re.compile(r"^\s*\d+[.)]\s+\S")
_SEQ_SPLIT_RE = re.compile(
    r"(?:;+\s*|(?:,\s*)?(?:после\s+чего|затем|далее|после\s+этого|при\s+этом)\s+)",
    re.I,
)
_WHO_RE = re.compile(r",\s+кото(?:рый|рая|рое|рые)\s+", re.I)
_A_CONJ_RE = re.compile(r",\s+а\s+(?!также\b)", re.I)


_PAGE_MARK_RE = re.compile(r"(?:стр\.?|страница)\s*\d+\s*(?:из\s*\d+)?", re.I)
_MULTI_NUM_RE = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+", re.M)
_ROMAN_LINE_RE = re.compile(
    r"^(?P<rom>(?=[IVXLCDM])M{0,4}(?:CM|CD|D?C{0,3})(?:XC|XL|L?X{0,3})(?:IX|IV|V?I{0,3}))[.)]\s+(?P<title>\S.*)$",
    re.I,
)
_LETTER_SUB_RE = re.compile(r"^([а-яёa-z])[).]\s+(\S.*)$", re.I)
_JOINT_LINK_RE = re.compile(r"^(?P<head>.+?)\s+(?P<link>совместно\s+с|вместе\s+с)\s+", re.I)
_CYR_SUB_LETTERS = "абвгдежзиклмнопрстуфхцчшщэюя"
_LAT_SUB_LETTERS = "abcdefghijklmnopqrstuvwxyz"
_SUB_LETTERS = set(_CYR_SUB_LETTERS + _LAT_SUB_LETTERS)
_PAR_WORD_RE = re.compile(r"\b(?:параллельно|одновременно)\b", re.I)
_PAR_GENERIC_RE = re.compile(
    r"(?:запуска\w*|старту\w*|начина\w*|выполня\w+ся|проводятся|ид[её]т|независим)",
    re.I,
)
_LET_TOKEN_RE = re.compile(r"(^|[\s;:])([а-яёa-z])[)]\s+", re.I)
_INTRO_CLAUSE_RE = re.compile(
    r"^(?:"
    r"после\s+(?:получения|утверждения|согласования|поступления|"
    r"завершения|окончания|оформления|подписания|выполнения|уведомления|"
    r"рассмотрения|регистрации)(?:\s+дефектной\s+ведомости|\s+наряда[\s-]*допуска|\s+\S+)?"
    r"|перед\s+(?:началом|стартом|проведением|выполнением)(?:\s+работ|\s+\S+)?"
    r"|запуска\w*(?:\s+(?:три|два|несколько|все|независим\w+))*\s+процесс\w*"
    r")[,:]?\s*",
    re.I,
)


def _letter_line_body(s: str) -> Optional[str]:
    lm = _LETTER_SUB_RE.match((s or "").strip())
    if not lm:
        return None
    if lm.group(1).lower() not in _SUB_LETTERS:
        return None
    return lm.group(2).strip()


def _split_inline_lettered(text: str) -> Tuple[str, List[str]]:
    src = text or ""
    found = [m for m in _LET_TOKEN_RE.finditer(src) if m.group(2).lower() in _SUB_LETTERS]
    if len(found) < 2:
        return src.strip(), []
    items: List[str] = []
    for i, m in enumerate(found):
        end = found[i + 1].start() if i + 1 < len(found) else len(src)
        chunk = src[m.end():end].strip(" ;,.\t")
        if chunk:
            items.append(chunk)
    if len(items) < 2:
        return src.strip(), []
    intro = src[: found[0].start()].strip(" :;,—–-\t")
    return intro, items


def _is_parallel_fork_intro(body: str) -> bool:
    t = (body or "").strip()
    if not t:
        return True
    low = t.lower()
    has_par = bool(_PAR_WORD_RE.search(low))
    generic = bool(_PAR_GENERIC_RE.search(low)) and bool(
        re.search(r"процесс|работ|ветк|поток|действи|независим", low)
    )
    stripped = _PAR_WORD_RE.sub("", t, count=1)
    stripped = re.sub(r"^[\s:,—–-]+", "", stripped)
    role, _, _ = _find_role(stripped)
    if has_par and generic:
        return True
    if generic and role is None:
        return True
    if has_par and role is None and len(stripped.split()) <= 12:
        return True
    return False


def _promote_lettered_parallels(text: str) -> str:
    """Буквенные подпункты внутри параллельного блока → отдельные шаги с префиксом «Параллельно»."""
    raw_lines = (text or "").split("\n")
    out: List[str] = []
    i = 0
    while i < len(raw_lines):
        raw = raw_lines[i]
        s = raw.strip()
        num_m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if not num_m:
            if _PAR_WORD_RE.search(s) and _is_parallel_fork_intro(s):
                kids: List[str] = []
                j = i + 1
                skipped_blank = 0
                while j < len(raw_lines):
                    stripped = raw_lines[j].strip()
                    if not stripped:
                        skipped_blank += 1
                        j += 1
                        continue
                    lb = _letter_line_body(stripped)
                    if lb is None:
                        j -= skipped_blank
                        break
                    kids.append(lb)
                    skipped_blank = 0
                    j += 1
                if len(kids) >= 2:
                    for k, kid in enumerate(kids):
                        mark = "Параллельно*: " if k == 0 else "Параллельно: "
                        let = _CYR_SUB_LETTERS[k] if k < len(_CYR_SUB_LETTERS) else _LAT_SUB_LETTERS[min(k, 25)]
                        out.append(f"{let}) {mark}{kid}")
                    i = j
                    continue
            out.append(raw)
            i += 1
            continue
        num, body = num_m.group(1), num_m.group(2)
        intro_inline, inline_items = _split_inline_lettered(body)
        kids: List[str] = list(inline_items)
        j = i + 1
        skipped_blank = 0
        while j < len(raw_lines):
            stripped = raw_lines[j].strip()
            if not stripped:
                skipped_blank += 1
                j += 1
                continue
            lb = _letter_line_body(stripped)
            if lb is None:
                j -= skipped_blank
                break
            kids.append(lb)
            skipped_blank = 0
            j += 1
        head = intro_inline if inline_items else body
        if kids and _is_parallel_fork_intro(head):
            for k, kid in enumerate(kids):
                mark = "Параллельно*: " if k == 0 else "Параллельно: "
                if k == 0:
                    out.append(f"{num}. {mark}{kid}")
                else:
                    let = _CYR_SUB_LETTERS[k] if k < len(_CYR_SUB_LETTERS) else _LAT_SUB_LETTERS[min(k, 25)]
                    out.append(f"{let}) {mark}{kid}")
            i = j
            continue
        if kids and _PAR_WORD_RE.search(head):
            out.append(f"{num}. {head.strip()}")
            for k, kid in enumerate(kids):
                let = _CYR_SUB_LETTERS[k] if k < len(_CYR_SUB_LETTERS) else _LAT_SUB_LETTERS[min(k, 25)]
                out.append(f"{let}) Параллельно: {kid}")
            i = j
            continue
        out.append(raw)
        i += 1
    return "\n".join(out)


def _expand_corporate_numbering(text: str) -> str:
    """Римские этапы (I./II.) и буквенные подпункты (а)/б)) → сквозные шаги с маркером стадии."""
    out: List[str] = []
    stage: Optional[str] = None
    last_n = 0
    extra = 0
    for line in text.split("\n"):
        s = line.strip()
        rm = _ROMAN_LINE_RE.match(s)
        if rm and rm.group("rom"):
            stage = rm.group("title").strip()
            continue
        lm = _LETTER_SUB_RE.match(s)
        if lm and lm.group(1).lower() in _SUB_LETTERS:
            last_n += 1
            extra += 1
            body = lm.group(2).strip()
            if stage and "этап" not in body.lower():
                body = f"{body} (этап «{stage}»)"
            out.append(f"{last_n}. {body}")
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            last_n = int(m.group(1)) + extra
            body = m.group(2)
            if stage and "этап" not in body.lower():
                body = f"{body} (этап «{stage}»)"
            out.append(f"{last_n}. {body}")
            continue
        out.append(line)
    return "\n".join(out)


def _count_explicit_numbers(text: str) -> int:
    n = 0
    for line in (text or "").splitlines():
        s = line.strip()
        if _NUMBERED_LINE_RE.match(s):
            n += 1
            continue
        rm = _ROMAN_LINE_RE.match(s)
        if rm and rm.group("rom"):
            n += 1
    return n


def _roles_mentioned(text: str) -> List[str]:
    found: List[Tuple[int, str]] = []
    for regex, name in _ROLE_RE:
        for m in regex.finditer(text or ""):
            found.append((m.start(), name))
    found.sort()
    out: List[str] = []
    for _, name in found:
        if not out or out[-1] != name:
            out.append(name)
    return out


def _role_at_start(text: str) -> Optional[str]:
    src = (text or "").strip()
    if not src:
        return None
    joint = _split_joint_role(src)
    if joint:
        return joint[0]
    inv = _match_inverted_role(src)
    if inv:
        return inv[0]
    role, rs, _ = _find_role(src)
    if role and rs == 0:
        return role
    subj = _subject_before_verb(src)
    if subj:
        name = _canonical_role(subj[0])
        if name.lower() in {"который", "которая", "которое", "которые", "которого", "которой"}:
            return None
        return name
    return None


def _chunk_role(text: str) -> Optional[str]:
    return _role_at_start(text) or (_find_role(text or "")[0])


def _split_seq_conjunctions(text: str) -> List[str]:
    bits = _SEQ_SPLIT_RE.split(text or "")
    return [b.strip(" ,;.—–-") for b in bits if len(b.strip(" ,;.—–-")) > 2]


def _split_role_handoff(text: str) -> List[str]:
    """Один пункт с передачей другой роли → отдельные шаги; однородные «и» не режем."""
    t = (text or "").strip()
    if len(t) < 12:
        return [t] if t else []
    m = _WHO_RE.search(t)
    if m:
        left, right = t[: m.start()].strip(), t[m.end() :].strip()
        actor = _role_at_start(left) or _chunk_role(left)
        mentioned = _roles_mentioned(left)
        referred = mentioned[-1] if mentioned else None
        if referred and actor != referred and right:
            prefixed = right if _role_at_start(right) else f"{referred} {right}"
            return _split_role_handoff(left) + _split_role_handoff(prefixed)
    m = _A_CONJ_RE.search(t)
    if m:
        left, right = t[: m.start()].strip(), t[m.end() :].strip()
        rrole = _role_at_start(right)
        lrole = _role_at_start(left) or _chunk_role(left)
        if rrole and rrole != lrole:
            return _split_role_handoff(left) + _split_role_handoff(right)
    for cm in re.finditer(r",\s+", t):
        right = t[cm.end() :]
        rrole = _role_at_start(right)
        if not rrole:
            continue
        left = t[: cm.start()].strip()
        lrole = _role_at_start(left) or _chunk_role(left)
        if lrole and rrole != lrole:
            return _split_role_handoff(left) + _split_role_handoff(right)
    return [t]


def _merge_homogeneous(parts: List[str]) -> List[str]:
    """«Проверить и подписать акт» одного исполнителя остаётся одной задачей."""
    if not parts:
        return []
    out = [parts[0]]
    for part in parts[1:]:
        prev = out[-1]
        if re.match(r"^и\s+", part, re.I):
            out[-1] = (prev.rstrip(" ,;") + " " + part).strip()
            continue
        pr = _role_at_start(prev) or _chunk_role(prev)
        cr = _role_at_start(part) or _chunk_role(part)
        if cr is None and pr is not None and len(part.split()) <= 5:
            words = part.split()
            if words and not _is_verb(words[0].strip(".,;:")):
                out[-1] = (prev.rstrip(" ,;") + " " + part).strip()
                continue
        out.append(part)
    return out


def _atomize_step_body(text: str) -> List[str]:
    """Смысловые шаги: ; / после чего / затем / смена роли. Не дробит «если» и однородные «и».

    Несколько длительностей в одном абзаце остаются на своих предложениях и не суммируются.
    """
    t = re.sub(r"\s+", " ", (text or "").strip())
    if not t:
        return []
    if len(t) < 12:
        return [t]
    if _CASE_RE.search(t):
        return [t]
    owned = _split_owned_durations(t)
    if owned:
        return owned
    parts: List[str] = []
    for seq in _split_seq_conjunctions(t):
        parts.extend(_split_role_handoff(seq))
    merged = _merge_homogeneous(parts)
    return [p for p in merged if len(p.strip()) > 8] or [t]


def _split_owned_durations(text: str) -> Optional[List[str]]:
    """«(5 минут). … (30 мин).» → два шага. Чужой срок в этот шаг не входит."""
    parts = [p.strip(" ;") for p in _SENTENCE_SPLIT_RE.split(text or "") if len(p.strip(" ;")) > 8]
    if len(parts) < 2:
        return None
    if any(_CASE_RE.search(p) or _RETURN_RE.search(p) for p in parts[1:]):
        return None
    if sum(1 for p in parts if _DUR_RE.search(p)) < 2:
        return None
    return parts


def _atomize_numbered_lines(text: str) -> str:
    extra = 0
    out: List[str] = []
    for line in (text or "").split("\n"):
        s = line.strip()
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if not m:
            out.append(line)
            continue
        n = int(m.group(1))
        bits = _atomize_step_body(m.group(2))
        if len(bits) <= 1:
            out.append(f"{n + extra}. {bits[0]}" if bits else line)
            continue
        for i, bit in enumerate(bits):
            out.append(f"{n + extra + i}. {bit}")
        extra += len(bits) - 1
    return "\n".join(out)


def _split_prose_sentences(text: str) -> List[str]:
    src = (text or "").strip()
    if not src:
        return []
    out: List[str] = []
    for line in src.splitlines():
        line = line.strip()
        if not line:
            continue
        for sent in _SENTENCE_SPLIT_RE.split(line):
            for semi in re.split(r";+\s*", sent):
                chunk = semi.strip()
                if len(chunk) > 8:
                    out.extend(_atomize_step_body(chunk))
    return out


def _sentence_is_step(sentence: str) -> bool:
    s = (sentence or "").strip()
    if not s:
        return False
    if re.match(r"^(?:регламент|название|процесс)\s*[:—–-]", s, re.I):
        return False
    if re.match(r"^порядок\s+", s, re.I) and not _CASE_RE.search(s) and not _RETURN_RE.search(s):
        if not any(_is_verb(w.strip(".,;:")) for w in s.split()[:8]):
            return False
    if _CASE_RE.search(s) or _BACK_KW.search(s) or _PAR_WORD_RE.search(s):
        return True
    if _find_role(s)[0] or _match_inverted_role(s):
        return True
    return any(_is_verb(w.strip(".,;:—–-")) for w in s.split()[:14])


def _segment_unnumbered_prose(text: str) -> str:
    """Сплошной текст без «1.» / «I.» → виртуально нумерованные шаги по предложениям."""
    src = (text or "").strip()
    if not src or _count_explicit_numbers(src) >= 2:
        return src
    if re.search(r"(?m)^[а-яёa-z]\)\s+", src, re.I):
        return src
    headers: List[str] = []
    body: List[str] = []
    for line in src.splitlines():
        s = line.strip()
        if not s:
            continue
        if re.match(r"^(?:регламент|название|процесс)\s*[:—–-]", s, re.I):
            headers.append(s)
            continue
        if re.match(r"^(?:целев\w+\s+(?:срок|sla)[^:—–]*|sla)\s*[:—–-]", s, re.I):
            headers.append(s)
            continue
        body.append(s)
    if body:
        first = body[0]
        looks_title = (
            not re.search(r"[.!?]$", first)
            and len(first.split()) <= 18
            and not _CASE_RE.search(first)
            and not _RETURN_RE.search(first)
            and not any(_is_verb(w.strip(".,;:")) for w in first.split())
        )
        if looks_title:
            if not headers:
                headers.append("Регламент: " + first.rstrip(" ."))
            body = body[1:]
    sentences: List[str] = []
    for chunk in body or ([src] if not headers else []):
        sentences.extend(_split_prose_sentences(chunk))
    steps = [s for s in sentences if _sentence_is_step(s)]
    if len(steps) < 2:
        return src
    if not headers:
        leftover = next((s for s in sentences if s not in steps), "")
        if leftover:
            headers.append("Регламент: " + leftover.rstrip(" ."))
    numbered = []
    for i, s in enumerate(steps, 1):
        body_s = s if s.endswith((".", "!", "?")) else s.rstrip(".") + "."
        numbered.append(f"{i}. {body_s}")
    return "\n".join([*headers, *numbered]).strip()


_WML = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_style_numbering(zf: Any) -> Dict[str, Tuple[str, str]]:
    """styleId → (numId, ilvl). Автонумерация часто висит на стиле, а не на абзаце."""
    direct: Dict[str, Tuple[str, str]] = {}
    based: Dict[str, str] = {}
    if "word/styles.xml" not in zf.namelist():
        return direct
    root = _ET.fromstring(zf.read("word/styles.xml"))
    for style in root.findall(f"{_WML}style"):
        sid = style.get(f"{_WML}styleId") or ""
        based_el = style.find(f"{_WML}basedOn")
        if based_el is not None and based_el.get(f"{_WML}val"):
            based[sid] = based_el.get(f"{_WML}val") or ""
        p_pr = style.find(f"{_WML}pPr")
        num_pr = p_pr.find(f"{_WML}numPr") if p_pr is not None else None
        if num_pr is None:
            continue
        num_id_el = num_pr.find(f"{_WML}numId")
        if num_id_el is None or not num_id_el.get(f"{_WML}val"):
            continue
        ilvl_el = num_pr.find(f"{_WML}ilvl")
        direct[sid] = (
            num_id_el.get(f"{_WML}val") or "",
            (ilvl_el.get(f"{_WML}val") if ilvl_el is not None else None) or "0",
        )

    resolved: Dict[str, Tuple[str, str]] = {}

    def walk(sid: str, seen: Tuple[str, ...]) -> Optional[Tuple[str, str]]:
        if sid in resolved:
            return resolved[sid]
        if sid in direct:
            resolved[sid] = direct[sid]
            return direct[sid]
        parent = based.get(sid)
        if not parent or parent in seen:
            return None
        found = walk(parent, seen + (sid,))
        if found:
            resolved[sid] = found
        return found

    for sid in set(direct) | set(based):
        walk(sid, ())
    return resolved


def _docx_numbering_maps(
    data: bytes,
) -> Tuple[Dict[str, str], Dict[Tuple[str, str], Tuple[int, str]], Dict[str, Tuple[str, str]]]:
    """numId → abstractNumId, (abstract, ilvl) → (start, fmt), styleId → (numId, ilvl)."""
    import zipfile

    num_to_abs: Dict[str, str] = {}
    levels: Dict[Tuple[str, str], Tuple[int, str]] = {}
    styles: Dict[str, Tuple[str, str]] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            styles = _docx_style_numbering(zf)
            if "word/numbering.xml" not in zf.namelist():
                return num_to_abs, levels, styles
            root = _ET.fromstring(zf.read("word/numbering.xml"))
    except Exception:  # noqa: BLE001
        return num_to_abs, levels, styles
    for abstract in root.findall(f"{_WML}abstractNum"):
        aid = abstract.get(f"{_WML}abstractNumId") or ""
        for lvl in abstract.findall(f"{_WML}lvl"):
            ilvl = lvl.get(f"{_WML}ilvl") or "0"
            start_el = lvl.find(f"{_WML}start")
            fmt_el = lvl.find(f"{_WML}numFmt")
            start = int(start_el.get(f"{_WML}val") or "1") if start_el is not None else 1
            fmt = (fmt_el.get(f"{_WML}val") if fmt_el is not None else None) or "decimal"
            levels[(aid, ilvl)] = (start, fmt)
    for num in root.findall(f"{_WML}num"):
        nid = num.get(f"{_WML}numId") or ""
        abs_el = num.find(f"{_WML}abstractNumId")
        if nid and abs_el is not None:
            num_to_abs[nid] = abs_el.get(f"{_WML}val") or ""
    return num_to_abs, levels, styles


def _docx_num_pr(paragraph: Any, styles: Dict[str, Tuple[str, str]]) -> Optional[Tuple[str, str]]:
    p_pr = getattr(paragraph._p, "pPr", None)
    if p_pr is not None and p_pr.numPr is not None:
        num_pr = p_pr.numPr
        num_id = num_pr.numId.val if num_pr.numId is not None else None
        if num_id is not None and str(num_id) != "0":
            ilvl = num_pr.ilvl.val if num_pr.ilvl is not None and num_pr.ilvl.val is not None else 0
            return str(num_id), str(ilvl)
    style = getattr(paragraph, "style", None)
    sid = getattr(style, "style_id", None) if style is not None else None
    found = styles.get(sid or "")
    if found and found[0] != "0":
        return found
    return None


def _docx_next_label(
    num_id: str,
    ilvl: str,
    counters: Dict[Tuple[str, str], int],
    num_to_abs: Dict[str, str],
    levels: Dict[Tuple[str, str], Tuple[int, str]],
) -> str:
    key = (num_id, ilvl)
    for deeper in [k for k in counters if k[0] == num_id and int(k[1]) > int(ilvl)]:
        counters.pop(deeper, None)
    start, _fmt = levels.get((num_to_abs.get(num_id, ""), ilvl), (1, "decimal"))
    if key not in counters:
        counters[key] = start
    else:
        counters[key] += 1
    return f"{counters[key]}. "


def read_docx_regulation(data: bytes) -> Tuple[Optional[str], Optional[str]]:
    """Текст .docx. Автонумерация Word не лежит в абзаце — номер берётся из numbering и пишется перед пунктом."""
    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError:
        return None, "Для чтения .docx установите пакет: `pip install python-docx>=1.0.0`"
    try:
        doc = Document(io.BytesIO(data))
        num_to_abs, levels, styles = _docx_numbering_maps(data)
        counters: Dict[Tuple[str, str], int] = {}
        parts: List[str] = []
        for paragraph in doc.paragraphs:
            raw = (paragraph.text or "").strip()
            if not raw:
                continue
            num = _docx_num_pr(paragraph, styles)
            if num and not re.match(r"^\d+[.)]\s+", raw):
                raw = _docx_next_label(num[0], num[1], counters, num_to_abs, levels) + raw
            parts.append(raw)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if not text:
            return None, "В файле .docx не найден текстовый слой (пустые абзацы и таблицы)."
        return text, None
    except Exception as exc:  # noqa: BLE001
        return None, f"Не удалось прочитать .docx: {type(exc).__name__}: {exc}"


def normalize_regulation(text: str) -> str:
    """Чистит текст, скопированный из PDF/Word, до вида «шаг на строке».

    - убирает колонтитулы («Стр. 5 из 17») и строки из одних номеров страниц;
    - склеивает переносы слов («регистри-\\nрует») и ссылки, разорванные строкой («п.\\n3.2.5»);
    - многоуровневую нумерацию «3.2.1.» переводит в сквозную «1.», «2.» … вместе со ссылками «п. 3.2.5».
    """
    text = text.replace("\r\n", "\n").replace("­", "")
    text = re.sub(r"(\w)-\n\s*(\w)", r"\1\2", text)
    text = re.sub(r"(\bп(?:\.|ункт\w*)|\bшаг\w*)\s*\n\s*(?=\d)", r"\1 ", text, flags=re.I)
    lines = []
    for line in text.split("\n"):
        line = _PAGE_MARK_RE.sub("", line) if _PAGE_MARK_RE.search(line) and len(line.strip()) < 120 else line
        if re.search(r"\S\s{8,}\S", line):  # колонтитул PDF: «ПАО …        СТО 123-2023»
            continue
        if not re.fullmatch(r"\s*[-–—]?\s*\d{1,3}\s*[-–—]?\s*", line):  # номер страницы
            lines.append(line.rstrip())
    text = "\n".join(lines)

    numbers = _MULTI_NUM_RE.findall(text)
    depth = max((n.count(".") for n in numbers), default=0)
    steps = [n for n in numbers if n.count(".") == depth]  # «3.2 Раздел» — заголовок, «3.2.1.» — шаг
    if len(steps) >= 2:
        mapping = {num: str(i) for i, num in enumerate(dict.fromkeys(steps), start=1)}
        text = _MULTI_NUM_RE.sub(lambda m: mapping[m.group(1)] + ". " if m.group(1) in mapping else m.group(0), text)
        for num in sorted(mapping, key=len, reverse=True):
            text = re.sub(r"(?<![\d.])" + re.escape(num) + r"(?![\d])", mapping[num], text)
    text = _promote_lettered_parallels(text)
    text = _expand_corporate_numbering(text)
    text = _segment_unnumbered_prose(text)
    return _atomize_numbered_lines(text)


def _to_hours(value: str, unit: str) -> float:
    num = float(value.replace(",", "."))
    unit = unit.lower()
    if unit.startswith("мин"):
        return num / 60.0
    if unit.startswith("час") or unit == "ч":
        return num
    if unit.startswith("сут") or unit.startswith("календарн"):
        return num * 24.0
    return num * 8.0  # рабочие дни


def _infinitive(word: str) -> str:
    low = word.lower()
    if low in VERB_MAP:
        return VERB_MAP[low]
    if len(low) >= 6:
        for suffix, repl in _VERB_SUFFIXES:
            if low.endswith(suffix):
                return low[: -len(suffix)] + repl
    return word


def _task_title(text: str) -> str:
    text = re.sub(r"\s+", " ", text).strip(" .;,:—–-")
    text = _CONNECTORS_RE.sub("", text).strip(" .;,:—–-")
    text = _INTRO_CLAUSE_RE.sub("", text).strip(" .;,:—–-")
    if not text:
        return "Выполнить действие"
    words = [_infinitive(w) if re.fullmatch(r"[А-Яа-яЁё]+", w) else w for w in text.split(" ")]
    title = " ".join(words)
    title = title[:1].upper() + title[1:] if title else "Выполнить действие"
    if len(title) > TITLE_MAX:
        cut = title[:TITLE_MAX].rsplit(" ", 1)[0]
        title = cut.rstrip(",;:—–- ") + "…"
    return title


TITLE_MAX = 110  # длиннее — обрезаем по слову; блок задачи растёт по высоте под текст
_CONNECTORS_RE = re.compile(
    r"^(?:затем|далее|потом|после этого|после проверки|также|при этом|"
    r"при согласовании|в случае согласования)[,\s]+",
    re.I,
)
_ADVERBS = {"автоматически", "затем", "далее", "также", "самостоятельно", "обязательно", "незамедлительно", "оперативно"}
_REL_PRONOUNS = {"который", "которая", "которое", "которые", "которого", "которой", "которым", "которыми"}


def _is_verb(word: str) -> bool:
    low = word.lower().strip(",.;")
    return low in VERB_MAP or (len(low) >= 6 and low.endswith(("ает", "яет", "ует")))


def _subject_before_verb(text: str) -> Optional[Tuple[str, int]]:
    """Подлежащее — слова перед первым глаголом («Начальник смены» оценивает …)."""
    m = _CONNECTORS_RE.match(text)
    offset = m.end() if m else 0
    tokens = list(re.finditer(r"\S+", text[offset:]))
    for i, tok in enumerate(tokens[:6]):
        if i >= 1 and _is_verb(tok.group()):
            words = [t.group() for t in tokens[:i]]
            while words and words[-1].lower().strip(",") in _ADVERBS:
                words.pop()
            if not words or len(words) > 4 or any(re.search(r"[\d:;()«»,]", w) for w in words):
                return None
            if all(w.lower().strip(",") in _REL_PRONOUNS for w in words):
                return None
            return " ".join(words), offset + tok.start()
        if re.search(r"[.;:]$", tok.group()):
            break
    return None


def _canonical_role(subject: str) -> str:
    for regex, name in _ROLE_RE:
        if regex.search(subject):
            return name
    return subject[:1].upper() + subject[1:]


_INVERSE_VERB_RE = re.compile(
    r"\s+(?P<verb>выполняет|выполняют|проводит|проводят|осуществляет|осуществляют|"
    r"ведёт|ведет|ведут|производит|производят|"
    r"готовит|готовят|направляет|направляют|рассматривает|рассматривают|"
    r"проверяет|проверяют|вносит|вносят|принимает|принимают)\s+",
    re.I,
)


def _match_inverted_role(text: str) -> Optional[Tuple[str, str, str]]:
    """Инверсия ТЭК: «Осмотр оборудования проводит начальник смены» → (роль, действие, глагол)."""
    src = (text or "").strip()
    if not src:
        return None
    m = _INVERSE_VERB_RE.search(src)
    if not m or m.start() < 3:
        return None
    action = src[: m.start()].strip(" .,;:—–-")
    tail = src[m.end() :].strip()
    if not action or not tail or len(action.split()) > 12:
        return None
    # «Ремонтная бригада готовит СИЗ» — обычный порядок, не инверсия.
    # Иначе объект «СИЗ» становится дорожкой, а бригада — названием задачи.
    if any((mm := rx.search(action)) and mm.start() <= 1 for rx, _ in _ROLE_RE):
        return None
    head = action.split()[0].lower().strip("«»\"'")
    if any(rx.search(head) for rx, _ in _ROLE_RE):
        return None
    role: Optional[str] = None
    for regex, name in _ROLE_RE:
        mm = regex.search(tail[:90])
        if mm and mm.start() <= 3:
            role = name
            break
    if role is None:
        rm = re.match(r"([А-ЯЁ][А-Яа-яЁё\-]+(?:\s+[а-яёА-ЯЁ\-]{3,24}){0,3})", tail)
        if not rm:
            return None
        cand = rm.group(1).strip()
        if len(cand.split()) > 5 or re.search(r"[\d:;()]", cand):
            return None
        role = _canonical_role(cand)
    return role, action, m.group("verb")


def _split_joint_role(text: str) -> Optional[Tuple[str, str]]:
    """«Диспетчер совместно с бригадой …» → первая роль + полный текст действия."""
    src = (text or "").strip()
    m = _JOINT_LINK_RE.match(src)
    if not m:
        return None
    head = m.group("head").strip(" ,")
    role: Optional[str] = None
    for regex, name in _ROLE_RE:
        mm = regex.search(head)
        if mm and mm.start() <= 2:
            role = name
            break
    if not role:
        return None
    action = src[len(head) :].strip(" ,:;.—–-")
    return role, action


def _find_role(text: str) -> Tuple[Optional[str], int, int]:
    """Роль-исполнитель шага: (название, начало, конец вырезаемого префикса)."""
    m = re.match(r"^([А-ЯЁ][А-Яа-яЁё\- ]{2,45}?)\s*[:—–]\s+", text)
    if m and len(m.group(1).split()) <= 5 and not re.match(r"(?i)^(если|параллельно|одновременно)", m.group(1)):
        return _canonical_role(m.group(1).strip()), 0, m.end()
    inverted = _match_inverted_role(text)
    if inverted:
        role, _action, _verb = inverted
        vm = _INVERSE_VERB_RE.search(text.strip())
        tail_start = vm.end() if vm else 0
        return role, tail_start, len(text.strip())
    subject = _subject_before_verb(text[:120])
    if subject:
        return _canonical_role(subject[0]), 0, subject[1]
    best: Optional[Tuple[int, int, str]] = None
    for regex, name in _ROLE_RE:
        mm = regex.search(text[:90])
        if mm and (best is None or mm.start() < best[0]):
            best = (mm.start(), mm.end(), name)
    if best:
        return best[2], best[0], best[1]
    return None, -1, -1


# --------------------------------------------------------------------------- #
# ИТ-ландшафт и документооборот: какие документы и системы фигурируют в регламенте
# --------------------------------------------------------------------------- #
_ENDINGS = r"(?:а|у|ом|е|ы|и|ой|ов|ам|ами|ах)?"
# (regex, каноническое название); более конкретные шаблоны стоят раньше общих и «съедают» свой фрагмент.
ARTIFACT_PATTERNS: List[Tuple[str, str]] = [
    (r"наряд[\s-]*допуск\w*", "Наряд-допуск"),
    (r"дефектн\w+\s+ведомост\w+", "Дефектная ведомость"),
    (r"техническ\w+\s+услови\w+|(?-i:\bТУ\b)", "Технические условия (ТУ)"),
    (r"техническ\w+\s+задани\w+|(?-i:\bТЗ\b)", "Техническое задание (ТЗ)"),
    (r"закупочн\w+\s+документаци\w+", "Закупочная документация"),
    (r"\bакт" + _ENDINGS + r"\s+о(?:б)?\s+[а-яё]+(?:ом|ем|ой|ей|ых)\s+[а-яё]{4,}", "Акт {tail}"),
    (r"\bакт" + _ENDINGS + r"\s+о(?:б)?\s+[а-яё]{4,}", "Акт {tail}"),
    (r"\bакт" + _ENDINGS + r"\s+(?:выполненных\s+работ|при[её]мки(?:-передачи)?|сдачи-при[её]мки|осмотра|допуска)", "Акт {tail}"),
    (r"\bакт" + _ENDINGS + r"\b", "Акт"),
    (r"\bдоговор" + _ENDINGS + r"\b", "Договор"),
    (r"\bзаяв(?:к(?:а|и|е|у|ой|ам|ами|ах)|ок|лени\w+)\b", "Заявка"),
    (r"\bсмет" + _ENDINGS + r"\b", "Смета"),
    (r"\bсч[её]т" + _ENDINGS + r"\b", "Счёт"),
    (r"\bпротокол" + _ENDINGS + r"\b", "Протокол"),
    (r"\bприказ" + _ENDINGS + r"\b|\bраспоряжени\w+", "Приказ / распоряжение"),
    (r"экспертн\w+\s+заключени\w+|заключени\w+\s+экспертизы", "Экспертное заключение"),
    (r"\bизвещени\w+", "Извещение о закупке"),
    (r"график\w*\s+ремонт\w*", "График ремонтов"),
    (r"\bуведомлени\w+", "Уведомление"),
    (r"пакет\w*\s+документ\w*", "Пакет документов"),
]
SYSTEM_PATTERNS: List[Tuple[str, str]] = [
    (r"(?-i:\bАСУ\s*ТП\b)|(?-i:\bSCADA\b)|\bскад[аеу]\b|(?-i:\bАСДУ\b)|(?-i:\bОИК\b)|телемеханик\w*", "АСУ ТП / SCADA"),
    (r"оперативн\w+\s+журнал\w*", "Оперативный журнал"),
    (r"\bв\s+журнал\w*", "Журнал"),
    (r"(?-i:\bCRM\b)", "CRM"),
    (r"электронн\w+\s+площадк\w+|(?-i:\bЭТП\b)", "Электронная площадка"),
    (r"биллинг\w*", "Биллинг"),
    (r"(?-i:\b1[СC]\b)|(?-i:\bSAP\b)|(?-i:\bERP\b)", "1С / SAP / ERP"),
    (r"(?-i:\bСЭД\b)|электронн\w+\s+документооборот\w*|систем\w+\s+электронного\s+документооборота", "СЭД (электронный документооборот)"),
    (r"личн\w+\s+кабинет\w*", "Личный кабинет"),
    (r"информационн\w+\s+систем\w+|автоматизированн\w+\s+систем\w+|(?-i:\bА?ИС\b)", "Информационная система (ИС / АС)"),
]
_ARTIFACT_RE = [(re.compile(p, re.I), n) for p, n in ARTIFACT_PATTERNS]
_SYSTEM_RE_LIST = [(re.compile(p, re.I), n) for p, n in SYSTEM_PATTERNS]


def _find_named(text: str, patterns: List[Tuple["re.Pattern[str]", str]]) -> List[str]:
    """Канонические названия всех найденных объектов (в порядке появления, без повторов)."""
    taken: List[Tuple[int, int]] = []
    found: List[Tuple[int, str]] = []
    for regex, name in patterns:
        for m in regex.finditer(text):
            span = m.span()
            if any(span[0] < e and s < span[1] for s, e in taken):  # уже учтено более конкретным шаблоном
                continue
            taken.append(span)
            if "{tail}" in name:
                head = re.match(r"\bакт\w*\s+", m.group(0), re.I)
                tail = m.group(0)[head.end():] if head else ""
                label = name.replace("{tail}", tail.strip().lower())
            else:
                label = name
            found.append((span[0], label))
    found.sort()
    result: List[str] = []
    for _, label in found:
        if label not in result:
            result.append(label)
    return result


def extract_artifacts(text: str) -> List[str]:
    """Документы и артефакты процесса: наряд-допуск, дефектная ведомость, ТУ, договор, заявка, акт…"""
    return _find_named(text, _ARTIFACT_RE)


def extract_it_systems(text: str) -> List[str]:
    """ИТ-системы: АСУ ТП / SCADA, оперативный журнал, CRM, электронная площадка, биллинг, 1С / SAP…"""
    return _find_named(text, _SYSTEM_RE_LIST)


def aggregate_landscape(steps: List["Step"], attr: str) -> List[Dict[str, Any]]:
    """Сводит найденное по шагам в список {name, mentions, steps, roles}, самые упоминаемые — первыми."""
    index: Dict[str, Dict[str, Any]] = {}
    for step in steps:
        for name in getattr(step, attr):
            item = index.setdefault(name, {"name": name, "mentions": 0, "steps": [], "roles": []})
            item["mentions"] += 1
            if step.num not in item["steps"]:
                item["steps"].append(step.num)
            if step.role and step.role not in item["roles"]:
                item["roles"].append(step.role)
    return sorted(index.values(), key=lambda it: (-it["mentions"], it["steps"][0] if it["steps"] else 0))


@dataclass
class Decision:
    yes_label: str
    no_label: str
    yes_ref: Optional[int] = None
    no_ref: Optional[int] = None
    yes_end: bool = False
    no_end: bool = False
    no_back: bool = False
    has_else: bool = False
    back_clause: str = ""


@dataclass
class Step:
    idx: int
    num: int
    role: str = ""
    title: str = ""
    parallel: bool = False
    fork_parallel: bool = False
    stage: Optional[str] = None
    hours: Optional[float] = None
    system: bool = False
    decision: Optional[Decision] = None
    back_ref: Optional[int] = None  # «на п.N» без развилки: одно ребро, шаг остаётся обычным
    action: bool = True  # есть ли собственное действие перед шлюзом
    artifacts: List[str] = field(default_factory=list)  # документы шага: наряд-допуск, акт, договор…
    systems: List[str] = field(default_factory=list)  # ИТ-системы шага: АСУ ТП, CRM, 1С…


@dataclass
class Block:
    kind: str  # step | decision | parallel | subprocess
    steps: List[Step]
    name: str = ""

    @property
    def role(self) -> str:
        return self.steps[0].role


@dataclass
class ParsedRegulation:
    title: str
    sla_hours: Optional[float]
    steps: List[Step] = field(default_factory=list)
    roles: List[str] = field(default_factory=list)
    artifacts: List[Dict[str, Any]] = field(default_factory=list)  # сводка по документам процесса
    it_systems: List[Dict[str, Any]] = field(default_factory=list)  # сводка по ИТ-системам процесса


def _derive_no_label(clause: str) -> str:
    low = clause.lower()
    if "замечан" in low:
        return "Замечания"
    if "отказ" in low or "отклон" in low:
        return "Отказ"
    if _BACK_KW.search(low):
        return "На доработку"
    return "Иначе"


def _explicit_alternative(src: str) -> Optional[re.Match]:
    """«либо … либо» и «или … или/иначе» — развилка. Одиночное «или» шлюзом не становится."""
    for word in ("либо", "или"):
        m = re.search(rf"\b{word}\b", src or "", re.I)
        if not m:
            continue
        rest = src[m.end():]
        if re.search(rf"\b(?:{word}|иначе|в противном случае)\b", rest, re.I):
            return m
    return None


def _parse_decision(text: str) -> Tuple[str, Optional[Decision]]:
    src = text or ""
    m = _CASE_RE.search(src)
    if m is None:
        m = _explicit_alternative(src)
    if m:
        action = src[: m.start()].strip(" .;,—–-")
        rest = src[m.end():]
        opener = m.group(0).lower()
        if opener in ("либо", "или"):
            else_m = re.search(r"[,;.]?\s*\b(?:иначе|в противном случае|либо|или)\b[,:]?", rest, re.I)
        else:
            else_m = re.search(r"[,;.]?\s*\b(?:иначе|в противном случае)\b[,:]?", rest, re.I)
        yes_part = rest[: else_m.start()] if else_m else rest
        no_part = rest[else_m.end():] if else_m else ""
        pieces = re.split(r"\s+[—–-]\s+|,\s+|:\s+", yes_part.strip(), maxsplit=1)
        cond = pieces[0].strip(" ,.;")
        yes_clause = pieces[1] if len(pieces) > 1 else ""
        quoted_yes = re.search(r"«([^»]+)»", cond)
        quoted_no = re.search(r"«([^»]+)»", no_part)
        yes_label = (quoted_yes.group(1) if quoted_yes else cond).strip()
        yes_label = (yes_label[:1].upper() + yes_label[1:]) if yes_label else "Да"
        no_label = quoted_no.group(1).strip() if quoted_no else _derive_no_label(no_part)
        no_label = (no_label[:1].upper() + no_label[1:]) if no_label else "Иначе"
        yes_ref = _REF_RE.search(yes_clause)
        no_ref_m = _REF_RE.search(no_part) if no_part else None
        back_src = no_part
        yes_end = bool(_END_KW.search(yes_clause))
        no_end = bool(_END_KW.search(no_part)) and not no_ref_m
        no_back = bool(_BACK_KW.search(no_part)) and not no_ref_m
        if not else_m and _BACK_KW.search(yes_part) and not _END_KW.search(yes_part):
            back_src = yes_part
            no_ref_m = no_ref_m or _REF_RE.search(yes_part)
            no_back = not bool(no_ref_m)
            no_label = _derive_no_label(yes_part)
            no_label = (no_label[:1].upper() + no_label[1:]) if no_label else "На доработку"
            low = yes_part.lower()
            if "замечан" in low:
                yes_label = "Замечаний нет"
            elif "отказ" in low or "отклон" in low:
                yes_label = "Согласовано"
            else:
                yes_label = "Продолжить"
            yes_end = False
            yes_ref = None
        if yes_end and not else_m:
            low_end = f"{cond} {yes_clause}".lower()
            if "отказ" in low_end or "отклон" in low_end:
                yes_label = "Отказ"
                no_label = "Согласовано"
        decision = Decision(
            yes_label=yes_label[:40],
            no_label=no_label[:40],
            yes_ref=int(yes_ref.group(1)) if yes_ref else None,
            no_ref=int(no_ref_m.group(1)) if no_ref_m else None,
            yes_end=yes_end,
            no_end=no_end,
            no_back=bool(no_back or (no_ref_m and _BACK_KW.search(back_src or ""))),
            has_else=bool(else_m),
            back_clause=back_src or "",
        )
        if no_ref_m and _BACK_KW.search(back_src or yes_part):
            decision.no_back = True
            decision.no_ref = int(no_ref_m.group(1))
        return action, decision

    return src, None


def _has_explicit_return(text: str) -> bool:
    """Одно обратное ребро: «на п.N» при глаголе возврата или «иначе … назад».

    Голые «возвращает», «вернуть» и «возврат» без номера пункта — обычный шаг, не цикл.
    """
    src = text or ""
    if _NAZAD_RETURN_RE.search(src):
        return True
    return bool(_EXPLICIT_RETURN_RE.search(src) and _REF_RE.search(src))


def _bare_return_ref(text: str) -> Optional[int]:
    """«Вернуть на п.N» без «если»: номер цели. Развилку это не создаёт."""
    src = text or ""
    if _CASE_RE.search(src) or _NAZAD_RETURN_RE.search(src):
        return None
    if not (_EXPLICIT_RETURN_RE.search(src) and _REF_RE.search(src)):
        return None
    ref = _REF_RE.search(src)
    return int(ref.group(1)) if ref else None


def _decision_points_back(decision: Optional[Decision], step_num: int) -> bool:
    if decision is None:
        return False
    if decision.no_back:
        return True
    if decision.no_ref is not None and decision.no_ref < step_num:
        return True
    if decision.yes_ref is not None and decision.yes_ref < step_num:
        return True
    return False


def _ensure_return_decision(text: str, decision: Optional[Decision], step_num: int) -> Optional[Decision]:
    """«Если … назад» дополняет уже найденную развилку. Голый глагол возврата шаг не превращает."""
    if not _has_explicit_return(text) or _decision_points_back(decision, step_num):
        return decision
    if decision is None:
        return None
    decision.no_back = True
    if not decision.back_clause:
        decision.back_clause = text
    if not decision.no_label or decision.no_label == "Иначе":
        decision.no_label = "На доработку"
    return decision


def _numbered_chunks(text: str) -> List[str]:
    raw = normalize_regulation(text or "")
    chunks: List[str] = []
    current = ""
    for line in raw.splitlines():
        if re.match(r"^\s*\d+[.)]\s+\S", line):
            if current:
                chunks.append(current)
            current = line
        elif current:
            current = f"{current} {line.strip()}"
    if current:
        chunks.append(current)
    return chunks


def _code_for_sandbox(code: str) -> str:
    """В exec попадает только код DIAGRAM. Дорисовка возвратов вызывается снаружи."""
    if "_complete_return_edges" not in (code or ""):
        return code or ""
    kept = [line for line in (code or "").splitlines() if "_complete_return_edges" not in line]
    return "\n".join(kept)


def _explicit_return_count(text: str) -> int:
    """Сколько пунктов регламента содержат явный возврат. Каждый такой пункт — одно ребро."""
    chunks = _numbered_chunks(text)
    if not chunks and _has_explicit_return(normalize_regulation(text or "")):
        return 1
    return sum(1 for chunk in chunks if _has_explicit_return(chunk))


def _return_diagram_not_ready(expected: int, actual: int) -> str:
    """После дорисовки рёбер всё ещё меньше, чем фраз: схему не отдаём как готовую."""
    return (
        f"Схема не готова: в регламенте {expected} явных возврата, "
        f"в графе обратных рёбер {actual}."
    )


def _model_return_incomplete(err: str) -> bool:
    """Дыра возврата в уже собранной схеме модели. Эмулятор вместо неё не подставляем."""
    return "Схема не готова" in (err or "") and "явных возврата" in (err or "")


def _return_phrase_specs(text: str) -> List[Dict[str, Any]]:
    """Каждая фраза «вернуть на п.N» или «иначе … назад» — источник, цель и подпись ветки."""
    raw = normalize_regulation(text or "")
    if not raw.strip() or _explicit_return_count(raw) <= 0:
        return []
    try:
        parsed = parse_regulation(raw)
    except Exception:  # noqa: BLE001
        return []
    by_num = {step.num: step for step in parsed.steps}

    def _step_for(chunk: str, num: Optional[int]) -> Any:
        if num is not None and num in by_num:
            return by_num[num]
        best, score = None, 0
        for step in parsed.steps:
            got = _sla_title_score(step.title or "", chunk)
            if got > score:
                best, score = step, got
        return best if score >= 2 else None

    specs: List[Dict[str, Any]] = []
    chunks = _numbered_chunks(raw)
    if not chunks and _has_explicit_return(raw):
        chunks = [raw]
    for chunk in chunks:
        if not _has_explicit_return(chunk):
            continue
        num_m = re.match(r"\s*(\d+)", chunk)
        num = int(num_m.group(1)) if num_m else None
        step = _step_for(chunk, num)
        target_num: Optional[int] = None
        label = "Возврат"
        if step and step.decision:
            decision = step.decision
            if decision.no_ref is not None and (decision.no_back or decision.no_ref < step.num):
                target_num = decision.no_ref
            elif decision.yes_ref is not None and decision.yes_ref < step.num:
                target_num = decision.yes_ref
            label = (decision.no_label or label)[:40]
        if target_num is None and step is not None and step.back_ref and step.back_ref < step.num:
            target_num = step.back_ref
            label = "Возврат"
        if target_num is None:
            ref_m = _REF_RE.search(chunk)
            if ref_m and step and int(ref_m.group(1)) < step.num:
                target_num = int(ref_m.group(1))
        target = by_num.get(target_num) if target_num is not None else None
        specs.append(
            {
                "source_num": step.num if step else num,
                "target_num": target.num if target else target_num,
                "label": label or "Возврат",
                "source_title": (step.title if step else "") or "",
                "target_title": (target.title if target else "") or "",
            }
        )
    return specs


def _bind_step_nodes(diagram: BPMNDiagramBuilder, titles: Dict[int, str]) -> Dict[int, Any]:
    """Шаг регламента → задача с именем из глагола, не контейнер подпроцесса."""
    task_kinds = WORK_KINDS | {"manualTask", "serviceTask", "sendTask", "receiveTask", "businessRuleTask"}
    work = [node for node in diagram.nodes.values() if node.kind in task_kinds]
    scored: List[Tuple[int, int, str]] = []
    for num, title in titles.items():
        if not title:
            continue
        for node in work:
            score = _sla_title_score(title, node.name or "")
            if score >= 2:
                scored.append((score, num, node.id))
    scored.sort(key=lambda item: (-item[0], item[1]))
    used_steps: set = set()
    used_nodes: set = set()
    bound: Dict[int, Any] = {}
    for _score, num, node_id in scored:
        if num in used_steps or node_id in used_nodes:
            continue
        bound[num] = diagram.nodes[node_id]
        used_steps.add(num)
        used_nodes.add(node_id)
    return bound


def _reachable_from(diagram: BPMNDiagramBuilder, start_id: str) -> set:
    seen: set = set()
    stack = [start_id]
    while stack:
        current = stack.pop()
        if current in seen:
            continue
        seen.add(current)
        for link in diagram.links:
            if link.source_id == current and link.target_id not in seen:
                stack.append(link.target_id)
    seen.discard(start_id)
    return seen


def _owner_holds(diagram: BPMNDiagramBuilder, node_id: str, owner_id: str) -> bool:
    node = diagram.nodes.get(node_id)
    seen: set = set()
    while node and node.owner_id and node.owner_id not in seen:
        if node.owner_id == owner_id:
            return True
        seen.add(node.owner_id)
        node = diagram.nodes.get(node.owner_id)
    return False


def _back_link_list(diagram: BPMNDiagramBuilder) -> List[Any]:
    found: List[Any] = []
    owners = [pool.process_id for pool in diagram.pools.values()][:1] or [diagram.process_id]
    owners += [node.id for node in diagram.nodes.values() if node.kind == "subProcess"]
    for owner in owners:
        nodes = diagram._children(owner)
        if not nodes:
            continue
        start_id = diagram._scope_start_end(owner)[0]
        back = diagram._back_links(nodes, start_id)
        ids = {node.id for node in nodes}
        for link in diagram._scope_links(ids):
            if link.id in back:
                found.append(link)
    return found


def _gateway_after(diagram: BPMNDiagramBuilder, node: Any) -> Any:
    outs = [link for link in diagram.links if link.source_id == node.id]
    if len(outs) != 1:
        return node
    nxt = diagram.nodes.get(outs[0].target_id)
    if nxt and nxt.kind in ("exclusiveGateway", "inclusiveGateway") and nxt.owner_id == node.owner_id:
        return nxt
    return node


def _pick_return_source(diagram: BPMNDiagramBuilder, target: Any, source: Optional[Any], reachable: set) -> Optional[Any]:
    """Узел того же процесса, из которого ребро на цель замыкает цикл."""
    task_kinds = WORK_KINDS | {"manualTask", "serviceTask", "sendTask", "receiveTask", "businessRuleTask"}

    def usable(node: Any) -> bool:
        if node is None or node.id == target.id or node.id not in reachable:
            return False
        if node.owner_id != target.owner_id or node.kind in ("startEvent", "endEvent", "subProcess", "parallelGateway"):
            return False
        return not any(link.source_id == node.id and link.target_id == target.id for link in diagram.links)

    ordered: List[Any] = []
    if source is not None:
        gate = _gateway_after(diagram, source)
        ordered.append(gate)
        if gate is not source:
            ordered.append(source)
    ordered.extend(
        node for node in diagram.nodes.values()
        if node.kind in ("exclusiveGateway", "inclusiveGateway")
    )
    ordered.extend(node for node in diagram.nodes.values() if node.kind in task_kinds)
    for node in ordered:
        if usable(node):
            return node
    return None


def _complete_return_edges(diagram: BPMNDiagramBuilder, regulation_text: str) -> None:
    """Недостающие явные возвраты дорисовываются в эту же схему, до analyze_bottlenecks.

    Цель ребра — задача с глаголом, в том числе внутри подпроцесса.
    Чужой граф эмулятора сюда не подставляется.
    """
    specs = _return_phrase_specs(regulation_text)
    if not specs:
        return
    titles: Dict[int, str] = {}
    for spec in specs:
        if spec.get("source_num") is not None and spec.get("source_title"):
            titles[int(spec["source_num"])] = str(spec["source_title"])
        if spec.get("target_num") is not None and spec.get("target_title"):
            titles[int(spec["target_num"])] = str(spec["target_title"])
    bound = _bind_step_nodes(diagram, titles)
    claimed: set = set()

    def _refresh() -> List[Any]:
        return _back_link_list(diagram)

    back = _refresh()

    def _covers(link: Any, target: Any, title: str) -> bool:
        node = diagram.nodes.get(link.target_id)
        if node is None:
            return False
        if target is not None and link.target_id == target.id:
            return True
        if title and node.kind != "subProcess" and _sla_title_score(title, node.name or "") >= 2:
            return True
        if target is not None and _owner_holds(diagram, target.id, link.target_id):
            return True
        if target is not None and node.kind in GATEWAY_KINDS:
            return any(item.source_id == node.id and item.target_id == target.id for item in diagram.links)
        return False

    pending: List[Dict[str, Any]] = []
    for spec in specs:
        target_num = spec.get("target_num")
        target = bound.get(int(target_num)) if target_num is not None else None
        title = str(spec.get("target_title") or "")
        hit = next((link for link in back if link.id not in claimed and _covers(link, target, title)), None)
        if hit is not None:
            claimed.add(hit.id)
            continue
        pending.append(spec)

    def _cp_hours() -> float:
        root = [pool.process_id for pool in diagram.pools.values()][:1] or [diagram.process_id]
        start_id = diagram.root_start_id if diagram.root_start_id in diagram.nodes else None
        hours, _ = diagram._critical_path(root[0], start_id, {})
        return float(hours)

    def _keep_link(link_id: Optional[str], cp_before: float) -> bool:
        if not link_id or not any(link.id == link_id for link in _refresh()):
            return False
        return abs(_cp_hours() - cp_before) <= 1e-6

    for spec in pending:
        if len(_refresh()) >= len(specs):
            break
        target_num = spec.get("target_num")
        target = bound.get(int(target_num)) if target_num is not None else None
        if target is None:
            continue
        source = bound.get(int(spec["source_num"])) if spec.get("source_num") is not None else None
        reachable = _reachable_from(diagram, target.id)
        picked = _pick_return_source(diagram, target, source, reachable)
        if picked is None:
            continue
        cp_before = _cp_hours()
        linked = diagram.add_link(picked.id, target.id, str(spec.get("label") or "Возврат"))
        if _keep_link(linked, cp_before):
            claimed.add(linked)
            continue
        if linked:
            diagram.links = [link for link in diagram.links if link.id != linked]

    def _forward_ok(cp_before: float) -> bool:
        if abs(_cp_hours() - cp_before) > 1e-6:
            return False
        start_id = diagram.root_start_id if diagram.root_start_id in diagram.nodes else None
        end_id = diagram.root_end_id if diagram.root_end_id in diagram.nodes else None
        if start_id and end_id and end_id not in _reachable_from(diagram, start_id):
            return False
        return True

    if len(_refresh()) > len(specs):
        cp_before = _cp_hours()
        for link in list(_refresh()):
            if len(_refresh()) <= len(specs):
                break
            if link.id in claimed:
                continue
            saved = list(diagram.links)
            diagram.links = [item for item in diagram.links if item.id != link.id]
            if not _forward_ok(cp_before):
                diagram.links = saved


def _limit_escalation_actions(
    actions: List[Dict[str, str]],
    asis_audit: Optional[dict],
    tobe_audit: Optional[dict],
) -> List[Dict[str, str]]:
    """«Цикл заменён эскалацией» только для ребра, которое было в As-Is и снято в To-Be."""
    before = len((asis_audit or {}).get("rework_loops") or [])
    after = len((tobe_audit or {}).get("rework_loops") or [])
    removed = max(0, before - after)
    kept: List[Dict[str, str]] = []
    used = 0
    for act in actions or []:
        detail = str(act.get("detail") or "")
        if re.search(r"цикл заменён эскалацией|петля возврата снята", detail, re.I):
            if used >= removed:
                continue
            used += 1
        kept.append(act)
    return kept


def _resolve_implicit_returns(steps: List[Step]) -> None:
    """Возврат без «п.N» → последний предыдущий шаг указанной (или текущей) роли."""
    for i, step in enumerate(steps):
        d = step.decision
        if not d or d.no_ref is not None or not d.no_back:
            continue
        clause = d.back_clause or d.no_label or ""
        hinted, _, _ = _find_role(clause)
        if hinted is None and _INITIATOR_RE.search(clause):
            hinted = steps[0].role if steps else None
        if hinted is None:
            hinted = step.role
        found: Optional[int] = None
        for prev in reversed(steps[:i]):
            if prev.role == hinted:
                found = prev.num
                break
        if found is None and _INITIATOR_RE.search(clause) and steps:
            found = steps[0].num
        if found is None and i:
            found = steps[i - 1].num
        if found is not None:
            d.no_ref = found


def parse_regulation(text: str) -> ParsedRegulation:
    text = _promote_lettered_parallels(text or "")
    text = _segment_unnumbered_prose(text)
    title: Optional[str] = None
    sla: Optional[float] = None
    raw_steps: List[Tuple[Optional[int], str]] = []
    current_stage: Optional[str] = None
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        m = re.match(r"^(?:регламент|название|процесс)\s*[:—–-]\s*(.+)$", s, re.I)
        if m and title is None:
            title = m.group(1).strip()
            continue
        m = re.match(r"^(?:целев\w+\s+(?:срок|sla)[^:—–]*|sla)\s*[:—–-]\s*(.+)$", s, re.I)
        if m:
            d = _DUR_RE.search(m.group(1))
            if d:
                sla = _to_hours(d.group(1), d.group(2))
            continue
        rm = _ROMAN_LINE_RE.match(s)
        if rm and rm.group("rom"):
            current_stage = rm.group("title").strip()
            rest = current_stage
            if _find_role(rest)[0] or _match_inverted_role(rest) or _split_joint_role(rest):
                body = rest if re.search(r"этап", rest, re.I) else f"{rest} (этап «{current_stage}»)"
                raw_steps.append((None, body))
            continue
        lm = _LETTER_SUB_RE.match(s)
        if lm and lm.group(1).lower() in _SUB_LETTERS:
            body = lm.group(2).strip()
            if current_stage and "этап" not in body.lower():
                body = f"{body} (этап «{current_stage}»)"
            raw_steps.append((None, body))
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            body = m.group(2)
            if current_stage and "этап" not in body.lower():
                body = f"{body} (этап «{current_stage}»)"
            raw_steps.append((int(m.group(1)), body))
            continue
        m = re.match(r"^[-•*]\s+(.*)$", s)
        if m:
            raw_steps.append((None, m.group(1)))
            continue
        if raw_steps:
            num, prev = raw_steps[-1]
            raw_steps[-1] = (num, prev + " " + s)
        elif title is None:
            title = s
        else:
            raw_steps.append((None, s))

    expanded: List[Tuple[Optional[int], str]] = []
    for num, body in raw_steps:
        bits = _atomize_step_body(body)
        if not bits:
            expanded.append((num, body))
            continue
        expanded.append((num, bits[0]))
        for extra in bits[1:]:
            expanded.append((None, extra))
    raw_steps = expanded

    if len(raw_steps) < 2:  # нет нумерации — режем на предложения
        sentences = _split_prose_sentences(" ".join(t for _, t in raw_steps) or text)
        raw_steps = [(None, s.strip()) for s in sentences if len(s.strip()) > 8]

    parsed = ParsedRegulation(title=title or "Бизнес-процесс по регламенту", sla_hours=sla)
    last_role = ""
    letter_shift = 0
    last_assigned = 0
    sequenced: List[Tuple[int, str]] = []
    for num, body in raw_steps:
        if num is None:
            last_assigned += 1
            letter_shift += 1
            sequenced.append((last_assigned, body))
        else:
            last_assigned = num + letter_shift
            sequenced.append((last_assigned, body))
    raw_steps = sequenced
    for idx, (num, body) in enumerate(raw_steps):
        step = Step(idx=idx, num=num if num is not None else idx + 1)
        t = body.strip()
        pm = re.match(r"^(параллельно|одновременно)\s*(\*)?\s*[:,—–-]?\s*", t, re.I)
        if pm:
            step.parallel = True
            step.fork_parallel = bool(pm.group(2))
            t = t[pm.end():]
        sm = re.search(r"\(\s*этап\s*«([^»]+)»\s*\)", t, re.I)
        if sm:
            step.stage = sm.group(1).strip()
            t = (t[: sm.start()] + t[sm.end():]).strip()
        dm = _DUR_RE.search(t)
        if dm:
            step.hours = _to_hours(dm.group(1), dm.group(2))
            t = (t[: dm.start()] + " " + t[dm.end():]).strip()
        t = _INTRO_CLAUSE_RE.sub("", t).strip(" .;,:—–-")
        t = _CONNECTORS_RE.sub("", t).strip(" .;,:—–-")
        bare_ref = _bare_return_ref(t)
        if bare_ref is not None:
            # Номер пункта — ребро, не часть названия и не имя этапа.
            t = _REF_RE.sub(" ", t)
            t = re.sub(r"\bна\s+(?=\s|$)", " ", t)
            t = re.sub(r"\s+", " ", t).strip(" .;,:—–-")
        action, decision = _parse_decision(t)
        step.decision = _ensure_return_decision(t, decision, step.num)
        if step.decision is None:
            step.back_ref = bare_ref
        step.action = bool(action)
        source = action if action else t
        joint = _split_joint_role(source)
        inverted = None if joint else _match_inverted_role(source)
        if inverted is None and decision is not None and joint is None:
            inverted = _match_inverted_role(t)
        if joint:
            step.role = joint[0] or last_role or "Исполнитель"
            last_role = step.role
            step.title = _task_title(joint[1]) if step.action else ""
        elif inverted:
            role, action_phrase, verb = inverted
            step.role = role or last_role or "Исполнитель"
            last_role = step.role
            inf = _infinitive(verb)
            rest = action_phrase[:1].lower() + action_phrase[1:] if action_phrase else ""
            step.title = _task_title(f"{inf} {rest}".strip()) if step.action else ""
        else:
            role, rs, re_ = _find_role(source)
            if role is None and decision is not None:
                role, rs, re_ = _find_role(t)
            if decision is not None and not step.action:
                role = last_role or role
            step.role = role or last_role or "Исполнитель"
            last_role = step.role
            if role and rs == 0 and step.action:
                source = source[re_:]
            step.title = _task_title(source) if step.action else ""
        step.system = bool(_SYSTEM_RE.search(source))
        step.artifacts = extract_artifacts(body)
        step.systems = extract_it_systems(body)
        parsed.steps.append(step)
        if step.role not in parsed.roles:
            parsed.roles.append(step.role)
    _resolve_implicit_returns(parsed.steps)
    parsed.artifacts = aggregate_landscape(parsed.steps, "artifacts")
    parsed.it_systems = aggregate_landscape(parsed.steps, "systems")
    return parsed


def _build_blocks(parsed: ParsedRegulation) -> Tuple[List[Block], Dict[int, int]]:
    steps = parsed.steps
    targets = {s.decision.yes_ref for s in steps if s.decision and s.decision.yes_ref}
    targets |= {s.decision.no_ref for s in steps if s.decision and s.decision.no_ref}
    targets |= {s.back_ref for s in steps if s.back_ref}

    blocks: List[Block] = []
    for step in steps:
        if step.decision is not None:
            blocks.append(Block("decision", [step]))
        elif step.parallel:
            if step.fork_parallel or not blocks or blocks[-1].kind not in ("step", "parallel"):
                blocks.append(Block("parallel", [step]))
            elif blocks[-1].kind == "parallel":
                blocks[-1].steps.append(step)
            else:
                prev = blocks[-1]
                blocks[-1] = Block("parallel", prev.steps + [step])
        else:
            blocks.append(Block("step", [step]))

    for block in blocks:
        if block.kind == "parallel" and len(block.steps) < 2:
            block.kind = "step"

    # Явная граница «(этап «…»)» на соседних шагах — тот же подпроцесс, даже если шагов меньше пяти.
    staged: List[Block] = []
    i = 0
    while i < len(blocks):
        stage = blocks[i].steps[0].stage if blocks[i].kind == "step" and blocks[i].steps else None
        if not stage:
            staged.append(blocks[i])
            i += 1
            continue
        j = i + 1
        while (
            j < len(blocks)
            and blocks[j].kind == "step"
            and blocks[j].steps
            and blocks[j].steps[0].stage == stage
            and blocks[j].role == blocks[i].role
        ):
            j += 1
        run = blocks[i:j]
        while run and run[0].steps[0].num in targets:
            staged.append(run.pop(0))
        if len(run) >= 2:
            staged.append(Block("subprocess", [b.steps[0] for b in run], name=stage))
        else:
            staged.extend(run)
        i = j
    blocks = staged

    merged: List[Block] = []
    i = 0
    while i < len(blocks):
        if blocks[i].kind != "step":
            merged.append(blocks[i])
            i += 1
            continue
        j = i + 1
        while (
            j < len(blocks)
            and blocks[j].kind == "step"
            and blocks[j].role == blocks[i].role
            and blocks[j].steps[0].num not in targets
        ):
            j += 1
        run = blocks[i:j]
        # Цель «вернуть на п.N» не прячем внутрь подпроцесса: иначе ребро садится на имя этапа.
        # Шлюз серию не режет и кусок из-за него в подпроцесс не кладём: шаг остаётся отдельным.
        while run and run[0].steps[0].num in targets:
            merged.append(run.pop(0))
        merged.extend(run)
        i = j

    number_to_block: Dict[int, int] = {}
    for b_idx, block in enumerate(merged):
        for s in block.steps:
            number_to_block[s.num] = b_idx
    return merged, number_to_block


def _q(text: str) -> str:
    return repr(text)


def _decision_branch_keys(b_idx: int, block: Block, n_blocks: int, num_to_block: Dict[int, int]) -> Tuple[Any, Any]:
    d = block.steps[0].decision
    nxt: Any = b_idx + 1 if b_idx + 1 < n_blocks else "END"
    if d is None:
        return nxt, nxt

    def resolve(end: bool, ref: Optional[int], back: bool, has_else: bool) -> Any:
        if end:
            return "END"
        if ref is not None:
            target = num_to_block.get(ref)
            if target is None or target == b_idx:
                return nxt
            return target
        if back:
            return b_idx - 1 if b_idx > 0 else "END"
        if has_else:
            return "END"
        return nxt

    yes = resolve(d.yes_end, d.yes_ref, False, False)
    no = resolve(d.no_end, d.no_ref, d.no_back, d.has_else)
    return yes, no


def _flatten_unary_blocks(blocks: List[Block], num_to_block: Dict[int, int]) -> Tuple[List[Block], Dict[int, int]]:
    """Убирает шлюзы без реального ветвления (1 исходящая ветка) и параллельные группы из 1 шага."""
    n_blocks = len(blocks)
    out: List[Block] = []
    for b_idx, block in enumerate(blocks):
        if block.kind == "parallel" and len(block.steps) < 2:
            out.append(Block("step", block.steps, name=block.name))
            continue
        if block.kind == "decision":
            yes, no = _decision_branch_keys(b_idx, block, n_blocks, num_to_block)
            if yes == no:
                if block.steps[0].action:
                    out.append(Block("step", block.steps, name=block.name))
                continue
        out.append(block)
    mapping: Dict[int, int] = {}
    for i, block in enumerate(out):
        for s in block.steps:
            mapping[s.num] = i
    return out, mapping


def emulate_generation(regulation_text: str) -> Tuple[str, Dict[str, Any]]:
    """Строит код для DIAGRAM по тексту регламента без внешних моделей."""
    parsed = parse_regulation(regulation_text)
    blocks, num_to_block = _build_blocks(parsed)
    blocks, num_to_block = _flatten_unary_blocks(blocks, num_to_block)
    lane_var = {role: f"lane_{i}" for i, role in enumerate(parsed.roles)}

    code: List[str] = [
        f"# Сгенерировано семантическим эмулятором: {parsed.title}",
        f"pool_id, lanes = DIAGRAM.add_pool(ROOT_PROCESS_ID, {json.dumps(parsed.roles, ensure_ascii=False)})",
    ]
    code += [f"lane_{i} = lanes[{i}]" for i in range(len(parsed.roles))]

    def emit_task(var: str, step: Step, parent: str) -> None:
        func = "add_script_task" if step.system else "add_user_task"
        code.append(f"{var} = DIAGRAM.{func}({_q(step.title)}, {parent})")
        if step.hours is not None:
            code.append(f"DIAGRAM.set_sla({var}, {round(step.hours, 3)})")

    entry: List[str] = []
    exits: List[str] = []
    gateways: Dict[int, str] = {}
    for b_idx, block in enumerate(blocks):
        lane = lane_var[block.role]
        if block.kind == "step":
            var = f"n{b_idx}"
            emit_task(var, block.steps[0], lane)
            entry.append(var)
            exits.append(var)
        elif block.kind == "decision":
            step = block.steps[0]
            gw = f"g{b_idx}"
            question = step.decision.yes_label.rstrip("?") + "?"  # type: ignore[union-attr]
            if step.action:
                var = f"n{b_idx}"
                emit_task(var, step, lane)
                code.append(f"{gw} = DIAGRAM.add_exclusive_gateway({_q(question)}, {lane})")
                code.append(f"DIAGRAM.add_link({var}, {gw})")
                entry.append(var)
            else:
                code.append(f"{gw} = DIAGRAM.add_exclusive_gateway({_q(question)}, {lane})")
                entry.append(gw)
            gateways[b_idx] = gw
            exits.append(gw)
        elif block.kind == "parallel":
            if len(block.steps) < 2:
                var = f"n{b_idx}"
                emit_task(var, block.steps[0], lane)
                entry.append(var)
                exits.append(var)
                continue
            split, join = f"ps{b_idx}", f"pj{b_idx}"
            code.append(f"{split} = DIAGRAM.add_parallel_gateway('Параллельные работы', {lane})")
            code.append(f"{join} = DIAGRAM.add_parallel_gateway('Работы завершены', {lane})")
            for k, step in enumerate(block.steps):
                var = f"n{b_idx}_{k}"
                emit_task(var, step, lane_var[step.role])
                code.append(f"DIAGRAM.add_link({split}, {var})")
                code.append(f"DIAGRAM.add_link({var}, {join})")
            entry.append(split)
            exits.append(join)
        else:  # subprocess
            sub = f"sp{b_idx}"
            code.append(f"{sub} = DIAGRAM.create_subprocess({_q(block.name)}, {lane})")
            prev = None
            for k, step in enumerate(block.steps):
                var = f"n{b_idx}_{k}"
                emit_task(var, step, sub)
                if prev:
                    code.append(f"DIAGRAM.add_link({prev}, {var})")
                prev = var
            entry.append(sub)
            exits.append(sub)

    end_expr = "ROOT_END_TASK_ID"
    if not entry:
        code.append("DIAGRAM.add_link(ROOT_START_TASK_ID, ROOT_END_TASK_ID)")
    else:
        code.append(f"DIAGRAM.add_link(ROOT_START_TASK_ID, {entry[0]})")

    def resolve(ref: Optional[int], own: int) -> Optional[str]:
        if ref is None or ref not in num_to_block:
            return None
        target = num_to_block[ref]
        return None if target == own else entry[target]

    for b_idx, block in enumerate(blocks):
        nxt = entry[b_idx + 1] if b_idx + 1 < len(blocks) else end_expr
        if block.kind != "decision":
            code.append(f"DIAGRAM.add_link({exits[b_idx]}, {nxt})")
            continue
        d = block.steps[0].decision
        assert d is not None
        gw = gateways[b_idx]
        yes_t = "ROOT_END_TASK_ID" if d.yes_end else (resolve(d.yes_ref, b_idx) or nxt)
        if d.no_end:
            no_t = end_expr
        elif d.no_ref is not None:
            no_t = resolve(d.no_ref, b_idx) or nxt
        elif d.no_back:
            no_t = entry[b_idx - 1] if b_idx > 0 else end_expr
        elif d.has_else:
            no_t = end_expr
        else:
            no_t = nxt
        code.append(f"DIAGRAM.add_link({gw}, {yes_t}, {_q(d.yes_label)})")
        if no_t != yes_t:
            code.append(f"DIAGRAM.add_link({gw}, {no_t}, {_q(d.no_label)})")

    info = {
        "title": parsed.title,
        "steps": len(parsed.steps),
        "roles": parsed.roles,
        "blocks": len(blocks),
        "subprocesses": sum(1 for b in blocks if b.kind == "subprocess"),
        "decisions": sum(1 for b in blocks if b.kind == "decision"),
        "parallel_groups": sum(1 for b in blocks if b.kind == "parallel"),
        "sla_hours": parsed.sla_hours,
    }
    return "\n".join(code) + "\n", info


# --------------------------------------------------------------------------- #
# Главная точка входа
# --------------------------------------------------------------------------- #
def _rate_limit_wait(exc: Exception) -> Optional[float]:
    """HTTP 429 → сколько секунд подождать (Retry-After или «try again in 7.5s»), не больше 20; иначе None."""
    response = getattr(exc, "response", None)
    if response is None or getattr(response, "status_code", None) != 429:
        return None
    header = (getattr(response, "headers", None) or {}).get("retry-after")
    try:
        seconds = float(header) if header else None
    except ValueError:
        seconds = None
    if seconds is None:
        m = re.search(r"try again in ([\d.]+)s", getattr(response, "text", "") or "")
        seconds = float(m.group(1)) if m else 6.0
    return min(20.0, seconds + 0.5)


def _error_reason(exc: Exception) -> str:
    """Короткая причина сбоя LLM: для HTTP-ошибок — код и сообщение провайдера («401: Invalid API Key»)."""
    if "Timeout" in type(exc).__name__:
        return "таймаут"
    response = getattr(exc, "response", None)  # requests.Response ложен при 4xx — сравниваем с None
    if response is None:
        response = getattr(exc, "fp", None)
    status = getattr(response, "status_code", None) or getattr(exc, "code", None)
    if status is None:
        return type(exc).__name__
    message = ""
    try:
        body = response.json() if hasattr(response, "json") else json.loads(response.read().decode("utf-8"))
        error = body.get("error", body) if isinstance(body, dict) else body
        message = error.get("message", "") if isinstance(error, dict) else str(error)
    except Exception:  # noqa: BLE001
        message = str(getattr(response, "text", ""))
    message = re.sub(r"\s+", " ", message).strip()[:160]
    return f"HTTP {status}: {message}" if message else f"HTTP {status}"


# Облачные провайдеры по порядку: OPENAI_* (основной) → FALLBACK_* → FALLBACK2_* …
# Бесплатные тарифы по отдельности ненадёжны: Groq отвечает 403 с части IP Streamlit Cloud,
# бесплатный канал OpenRouter бывает перегружен (429) — поэтому цепочка из нескольких.
CLOUD_PREFIXES = ("OPENAI", "FALLBACK", "FALLBACK2", "FALLBACK3")


def _cloud_prefixes() -> List[str]:
    return [p for p in CLOUD_PREFIXES if os.getenv(f"{p}_API_KEY")]


_SET_SLA_RE = re.compile(r"(DIAGRAM\.set_sla\(\s*[^,]+,\s*)([\d.]+)(\s*\))")


def _fix_sla_units(code: str, text_hours: List[float]) -> Tuple[str, str]:
    """Модель иногда пишет сроки в днях («5 рабочих дней» → 5), а set_sla ждёт часы (40).

    Сверяем сумму её сроков с суммой сроков, разобранных из текста: расхождение ровно ×8
    (рабочие дни) или ×24 (сутки) — это перепутанные единицы, пересчитываем. Иначе код не трогаем.
    """
    values = [float(m.group(2)) for m in _SET_SLA_RE.finditer(code or "")]
    if len(text_hours) < 2 or len(values) < 2 or not sum(values):
        return code, ""
    ratio = sum(text_hours) / sum(values)
    factor = next((f for f in (8.0, 24.0) if abs(ratio - f) / f < 0.2), None)
    if factor is None:
        return code, ""
    fixed = _SET_SLA_RE.sub(lambda m: f"{m.group(1)}{round(float(m.group(2)) * factor, 3)}{m.group(3)}", code)
    return fixed, f"сроки шагов переведены из дней в часы (×{factor:.0f}): модель указала дни вместо часов"


_UNRECOGNIZED_PROCESS = (
    "Не удалось распознать шаги процесса. Пожалуйста, опишите регламент по пунктам "
    "(например: 1. Диспетчер принимает заявку...)."
)


def cloud_engine_status() -> str:
    """Строка для интерфейса: подключённые облачные модели по порядку (без раскрытия ключей)."""
    parts = []
    for prefix in _cloud_prefixes():
        host = re.sub(r"^https?://([^/]+).*$", r"\1", os.getenv(f"{prefix}_BASE_URL", OPENAI_DEFAULT_BASE))
        parts.append(f"{os.getenv(f'{prefix}_MODEL', 'gpt-4o-mini')} · {host}")
    if not parts:
        return "не подключена (OPENAI_API_KEY не задан)"
    return parts[0] + "".join(f" · запасная: {p}" for p in parts[1:])


def _engines() -> List[Tuple[str, Callable[[str], Tuple[str, str]]]]:
    mode = os.getenv("BPMN_AI_MODE", "auto").lower()
    if mode == "emulator":
        return []
    # Облачные модели отвечают за секунды — они первые; Ollama — офлайн-резерв.
    engines: List[Tuple[str, Callable[[str], Tuple[str, str]]]] = [
        (prefix.lower(), (lambda prompt, _p=prefix: _call_openai(prompt, _p))) for prefix in _cloud_prefixes()
    ]
    engines.append(("ollama", _call_ollama))
    return engines


def _reject_cloud_diagram(err: str, quality: Optional[Dict[str, Any]]) -> bool:
    """Ответ облака не берём, если код не собрался, граф битый или возвратов меньше, чем фраз.

    Нехватку рёбер чинит дорисовка в execute_generated_code. Если после неё фраз
    всё ещё больше, схема не готова: целиком подменять её эмулятором нельзя.
    """
    if err:
        return True
    critical = [str(item) for item in ((quality or {}).get("critical") or [])]
    return bool(critical)


def generate_bpmn_from_text(regulation_text: str, use_llm: bool = True) -> Tuple[str, Dict[str, Any], str]:
    """Регламент (RU) → (bpmn_xml, audit_data, error).

    Порядок: облачный API (если задан OPENAI_API_KEY) → локальная Ollama (qwen2.5-coder / llama3) →
    встроенный семантический эмулятор. Исключения сети наружу не выходят.
    Ответ модели отклоняется по графу после heal (нет старта/конца, тупик или дыра) или по XSD,
    не по тексту журнала лечения. Нет ключа, ошибка сети или пустой ответ — эмулятор.
    модель получает список ошибок и одну попытку исправиться (LLM_MAX_ATTEMPTS).
    """
    started = time.time()
    text = normalize_regulation(regulation_text or "").strip()
    if len(text) < 20:
        return "", {}, _UNRECOGNIZED_PROCESS

    artifacts: List[Dict[str, Any]] = []
    it_systems: List[Dict[str, Any]] = []
    text_hours: List[float] = []
    try:
        header = parse_regulation(text)
        process_name, sla = header.title, header.sla_hours
        artifacts, it_systems = header.artifacts, header.it_systems
        text_hours = [s.hours for s in header.steps if s.hours]
    except Exception:  # noqa: BLE001
        process_name, sla = "Бизнес-процесс по регламенту", None

    trace: List[str] = []
    rejected: List[Dict[str, Any]] = []  # отклонённые ответы LLM — для разбора в «Технических деталях»
    incomplete_model = ""
    if use_llm:
        max_attempts = max(1, int(os.getenv("LLM_MAX_ATTEMPTS", "2")))
        retry_limit_s = float(os.getenv("LLM_RETRY_MAX_CALL_S", "90"))
        if not os.getenv("OPENAI_API_KEY"):
            trace.append("облачный API: OPENAI_API_KEY не задан")
        for name, call in _engines():
            prompt = build_prompt(text)
            for attempt in range(1, max_attempts + 1):
                call_started = time.time()
                try:
                    engine_label, raw = call(prompt)
                except Exception as exc:  # noqa: BLE001 — недоступность модели не должна ронять приложение
                    trace.append(f"{name}: недоступен ({_error_reason(exc)})")
                    break
                call_s = time.time() - call_started
                raw, unit_note = _fix_sla_units(raw, text_hours)
                if unit_note:
                    trace.append(f"{engine_label}: {unit_note}")
                xml, audit, err = execute_generated_code(raw, process_name, sla, regulation_text=text)
                quality = audit.get("quality", {}) if not err else {}
                if not _reject_cloud_diagram(err, quality):
                    audit["artifacts"], audit["it_systems"] = artifacts, it_systems
                    audit["generation"] = {
                        "engine": engine_label,
                        "fallback": False,
                        "attempts": attempt,
                        "trace": trace,
                        "rejected": rejected,
                        "code": _strip_markdown(raw),
                        "elapsed_s": round(time.time() - started, 2),
                    }
                    return xml, audit, ""
                if _model_return_incomplete(err):
                    incomplete_model = err
                problems = [err] if err else quality["critical"]
                if err and (
                    err == _UNRECOGNIZED_PROCESS
                    or "не удалось распознать" in err.lower()
                    or "не распознан" in err.lower()
                ):
                    # Модель сама не нашла процесс в тексте: «исправлять» — значит заставить её выдумывать.
                    trace.append(f"{engine_label}: регламент не распознан — повтор не делаем")
                    return "", {}, _UNRECOGNIZED_PROCESS
                rejected.append({"engine": engine_label, "attempt": attempt, "problems": problems, "code": _strip_markdown(raw)})
                trace.append(
                    f"{engine_label}, попытка {attempt} ({call_s:.0f} с): результат отклонён — "
                    + "; ".join(problems[:5])
                )
                if attempt < max_attempts and call_s > retry_limit_s:
                    trace.append(f"{engine_label}: повтор пропущен — модель отвечала дольше {retry_limit_s:.0f} с")
                    break
                prompt = build_repair_prompt(text, _strip_markdown(raw), problems)

    if incomplete_model:
        return "", {}, incomplete_model

    try:
        code, info = emulate_generation(text)
    except Exception as exc:  # noqa: BLE001
        return "", {}, _UNRECOGNIZED_PROCESS
    xml, audit, err = execute_generated_code(code, process_name, sla, regulation_text=text)
    if err:
        if (
            err == _UNRECOGNIZED_PROCESS
            or "не удалось распознать" in err.lower()
            or "не распознан" in err.lower()
            or "меньше двух" in err
        ):
            return "", {}, _UNRECOGNIZED_PROCESS
        return "", {}, f"Не удалось построить диаграмму: {err}"
    audit["artifacts"], audit["it_systems"] = artifacts, it_systems
    audit["generation"] = {
        "engine": "semantic-emulator",
        "fallback": use_llm,
        "trace": trace or (["LLM отключён"] if not use_llm else ["LLM недоступна — включён встроенный эмулятор"]),
        "rejected": rejected,
        "code": code,
        "elapsed_s": round(time.time() - started, 2),
        "parsed": info,
    }
    return xml, audit, ""


# =========================================================================== #
# AI-АССИСТЕНТ БИЗНЕС-АРХИТЕКТОРА: диалог над активным процессом
# =========================================================================== #
#   1) аналитика и объяснение      — «В чём причина срыва SLA?», «Как разгрузить диспетчера?»
#   2) правка процесса на лету     — «Добавь согласование с экологами после шага 3»
#                                     → правка текста регламента → перестроение через execute_generated_code
#   3) реверс-генерация            — «Составь инструкцию для роли Диспетчер» (по схеме BPMN)
# Цепочка: облачная модель (OPENAI_* / FALLBACK_*) → локальная Ollama → локальные эвристики (без сети).
# --------------------------------------------------------------------------- #
_NS_BPMN = "http://www.omg.org/spec/BPMN/20100524/MODEL"
_NS_DI = "http://www.omg.org/spec/BPMN/20100524/DI"
_NS_DC = "http://www.omg.org/spec/DD/20100524/DC"
_TASK_TAGS = {"task", "userTask", "scriptTask", "serviceTask", "manualTask", "sendTask", "receiveTask", "businessRuleTask"}
_GATE_TAGS = {"exclusiveGateway", "parallelGateway", "inclusiveGateway", "eventBasedGateway", "complexGateway"}

CHAT_SYSTEM = """Ты — «AI-Ассистент Бизнес-Архитектора» ПАО «Интер РАО». Собеседник — эксперт Дирекции бизнес-архитектуры.
Правила:
- Отвечай по-русски, не более 10 строк. Сначала вывод с цифрой из блока ФАКТЫ, затем шаги и роли, в конце одна команда в «ёлочках», если уместна правка.
- Часы, проценты и число циклов бери ТОЛЬКО из ФАКТОВ. Не выдумывай числа.
- На SLA / As-Is/To-Be / циклы / bus-factor: цифры только из ФАКТОВ. «Без ускорения» — только если не укоротились ни голый путь, ни путь с возвратами.
- На «сравни as-is / to-be»: путь с возвратами, циклы N → M, голый КП.
- Предлог «со» перед творительным на с, з, ж, ш, щ («со службой экологии», не «с службой»).
- Не начинай с заглушки «N шагов / M узлов». Если данных нет — так и скажи.

ДАННЫЕ АКТИВНОГО ПРОЦЕССА:
"""

EDIT_SYSTEM = """Ты — редактор регламентов ПАО «Интер РАО». Тебе дан текущий регламент и команда пользователя.
Формат регламента: заголовок «Регламент: …», строка «Целевой срок …: …», затем нумерованные шаги:
«N. Роль глагол … (срок). Если условие — перейти к п.K, иначе «метка» — вернуть на п.M.»; «Параллельно: …» — шаг,
выполняемый одновременно с предыдущим.
Верни ТОЛЬКО полный обновлённый регламент в том же формате: перенумеруй шаги подряд и поправь ссылки «п.N»;
не меняй то, что пользователь не просил; каждый шаг начинается с названия роли (подразделения); без пояснений и markdown."""

POLISH_SYSTEM = """Оформи приведённые факты как официальную должностную инструкцию на русском языке.
Используй ТОЛЬКО факты из входных данных: ничего не добавляй и не выдумывай (сроки, системы, роли, документы).
Структура: 1. Общие положения, 2. Входные данные и триггеры, 3. Порядок действий (нумерованный список, сохрани все шаги
и условия ветвления), 4. Передача результата, 5. Используемые системы и документы, 6. Контроль сроков и риски. Markdown."""


# --------------------------- вызовы LLM в диалоговом режиме --------------------------- #
def _chat_openai(messages: List[Dict[str, str]], prefix: str) -> Tuple[str, str]:
    key = os.getenv(f"{prefix}_API_KEY")
    if not key:
        raise RuntimeError(f"{prefix}_API_KEY не задан")
    base = os.getenv(f"{prefix}_BASE_URL", OPENAI_DEFAULT_BASE).rstrip("/")
    model = os.getenv(f"{prefix}_MODEL", "gpt-4o-mini")
    payload: Dict[str, Any] = {"model": model, "temperature": 0.2, "messages": messages}
    effort = os.getenv(f"{prefix}_REASONING_EFFORT", "")
    if effort == "none" and "openrouter.ai" in base:
        payload["reasoning"] = {"enabled": False}
    elif effort:
        payload["reasoning_effort"] = effort
    data = _http_json(
        f"{base}/chat/completions",
        payload,
        headers={"Authorization": f"Bearer {key}"},
        timeout=float(os.getenv("CHAT_TIMEOUT", "60")),
    )
    content = str(data["choices"][0]["message"].get("content") or "").strip()
    if not content:
        raise RuntimeError("пустой ответ модели")
    return ("openai" if prefix == "OPENAI" else prefix.lower()) + f":{model}", content


def _chat_ollama(messages: List[Dict[str, str]]) -> Tuple[str, str]:
    model = _pick_ollama_model(_ollama_models())
    if not model:
        raise RuntimeError("в Ollama нет моделей qwen2.5-coder / llama3")
    data = _http_json(
        f"{_ollama_base()}/api/chat",
        {
            "model": model,
            "messages": messages,
            "stream": False,
            "keep_alive": "30m",
            "options": {"temperature": 0.2, "num_ctx": int(os.getenv("OLLAMA_NUM_CTX", "8192"))},
        },
        timeout=float(os.getenv("OLLAMA_CHAT_TIMEOUT", "120")),
    )
    content = str((data.get("message") or {}).get("content") or "").strip()
    if not content:
        raise RuntimeError("пустой ответ модели")
    return f"ollama:{model}", content


def _chat_llm(messages: List[Dict[str, str]], trace: List[str]) -> Optional[Tuple[str, str]]:
    """Первая доступная модель отвечает; сбои пишутся в trace, наружу исключения не выходят."""
    if os.getenv("BPMN_AI_MODE", "auto").lower() == "emulator":
        return None
    attempts: List[Tuple[str, Callable[[], Tuple[str, str]]]] = []
    for prefix in _cloud_prefixes():
        attempts.append((prefix.lower(), lambda _p=prefix: _chat_openai(messages, _p)))
    attempts.append(("ollama", lambda: _chat_ollama(messages)))
    for name, call in attempts:
        try:
            return call()
        except Exception as exc:  # noqa: BLE001 — недоступность модели не должна ронять диалог
            trace.append(f"{name}: {_error_reason(exc)}")
    return None


# --------------------------- контекст активного процесса --------------------------- #
def _fh(hours: float) -> str:
    if hours < 1:
        return f"{hours * 60:.0f} мин"
    if hours < 48:
        return f"{hours:.1f} ч".replace(".0 ", " ")
    return f"{hours:.0f} ч (≈ {hours / 8:.0f} раб. дн.)"


def process_facts(audit: Optional[Dict[str, Any]] = None, tobe_delta: Optional[dict] = None) -> Dict[str, Any]:
    """Единый набор цифр для копайлота и сайдбара: КП, срок с возвратами, циклы, bus-factor, To-Be."""
    data = audit or {}
    sla = data.get("sla") or {}
    bus = data.get("bus_factor") or {}
    loops = list(data.get("rework_loops") or [])
    delta = tobe_delta if isinstance(tobe_delta, dict) else {}
    if not delta:
        cmp = data.get("tobe_compare")
        delta = cmp if isinstance(cmp, dict) else {}

    cp = float(sla.get("critical_path_hours") or 0)
    rw = float(sla.get("with_rework_hours") or cp)
    cp_b = delta.get("sla_before_hours")
    cp_a = delta.get("sla_after_hours")
    rw_b = delta.get("with_rework_before")
    rw_a = delta.get("with_rework_after")
    lb = delta.get("rework_before")
    la = delta.get("rework_after")
    tobe_ready = cp_a is not None or rw_a is not None or lb is not None

    cp_before = round(float(cp_b if cp_b is not None else cp), 1)
    rw_before = round(float(rw_b if rw_b is not None else rw), 1)
    loops_before = int(lb if lb is not None else len(loops))
    cp_after = round(float(cp_a), 1) if cp_a is not None else (round(cp_before, 1) if tobe_ready else None)
    rw_after = round(float(rw_a), 1) if rw_a is not None else (round(rw_before, 1) if tobe_ready else None)
    loops_after = int(la) if la is not None else (0 if tobe_ready else None)

    rw_saved = round(rw_before - float(rw_after or 0), 1) if tobe_ready else 0.0
    cp_ok = cp_after is None or float(cp_after) <= float(cp_before) + 0.05
    speedup = bool(tobe_ready and rw_saved > 0 and cp_ok)
    return {
        "critical_path_hours": round(cp, 1),
        "with_rework_hours": round(rw, 1),
        "rework_loops_n": len(loops),
        "rework_loops": loops,
        "critical_path": list(data.get("critical_path") or []),
        "lane_load": list(data.get("lane_load") or []),
        "bus_role": str(bus.get("top_role") or ""),
        "bus_share": float(bus.get("max_share") or 0),
        "bus_threshold": float(bus.get("threshold") or 0.45),
        "bus_status": str(bus.get("status") or ""),
        "tobe_ready": bool(tobe_ready),
        "cp_before": cp_before,
        "cp_after": cp_after,
        "rw_before": rw_before,
        "rw_after": rw_after,
        "loops_before": loops_before,
        "loops_after": loops_after,
        "rw_saved_hours": rw_saved,
        "rw_saved_pct": int(round(100.0 * rw_saved / rw_before)) if tobe_ready and rw_before else 0,
        "speedup_via_rework": speedup,
        "cp_grew": bool(tobe_ready and cp_after is not None and cp_after > cp_before),
        "quality_before": int(delta.get("quality_before") or (data.get("methodology") or {}).get("score") or 0),
        "quality_after": int(delta.get("quality_after") or 0),
        "tobe_actions": [a for a in (delta.get("actions") or []) if isinstance(a, dict)],
    }


def _limit_lines(text: str, n: int) -> str:
    lines = [ln.rstrip() for ln in (text or "").splitlines() if ln.strip()]
    return "\n".join(lines[:n])


def build_process_context(
    current_text: str,
    current_xml: str,
    current_audit: Dict[str, Any],
    tobe_delta: Optional[dict] = None,
) -> Dict[str, Any]:
    """Сводка активного процесса: шаги, роли, SLA, критический путь, bus-factor, циклы, ИТ-ландшафт, To-Be."""
    audit = current_audit or {}
    title, sla_target, steps = "Бизнес-процесс", None, []
    try:
        parsed = parse_regulation(normalize_regulation(current_text or ""))
        title, sla_target = parsed.title, parsed.sla_hours
        for st in parsed.steps:
            decision = None
            if st.decision:
                decision = {"yes": st.decision.yes_label, "no": st.decision.no_label, "back": st.decision.no_back}
            steps.append(
                {
                    "num": st.num, "role": st.role, "title": st.title or "(решение)", "hours": st.hours,
                    "parallel": st.parallel, "stage": st.stage, "decision": decision,
                    "systems": st.systems, "artifacts": st.artifacts,
                }
            )
    except Exception:  # noqa: BLE001
        pass
    sla = audit.get("sla") or {}
    if sla_target is None:
        sla_target = sla.get("target_hours")
    ctx: Dict[str, Any] = {
        "title": title,
        "sla_target_hours": sla_target,
        "steps": steps,
        "lane_load": audit.get("lane_load") or [],
        "bus_factor": audit.get("bus_factor") or {},
        "sla": sla,
        "critical_path": audit.get("critical_path") or [],
        "rework_loops": audit.get("rework_loops") or [],
        "sla_risks": audit.get("sla_risks") or [],
        "it_systems": audit.get("it_systems") or [],
        "artifacts": audit.get("artifacts") or [],
        "recommendations": audit.get("recommendations") or [],
        "stats": audit.get("stats") or {},
        "methodology": audit.get("methodology") or {},
        "xml_available": bool(current_xml),
        "tobe_ready": False,
        "sla_hours_as_is": None,
        "sla_hours_to_be": None,
        "delta_sla_percent": None,
        "rework_loops_as_is": None,
        "rework_loops_to_be": None,
        "quality_score_as_is": None,
        "quality_score_to_be": None,
        "tobe_actions": [],
        "facts": {},
        "cp_hours_as_is": None,
        "cp_hours_to_be": None,
        "regulation_text": current_text or "",
    }
    facts = process_facts(audit, tobe_delta)
    ctx["facts"] = facts
    _attach_tobe_compare(ctx, audit, tobe_delta)
    return ctx


def _attach_tobe_compare(ctx: Dict[str, Any], audit: Dict[str, Any], tobe_delta: Optional[dict]) -> None:
    """Пишет в контекст цифры To-Be из process_facts: срок с возвратами отдельно от голого КП."""
    facts = ctx.get("facts") or process_facts(audit, tobe_delta)
    ctx["facts"] = facts
    sla = (audit or {}).get("sla") or {}
    meth = (audit or {}).get("methodology") or {}
    d = tobe_delta if isinstance(tobe_delta, dict) else {}
    if not facts.get("tobe_ready"):
        return
    ctx["tobe_ready"] = True
    ctx["sla_hours_as_is"] = facts["rw_before"]
    ctx["sla_hours_to_be"] = facts["rw_after"]
    ctx["delta_sla_percent"] = facts["rw_saved_pct"]
    ctx["sla_saved_hours"] = facts["rw_saved_hours"]
    ctx["cp_hours_as_is"] = facts["cp_before"]
    ctx["cp_hours_to_be"] = facts["cp_after"]
    ctx["rework_loops_as_is"] = facts["loops_before"]
    ctx["rework_loops_to_be"] = facts["loops_after"]
    q_as = d.get("quality_before")
    if q_as is None:
        q_as = int(meth.get("score") or 0)
    q_to = d.get("quality_after")
    if q_to is None:
        q_to = q_as
    ctx["quality_score_as_is"] = int(q_as or 0)
    ctx["quality_score_to_be"] = int(q_to or 0)
    ctx["tobe_actions"] = facts.get("tobe_actions") or []
    ctx["rework_hours_as_is"] = round(float(d.get("rework_hours_before") or sla.get("rework_hours") or 0), 3)
    ctx["rework_hours_to_be"] = round(float(d.get("rework_hours_after") or 0), 3)


def format_context_for_prompt(ctx: Dict[str, Any], max_steps: int = 60) -> str:
    lines = [f"Процесс: {ctx['title']}"]
    if ctx.get("sla_target_hours"):
        lines.append(f"Целевой срок (SLA): {_fh(float(ctx['sla_target_hours']))}")
    sla = ctx.get("sla") or {}
    if sla:
        lines.append(
            f"Критический путь: {_fh(float(sla.get('critical_path_hours', 0)))}; "
            f"с худшим возвратом: {_fh(float(sla.get('with_rework_hours', 0)))}; "
            f"срыв SLA: {'да' if sla.get('breach') else 'нет'}"
        )
    lines.append("Шаги:")
    for s in ctx["steps"][:max_steps]:
        extra = []
        if s["hours"]:
            extra.append(_fh(float(s["hours"])))
        if s["parallel"]:
            extra.append("параллельно с предыдущим")
        if s["decision"]:
            extra.append(f"ветвление: «{s['decision']['yes']}» / «{s['decision']['no']}»")
        lines.append(f"  {s['num']}. [{s['role']}] {s['title']}" + (f" ({'; '.join(extra)})" if extra else ""))
    if len(ctx["steps"]) > max_steps:
        lines.append(f"  … ещё {len(ctx['steps']) - max_steps} шагов")
    if ctx["lane_load"]:
        lines.append("Нагрузка по ролям: " + "; ".join(
            f"{r['role']} {float(r['share']):.0%} ({int(r['tasks'])} шаг.)" for r in ctx["lane_load"]))
    bus = ctx.get("bus_factor") or {}
    if bus:
        lines.append(f"Bus-factor: макс. доля {float(bus.get('max_share', 0)):.0%} у роли «{bus.get('top_role')}» "
                     f"(порог {float(bus.get('threshold', 0.45)):.0%}, статус {bus.get('status')})")
    if ctx["critical_path"]:
        lines.append("Критический путь: " + " → ".join(
            f"{c['name']} [{c['role']}, {_fh(float(c['hours']))}]" for c in ctx["critical_path"][:12]))
    for loop in ctx["rework_loops"]:
        lines.append(f"Цикл возврата «{loop['label']}»: {loop['from']} → {loop['to']} ({loop['lane']}), "
                     f"стоимость цикла {_fh(float(loop['cycle_hours']))}")
    if ctx["it_systems"]:
        lines.append("ИТ-системы: " + ", ".join(f"{i['name']} (шаги {i['steps']})" for i in ctx["it_systems"]))
    if ctx["artifacts"]:
        lines.append("Документы: " + ", ".join(f"{i['name']} (шаги {i['steps']})" for i in ctx["artifacts"]))
    for risk in ctx["sla_risks"][:6]:
        lines.append(f"Риск [{risk.get('severity')}]: {risk.get('message')}")
    meth = ctx.get("methodology") or {}
    if meth.get("score") is not None:
        lines.append(f"Quality Score (линтер нотации): {int(meth.get('score') or 0)}%")
        for chk in (meth.get("checks") or [])[:4]:
            mark = "ok" if chk.get("passed") else "fail"
            lines.append(f"  [{mark}] {chk.get('title')}: {chk.get('detail')}")
    facts = ctx.get("facts") or {}
    if facts:
        lines.append(
            f"ФАКТЫ: КП {_fh(float(facts.get('critical_path_hours') or 0))}; "
            f"с возвратами {_fh(float(facts.get('with_rework_hours') or 0))}; "
            f"циклов {int(facts.get('rework_loops_n') or 0)}; "
            f"bus-factor «{facts.get('bus_role') or '—'}» {float(facts.get('bus_share') or 0):.0%}"
        )
    if ctx.get("tobe_ready"):
        lines += [
            "СРАВНЕНИЕ AS-IS / TO-BE:",
            f"  путь_с_возвратами: {ctx.get('sla_hours_as_is')} → {ctx.get('sla_hours_to_be')} ч (−{ctx.get('delta_sla_percent')}%)",
            f"  голый_критический_путь: {ctx.get('cp_hours_as_is')} → {ctx.get('cp_hours_to_be')} ч",
            f"  циклы: {ctx.get('rework_loops_as_is')} → {ctx.get('rework_loops_to_be')}",
            f"  quality_score_as_is: {ctx.get('quality_score_as_is')}",
            f"  quality_score_to_be: {ctx.get('quality_score_to_be')}",
        ]
        if facts.get("speedup_via_rework"):
            lines.append("  НЕ говорить «Без ускорения»: путь с возвратами короче, голый КП не длиннее As-Is.")
        elif facts.get("cp_grew"):
            lines.append("  Голый КП вырос — не выдавай это за оптимум To-Be.")
        for act in (ctx.get("tobe_actions") or [])[:6]:
            lines.append(f"  действие To-Be [{act.get('kind')}]: {act.get('detail')}")
        lines.append(f"  rework_hours_as_is: {ctx.get('rework_hours_as_is')}")
        lines.append(f"  rework_hours_to_be: {ctx.get('rework_hours_to_be')}")
        lines.append(f"  sla_saved_hours: {ctx.get('sla_saved_hours')}")
    return "\n".join(lines)


# --------------------------- классификация намерения --------------------------- #
_EDIT_VERB_RE = re.compile(
    r"\b(добав\w+|вставь\w*|вставить|удал\w+|убер\w+|убрать|исключ\w+|"
    r"перепиши\w*|переписать|"
    r"сделай(?:те)?\s+.{0,80}параллел|"
    r"измени(?:ть)?\s+(?:срок|шаг|роль|название|пункт)|"
    r"замени\w*|заменить|перенес\w+|перемест\w+|поменя\w+|"
    r"передай\w+|назначь\w+|установи\w*\s+целев\w+\s+срок)\b",
    re.I,
)
_NO_EDIT_RE = re.compile(
    r"не\s+изменяй|не\s+менять|не\s+правь|не\s+править|не\s+трогай|"
    r"не\s+перестраивай|не\s+редактир|без\s+изменен|не\s+меняй\s+bpmn",
    re.I,
)
_ANALYSIS_HINT_RE = re.compile(
    r"аудит|анализ|аналитическ|сравни|as-is|to-be|tobe|риск|"
    r"что\s+изменил|подтвержд|разбор",
    re.I,
)
_QUESTION_START_RE = re.compile(r"^\s*(?:как|почему|что|зачем|можно ли|стоит ли|какие|какой|какая|сколько|где|когда|кто|в чем|в чём|есть ли)\b", re.I)
_INSTRUCTION_RE = re.compile(
    r"инструкци|памятк|регламент для исполнител|для исполнител|должностн|чем занимается|что делает|опиши работу|обязанност", re.I)
_NEXT_STEP_RE = re.compile(r"следующ\w+ шаг|что дальше|чего не хватает|предложи\w* шаг|что добавить|каких шагов", re.I)
_OPT_AUDIT_Q_RE = re.compile(
    r"аудит|аналитическ|что\s+изменил|подтвержд|"
    r"не\s+изменяй|не\s+правь|не\s+менять|"
    r"разбор.{0,24}оптимиз|оптимизац.{0,40}(аудит|анализ)",
    re.I,
)


def classify_intent(message: str) -> str:
    """'instruction' | 'edit' | 'next_step' | 'analysis'."""
    msg = (message or "").strip()
    if _INSTRUCTION_RE.search(msg) and not re.match(r"^\s*(?:удал|убер)", msg, re.I):
        return "instruction"
    if _NEXT_STEP_RE.search(msg):
        return "next_step"
    if _NO_EDIT_RE.search(msg) or (_ANALYSIS_HINT_RE.search(msg) and not _EDIT_VERB_RE.search(msg)):
        return "analysis"
    if _EDIT_VERB_RE.search(msg) and not msg.endswith("?") and not _QUESTION_START_RE.match(msg):
        return "edit"
    return "analysis"


def _roles_of(ctx: Dict[str, Any]) -> List[str]:
    roles: List[str] = [r["role"] for r in ctx["lane_load"]]
    for s in ctx["steps"]:
        if s["role"] not in roles:
            roles.append(s["role"])
    return roles


def _roles_in_message(message: str, roles: List[str]) -> List[str]:
    """Роли процесса, упомянутые в сообщении (в порядке упоминания)."""
    low = message.lower()
    hits: List[Tuple[int, str]] = []
    canonical = _find_role(message)[0]
    for role in roles:
        pos = -1
        if canonical == role:
            pos = 0
        else:
            stems = [w[:6].lower() for w in re.findall(r"[А-Яа-яЁё]{5,}", role)]
            found = [low.find(st) for st in stems if st in low]
            if found:
                pos = min(found)
        if pos >= 0:
            hits.append((pos, role))
    return [r for _, r in sorted(hits)]


# --------------------------- режим 1: аналитика и объяснение --------------------------- #
def _analysis_sla(ctx: Dict[str, Any]) -> str:
    sla, target = ctx["sla"], ctx.get("sla_target_hours")
    if not sla:
        return "Данных аудита пока нет: сгенерируйте диаграмму."
    base, total = float(sla["critical_path_hours"]), float(sla["with_rework_hours"])
    out = [f"**Критический путь — {_fh(base)}**" + (f" при цели {_fh(float(target))}" if target else "") + "."]
    if sla.get("breach"):
        if target and base > float(target):
            out.append(f"🔴 Срыв SLA заложен в самом процессе: даже без возвратов путь длиннее цели на {_fh(base - float(target))}.")
        else:
            out.append(f"🔴 Без возвратов процесс укладывается в срок, но **с худшим циклом возврата** — {_fh(total)}: срыв из-за доработок.")
    elif float(sla.get("rework_share", 0)) > 0.25:
        out.append(f"🟠 В срок укладываемся, но худший возврат добавляет {float(sla['rework_share']):.0%} ({_fh(total)} с возвратом) — запас съедается.")
    else:
        out.append("🟢 Процесс укладывается в срок, риски умеренные.")
    heavy = sorted(ctx["critical_path"], key=lambda c: -float(c["hours"]))[:3]
    if heavy and base > 0:
        out.append("Что занимает больше всего времени на критическом пути:")
        out += [f"- «{c['name']}» ({c['role']}) — {_fh(float(c['hours']))}, {float(c['hours']) / base:.0%} пути" for c in heavy]
    loops = sorted(ctx["rework_loops"], key=lambda l: -float(l["cycle_hours"]))
    if loops:
        l0 = loops[0]
        out.append(f"Самый дорогой возврат — «{l0['label']}» ({l0['from']} → {l0['to']}): {_fh(float(l0['cycle_hours']))}.")
    out.append("Что делать: сократить самые долгие шаги пути, вынести независимые шаги в параллель, ужесточить входной контроль, "
               "чтобы реже возвращать на доработку.")
    if len(ctx["steps"]) >= 5:
        out.append("Например: «Сделай шаги 4 и 5 параллельными» (если они независимы).")
    return "\n".join(out)


def _analysis_load(ctx: Dict[str, Any], roles: List[str]) -> str:
    load = {r["role"]: r for r in ctx["lane_load"]}
    if not load:
        return "Данных о нагрузке пока нет."
    role = roles[0] if roles and roles[0] in load else max(load, key=lambda r: float(load[r]["share"]))
    info, bus = load[role], ctx["bus_factor"]
    thr = float(bus.get("threshold", 0.45))
    mine = [s for s in ctx["steps"] if s["role"] == role]
    out = [f"**{role}: {float(info['share']):.0%} шагов процесса ({int(info['tasks'])} задач)**, порог bus-factor — {thr:.0%}."]
    if float(info["share"]) > thr:
        out.append("🔴 Процесс держится на этой роли: отсутствие одного сотрудника остановит значительную часть шагов.")
    else:
        out.append("🟢 Нагрузка в пределах нормы.")
    if mine:
        out.append("Шаги роли:")
        out += [f"- {s['num']}. {s['title']}" + (f" ({_fh(float(s['hours']))})" if s["hours"] else "") for s in mine[:8]]
    routine = [s for s in mine if re.search(r"регистр|фиксир|вносит|вести|журнал|уведом|направ|передач|заказ", s["title"], re.I)]
    out.append("Как разгрузить:")
    if routine:
        out.append(f"- автоматизировать или передать рутину: «{routine[0]['title']}» (шаг {routine[0]['num']});")
    out.append("- назначить заместителя/резерв на ключевые шаги и описать их в инструкции (реверс-генерация: «Составь инструкцию для роли "
               f"{role}»);")
    others = sorted((r for r in load if r != role), key=lambda r: float(load[r]["share"]))
    if others:
        out.append(f"- перераспределить часть проверок на роль с минимальной загрузкой — «{others[0]}» ({float(load[others[0]]['share']):.0%}).")
    return "\n".join(out)


def _analysis_speedup(ctx: Dict[str, Any]) -> str:
    out = ["**Как ускорить процесс — по убыванию эффекта:**"]
    sla = ctx["sla"]
    heavy = sorted(ctx["critical_path"], key=lambda c: -float(c["hours"]))[:2]
    for i, c in enumerate(heavy, 1):
        out.append(f"{i}. Сократить/автоматизировать «{c['name']}» ({c['role']}, {_fh(float(c['hours']))}) — это на критическом пути.")
    loops = sorted(ctx["rework_loops"], key=lambda l: -float(l["cycle_hours"]))
    if loops:
        out.append(f"{len(heavy) + 1}. Снизить долю возвратов «{loops[0]['label']}»: чек-лист на входе шага «{loops[0]['from']}» "
                   f"(цикл стоит {_fh(float(loops[0]['cycle_hours']))}).")
    pairs = []
    steps = ctx["steps"]
    for a, b in zip(steps, steps[1:]):
        if a["role"] != b["role"] and not a["decision"] and not b["parallel"] and not b["decision"]:
            pairs.append((a, b))
    if pairs:
        a, b = max(pairs, key=lambda p: (p[1]["hours"] or 0))
        out.append(f"{len(heavy) + (1 if loops else 0) + 1}. Проверить независимость шагов {a['num']} и {b['num']} "
                   f"(«{a['role']}» и «{b['role']}») — возможна параллель: «Сделай шаги {a['num']} и {b['num']} параллельными».")
    if sla:
        out.append(f"\nСейчас критический путь — {_fh(float(sla['critical_path_hours']))}.")
    return "\n".join(out)


def _analysis_loops(ctx: Dict[str, Any]) -> str:
    loops = ctx["rework_loops"]
    if not loops:
        return "Циклов возврата на доработку в модели нет — процесс линейный."
    out = [f"**Циклов возврата: {len(loops)}**"]
    out += [f"- «{l['label']}»: {l['from']} → {l['to']} ({l['lane']}), стоимость цикла {_fh(float(l['cycle_hours']))}" for l in loops]
    out.append("Риск: каждый возврат повторяет участок пути. Рекомендация — входной контроль и лимит возвратов с эскалацией руководителю.")
    return "\n".join(out)


def _analysis_landscape(ctx: Dict[str, Any]) -> str:
    sys_, docs = ctx["it_systems"], ctx["artifacts"]
    out = ["**ИТ-ландшафт процесса**"]
    out.append("💻 Системы: " + (", ".join(f"{i['name']} (шаги {', '.join(map(str, i['steps']))})" for i in sys_) or "не упомянуты"))
    out.append("📄 Документы: " + (", ".join(f"{i['name']} (шаги {', '.join(map(str, i['steps']))})" for i in docs) or "не обнаружены"))
    if not sys_:
        out.append("⚠️ Ни один шаг не привязан к информационной системе — риск ручной обработки и потери данных.")
    elif len(sys_) < 2 and len(ctx["steps"]) > 8:
        out.append("Для процесса такого размера систем мало: проверьте, нет ли шагов, которые ведутся «на бумаге».")
    return "\n".join(out)


def _analysis_summary(ctx: Dict[str, Any]) -> str:
    st_ = ctx["stats"]
    sla = ctx["sla"]
    out = [f"**{ctx['title']}**: {len(ctx['steps'])} шагов, {len(_roles_of(ctx))} ролей"
           + (f", подпроцессов: {st_.get('subprocesses')}" if st_ else "") + "."]
    if sla:
        out.append(f"Критический путь {_fh(float(sla['critical_path_hours']))}, срыв SLA: {'да' if sla.get('breach') else 'нет'}.")
    if ctx["bus_factor"]:
        b = ctx["bus_factor"]
        out.append(f"Самая загруженная роль — «{b.get('top_role')}» ({float(b.get('max_share', 0)):.0%}).")
    out.append("Спросите: «В чём причина срыва SLA?», «Как разгрузить диспетчера?», «Как ускорить процесс?»; "
               "попросите изменить процесс («Добавь согласование с экологами после шага 3») "
               "или составить инструкцию для роли.")
    return "\n".join(out)


def _analysis_tobe(ctx: Dict[str, Any]) -> str:
    f = ctx.get("facts") or {}
    if not f.get("tobe_ready"):
        sla = ctx.get("sla") or {}
        loops = ctx.get("rework_loops") or []
        return (
            f"**As-Is: срок с возвратами {_fh(float(sla.get('with_rework_hours') or sla.get('critical_path_hours') or 0))}, "
            f"циклов {len(loops)}.** To-Be ещё не посчитан — откройте вкладку оптимизации."
        )
    rw_b, rw_a = float(f.get("rw_before") or 0), float(f.get("rw_after") or 0)
    cp_b, cp_a = float(f.get("cp_before") or 0), float(f.get("cp_after") or 0)
    rb, ra = int(f.get("loops_before") or 0), int(f.get("loops_after") or 0)
    pct = int(f.get("rw_saved_pct") or 0)
    out: List[str] = []
    if f.get("speedup_via_rework") or rw_a < rw_b:
        out.append(
            f"**Экономия пути с возвратами: {_fh(rw_b)} → {_fh(rw_a)} (−{pct}%), циклы {rb} → {ra}.**"
        )
    else:
        out.append(f"**To-Be: путь с возвратами {_fh(rw_b)} → {_fh(rw_a)}, циклы {rb} → {ra}.**")
    if f.get("cp_grew"):
        out.append(f"Голый критический путь {_fh(cp_b)} → {_fh(cp_a)} — не оптимален, если есть кандидат без удлинения.")
    else:
        out.append(f"Голый критический путь: {_fh(cp_b)} → {_fh(cp_a)}.")
    kinds = {str(a.get("kind")) for a in (ctx.get("tobe_actions") or [])}
    if "zero_rework" in kinds:
        out.append("Zero-Rework: циклы заменены эскалацией на исключительной ветке, не на счастливом пути.")
    if "parallel" in kinds:
        out.append("Параллель: независимые роли идут одновременно, если нет стоп-листа охраны труда.")
    if "automation" in kinds:
        out.append("Журналы переведены в scriptTask.")
    bus_role = f.get("bus_role")
    if bus_role:
        out.append(f"Bus-factor: «{bus_role}» держит {float(f.get('bus_share') or 0):.0%} шагов.")
    if abs(rw_a - cp_a) < 0.051:
        out.append("Одинаковые цифры справа: циклы сняты, добавки за возврат нет.")
    return "\n".join(out)


_ASK_STOP = {
    "что", "как", "про", "для", "это", "шаг", "шага", "шаге", "шаги", "шагов", "пункт", "пункта",
    "делает", "делают", "почему", "цифры", "цифра", "одинаковые", "одинаковых", "время", "срок",
    "процесс", "схемы", "схеме", "схему", "расскажи", "текущей", "какие", "какой", "какая",
    "роль", "роли", "ролей", "часов", "часы", "минут",
}
_FACTS_Q_RE = re.compile(
    r"одинаков|сравни|as-is|as is|to-be|tobe|до и после|экономи|rework|"
    r"цикл|возврат|сколько\s+времени",
    re.I,
)
_ROLE_Q_RE = re.compile(r"рол|какие\s+шаги|что\s+делает|перечисл|шаги\s+|нагрузк|разгруз", re.I)
_STEP_WORD_RE = re.compile(r"шаг\w*|пункт\w*|п\.\s*\d+|этап\w*", re.I)


def _explicit_step_num(message: str) -> Optional[int]:
    m = re.search(r"(?:шаг\w*|пункт\w*|этап\w*|п\.)\s*№?\s*(\d+)", message or "", re.I)
    if not m:
        m = re.search(r"\b(\d+)\s*(?:-?(?:й|ый|ой|го|м))?\s*шаг", message or "", re.I)
    return int(m.group(1)) if m else None


def _step_bodies(ctx: Dict[str, Any]) -> Dict[int, str]:
    reg = ctx.get("regulation_text") or ""
    if not str(reg).strip():
        return {}
    try:
        _, raw = _split_steps(normalize_regulation(reg))
    except Exception:  # noqa: BLE001
        return {}
    return {int(s["num"]): str(s.get("body") or "") for s in raw if s.get("num") is not None}


def _ask_tokens(message: str) -> List[str]:
    toks = []
    for w in re.findall(r"[А-Яа-яЁёA-Za-z]{3,}", message or ""):
        low = w.lower().replace("ё", "е")
        if low in _ASK_STOP:
            continue
        toks.append(low)
    return toks


def _blob_has_token(blob: str, token: str) -> bool:
    b = (blob or "").lower().replace("ё", "е")
    if token in b:
        return True
    stem = token[:5]
    return len(stem) >= 4 and stem in b


def _match_asked_step(message: str, ctx: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    steps = list(ctx.get("steps") or [])
    if not steps:
        return None
    bodies = _step_bodies(ctx)
    tokens = _ask_tokens(message)
    scored: List[Tuple[int, int, Dict[str, Any]]] = []
    for s in steps:
        blob = " ".join([str(s.get("title") or ""), str(s.get("role") or ""), bodies.get(int(s["num"]), "")])
        hit = [t for t in tokens if _blob_has_token(blob, t)]
        if hit:
            scored.append((len(hit), max(len(t) for t in hit), s))
    scored.sort(key=lambda pair: (-pair[0], -pair[1], int(pair[2]["num"])))
    if not scored:
        return None
    if len(scored) > 1 and scored[0][0] == scored[1][0]:
        return None
    best_n, best_len, best = scored[0]
    if best_n >= 2:
        return best
    if best_len >= 6 and re.search(r"шаг|что\s+делает|про\s+|подготов", message or "", re.I):
        return best
    return None


def _one_step_reply(step: Dict[str, Any], ctx: Dict[str, Any]) -> str:
    num = int(step["num"])
    title = str(step.get("title") or "").strip() or "—"
    role = str(step.get("role") or "").strip() or "—"
    hours = step.get("hours")
    tail = f", {_fh(float(hours))}" if hours else ""
    body = _step_bodies(ctx).get(num, "")
    lines = [f"**Шаг {num}. «{title}»** — {role}{tail}."]
    if body and title.lower() not in body.lower():
        snip = re.sub(r"^(?:параллельно|одновременно)\s*[:,—–-]?\s*", "", body, flags=re.I)
        snip = re.sub(r"\s*\([^)]*\)\s*$", "", snip).strip()
        if snip:
            lines.append(f"В регламенте: «{snip[:140]}».")
    return "\n".join(lines)


def _role_steps_reply(ctx: Dict[str, Any], role: str) -> str:
    mine = [s for s in (ctx.get("steps") or []) if s.get("role") == role]
    lines = [f"**«{role}»** — шаги по порядку:"]
    if not mine:
        lines.append("Шагов этой роли на схеме нет.")
        return "\n".join(lines)
    for s in mine:
        hours = s.get("hours")
        tail = f" — {_fh(float(hours))}" if hours else ""
        lines.append(f"{s['num']}. «{s.get('title') or '—'}»{tail}")
    return "\n".join(lines)


def _steps_miss_reply(ctx: Dict[str, Any]) -> str:
    f = ctx.get("facts") or {}
    n = len(ctx.get("steps") or [])
    cp = float(f.get("critical_path_hours") or (ctx.get("sla") or {}).get("critical_path_hours") or 0)
    title = str(ctx.get("title") or "процесс")
    return f"Шаг не найден. В схеме «{title}» {n} шагов, критический путь {_fh(cp)}."


def _facts_time_reply(ctx: Dict[str, Any]) -> str:
    """Срок, экономия, циклы — только числа process_facts, без подстановки шага 1."""
    f = ctx.get("facts") or {}
    if not f.get("tobe_ready"):
        rw = float(f.get("with_rework_hours") or 0)
        cp = float(f.get("critical_path_hours") or 0)
        n = int(f.get("rework_loops_n") or 0)
        return (
            f"**As-Is: срок с возвратами {_fh(rw)}, голый КП {_fh(cp)}, циклов {n}.** "
            "To-Be ещё не посчитан."
        )
    rw_b, rw_a = float(f.get("rw_before") or 0), float(f.get("rw_after") or 0)
    cp_b, cp_a = float(f.get("cp_before") or 0), float(f.get("cp_after") or 0)
    rb, ra = int(f.get("loops_before") or 0), int(f.get("loops_after") or 0)
    pct = int(f.get("rw_saved_pct") or 0)
    lines = [
        f"**Экономия пути с возвратами: {_fh(rw_b)} → {_fh(rw_a)} (−{pct}%), циклы «{rb} → {ra}».**",
        f"Голый критический путь: {_fh(cp_b)} → {_fh(cp_a)}.",
    ]
    if abs(rw_a - cp_a) < 0.051 or ra == 0:
        lines.append("Одинаковые цифры справа: циклы сняты, добавки за возврат нет.")
    return "\n".join(lines)


def _visible_answer(message: str, ctx: Dict[str, Any]) -> Optional[str]:
    """Шаг, роль или цифры To-Be по тексту вопроса. None — вопрос не про видимый шаг и не про срок."""
    steps = list(ctx.get("steps") or [])
    num = _explicit_step_num(message)
    if num is not None:
        hit = next((s for s in steps if int(s["num"]) == num), None)
        return _one_step_reply(hit, ctx) if hit else _steps_miss_reply(ctx)
    hit = _match_asked_step(message, ctx)
    if hit is not None and not _FACTS_Q_RE.search(message or ""):
        return _one_step_reply(hit, ctx)
    if _FACTS_Q_RE.search(message or ""):
        return _facts_time_reply(ctx)
    roles = _roles_in_message(message, _roles_of(ctx))
    if roles and _ROLE_Q_RE.search(message or ""):
        return _role_steps_reply(ctx, roles[0])
    if _STEP_WORD_RE.search(message or "") and hit is None:
        return _steps_miss_reply(ctx)
    return None


def _analysis_readability(ctx: Dict[str, Any]) -> str:
    meth = ctx.get("methodology") or {}
    stats = ctx.get("stats") or {}
    score = int(meth.get("score") or 0)
    subs = int(stats.get("subprocesses") or 0)
    layers = int(stats.get("critical_path_layers") or 0)
    out = [f"**Читаемость BPMN «{ctx.get('title')}».** Quality Score: **{score}%**."]
    if subs:
        out.append(f"Декомпозиция: {subs} подпроцесс(а) — это как раз против «метро Токио».")
    else:
        out.append("Подпроцессов нет: монотонные цепочки одной роли лучше свернуть в subprocess.")
    if layers:
        out.append(f"Слоёв на критическом пути: {layers}.")
    for chk in meth.get("checks") or []:
        mark = "✓" if chk.get("passed") else "✗"
        out.append(f"- {mark} {chk.get('title')}: {chk.get('detail')}")
    if ctx.get("tobe_ready"):
        out.append(
            f"To-Be: Quality {ctx.get('quality_score_as_is')}% → {ctx.get('quality_score_to_be')}%, "
            f"петли {ctx.get('rework_loops_as_is')} → {ctx.get('rework_loops_to_be')} — меньше возвратных стрелок, схема читается линейнее."
        )
    return "\n".join(out)


def _analysis_open(message: str, ctx: Dict[str, Any]) -> str:
    """Развёрнутый ответ по архитектуре процесса — не статичная заглушка с числом шагов."""
    q = (message or "").strip()
    sla = ctx.get("sla") or {}
    bus = ctx.get("bus_factor") or {}
    meth = ctx.get("methodology") or {}
    recs: List[str] = []
    for item in ctx.get("recommendations") or []:
        recs.append(item if isinstance(item, str) else str((item or {}).get("message") or ""))
    for item in ctx.get("sla_risks") or []:
        if isinstance(item, dict) and item.get("message"):
            recs.append(str(item["message"]))
    recs = [r for r in recs if r]
    out: List[str] = []
    if q:
        out.append(f"По запросу «{q[:140]}» — факты аудита, не общая заглушка.")
    else:
        out.append(f"Архитектурный разбор «{ctx.get('title')}».")
    if sla:
        out.append(
            f"SLA: критический путь {_fh(float(sla.get('critical_path_hours') or 0))}, "
            f"с худшим возвратом {_fh(float(sla.get('with_rework_hours') or 0))}, "
            f"срыв: {'да' if sla.get('breach') else 'нет'}."
        )
    if bus:
        out.append(f"Ключевая роль: «{bus.get('top_role')}» ({float(bus.get('max_share') or 0):.0%} шагов).")
    if meth.get("score") is not None:
        out.append(f"Линтер нотации: Quality Score {int(meth.get('score') or 0)}%.")
    loops = ctx.get("rework_loops") or []
    if loops:
        l0 = max(loops, key=lambda x: float((x or {}).get("cycle_hours") or 0))
        out.append(f"Дороже всего цикл «{l0.get('label')}»: {_fh(float(l0.get('cycle_hours') or 0))}.")
    if ctx.get("tobe_ready"):
        f = ctx.get("facts") or {}
        out.append(
            f"To-Be: путь с возвратами {_fh(float(f.get('rw_before') or ctx.get('sla_hours_as_is') or 0))} → "
            f"{_fh(float(f.get('rw_after') or ctx.get('sla_hours_to_be') or 0))} "
            f"(циклы {f.get('loops_before')} → {f.get('loops_after')})."
        )
    if recs:
        out.append("Что делать архитектору:")
        out += [f"- {r}" for r in recs[:3]]
    out.append("Уточните: сравнение As-Is/To-Be, читаемость (анти-метро), роли или ускорение SLA.")
    return "\n".join(out)


_QUALITY_Q_RE = re.compile(
    r"quality|качеств\w+\s+(нотац|score)|почему.{0,50}(score|балл|нотац|линтер)|"
    r"(упал|изменил|просел)\w*.{0,30}(quality|балл|нотац)|линтер\w*\s+нотац",
    re.I,
)
_WHY_SAVED_RE = re.compile(
    r"как\s+(мы\s+)?сократ|за сч[её]т чего|за счет чего|объясни подробн|"
    r"почему.{0,40}(сократ|ускор|экономи|to-be|tobe)|"
    r"как.{0,40}(ускор\w+\s+процесс|экономи)|декомпозиц\w+\s+экономи|"
    r"в ч[её]м причина.{0,30}(сократ|экономи|to-be|tobe)",
    re.I,
)
_FACT_OVERRIDE_RE = re.compile(
    r"sla|срок|срыв|as-is|as is|to-be|tobe|цикл|возврат|bus.?factor|ускор|экономи",
    re.I,
)


def _suggested_command(ctx: Dict[str, Any]) -> str:
    loops = ctx.get("rework_loops") or []
    steps = ctx.get("steps") or []
    if loops:
        return "Добавь входной контроль перед возвратом на доработку"
    if len(steps) >= 5:
        return "Сделай шаги 4 и 5 параллельными"
    return "Добавь согласование с экологами после шага 3"


def _sidebar_shape(body: str, ctx: Dict[str, Any], intent: str) -> str:
    """Сайдбар: вывод → шаги/роли → одна команда в ёлочках. Не больше 10 строк."""
    lines = [ln.rstrip() for ln in (body or "").splitlines() if ln.strip() and not ln.strip().startswith("<sub>")]
    table = any("| Изменение |" in ln or "Подтверждено регламентом" in ln for ln in lines)
    if table:
        return "\n".join(lines[:24])
    if intent in ("analysis", "next_step"):
        lines = lines[:9]
        if not any("«" in ln and "»" in ln for ln in lines):
            lines.append(f"«{_suggested_command(ctx)}»")
        return "\n".join(lines[:10])
    return "\n".join(lines[:12])


def _analysis_why_saved(ctx: Dict[str, Any]) -> str:
    """Декомпозиция экономии SLA To-Be: петли, параллель, scriptTask — не заглушка читаемости."""
    if not ctx.get("tobe_ready"):
        return (
            "**To-Be ещё не посчитан.** Откройте вкладку «Оптимизация As-Is → To-Be» — "
            "разложу экономию по петлям доработки, параллелизации шагов и автоматизации (scriptTask)."
        )
    as_h = float(ctx.get("sla_hours_as_is") or 0)
    to_h = float(ctx.get("sla_hours_to_be") or 0)
    saved = float(ctx.get("sla_saved_hours") or max(0.0, as_h - to_h))
    pct = float(ctx.get("delta_sla_percent") or 0)
    rw_h_as = float(ctx.get("rework_hours_as_is") or 0)
    rw_h_to = float(ctx.get("rework_hours_to_be") or 0)
    loop_saved = max(0.0, rw_h_as - rw_h_to)
    rb, ra = int(ctx.get("rework_loops_as_is") or 0), int(ctx.get("rework_loops_to_be") or 0)
    actions = [a for a in (ctx.get("tobe_actions") or []) if isinstance(a, dict)]
    par = [a for a in actions if a.get("kind") == "parallel"]
    auto = [a for a in actions if a.get("kind") == "automation"]
    ctrl = [a for a in actions if a.get("kind") == "zero_rework"]
    rest = max(0.0, saved - loop_saved)
    out = [
        f"**Как сократили SLA «{ctx.get('title')}»: {_fh(as_h)} → {_fh(to_h)} "
        f"(экономия {_fh(saved)}, {pct:g}%).**",
        "Декомпозиция выигрыша:",
        f"1) **Устранение петель доработки:** {rb} → {ra} "
        f"(снято {max(0, rb - ra)} цикл.). Экономия циклов: **{_fh(loop_saved)}** "
        f"(худший возврат {_fh(rw_h_as)} → {_fh(rw_h_to)}).",
    ]
    f = ctx.get("facts") or {}
    if f.get("cp_grew"):
        out.append(
            f"   Голый путь {_fh(float(f.get('cp_before') or 0))} → {_fh(float(f.get('cp_after') or 0))}."
        )
    loops = sorted(ctx.get("rework_loops") or [], key=lambda x: -float((x or {}).get("cycle_hours") or 0))
    for item in loops[:3]:
        out.append(f"   - цикл As-Is «{item.get('label')}»: {_fh(float(item.get('cycle_hours') or 0))}")
    par_txt = "; ".join(str(a.get("detail") or "").rstrip(".") for a in par if a.get("detail")) or (
        "независимые шаги разных ролей вынесены в параллельный AND-шлюз"
    )
    out.append(f"2) **Параллелизация конкретных шагов:** {par_txt}. "
               f"Выигрыш на критическом пути (без петель): **{_fh(rest)}**.")
    if auto:
        auto_txt = "; ".join(str(a.get("detail") or "") for a in auto if a.get("detail"))
        out.append(
            f"3) **Автоматизация рутины (scriptTask):** {auto_txt}. "
            "Фиксация в журналах выполняется информационной системой (обычно 5 мин вместо ручного шага)."
        )
    else:
        out.append("3) **Автоматизация рутины (scriptTask):** отдельных журнальных шагов к автоматизации не было.")
    for act in ctrl[:2]:
        if act.get("detail"):
            out.append(f"- Zero-Rework: {act['detail']}")
    return "\n".join(out)


def _analysis_quality_score(ctx: Dict[str, Any]) -> str:
    """Почему изменился Quality Score — не анти-метро и не заглушка читаемости."""
    meth = ctx.get("methodology") or {}
    now = int(meth.get("score") or 0)
    qb = int(ctx.get("quality_score_as_is") or now)
    qa = int(ctx.get("quality_score_to_be") or now)
    naming = next((c for c in (meth.get("checks") or []) if c.get("key") == "naming"), None)
    if ctx.get("tobe_ready"):
        if qa >= 100 and qb >= 100:
            return (
                f"**Quality Score нотации: {qb}% → {qa}%.** Стандарты соблюдены на 100%: "
                "все задачи To-Be в форме «глагол в инфинитиве + объект», шлюзы подписаны. "
                "Параллельные шлюзы и входной контроль добавлены без поломки Naming Compliance."
            )
        extra = f"\nNaming: {naming.get('detail')}" if naming else ""
        if qa < qb:
            return (
                f"**Quality Score: {qb}% → {qa}%.** При реинжиниринге добавлены шлюзы параллелизации "
                "и шаги входного контроля. Правило «глагол + объект» для задач сохранено; "
                "расхождение даёт линтер на новых элементах, а не отказ от стандарта."
                + extra
            )
        return (
            f"**Quality Score: {qb}% → {qa}%.** Нотация не просела: глагол + объект сохранены, "
            "добавлены только параллельные шлюзы и контроли."
        )
    if now >= 100:
        return (
            f"**Quality Score текущей схемы: {now}%.** Стандарты нотации соблюдены на 100% "
            "(глагол в инфинитиве + объект)."
        )
    return _analysis_readability(ctx)


_OPTIMIZER_MOVE = "это ход оптимизатора, в регламенте такой формулировки нет"


def _audit_cell(value: Any) -> str:
    return str(value if value is not None and value != "" else "—").replace("|", "\\|").replace("\n", " ").strip() or "—"


def _audit_snip(text: str, limit: int = 88) -> str:
    t = re.sub(r"\s+", " ", (text or "")).strip()
    t = re.sub(r"^\d+[.)]\s*", "", t)
    if len(t) > limit:
        t = t[: limit - 1].rstrip(" ,;:") + "…"
    return t


def _analysis_opt_audit(ctx: Dict[str, Any]) -> str:
    """Сверка To-Be с исходным регламентом: таблица, без команд правки XML."""
    facts = ctx.get("facts") or {}
    steps = ctx.get("steps") or []
    reg = ctx.get("regulation_text") or ""
    try:
        _, raw = _split_steps(normalize_regulation(reg)) if reg.strip() else ([], [])
    except Exception:  # noqa: BLE001
        raw = []
    bodies = {int(s["num"]): s["body"] for s in raw if s.get("num") is not None}
    titles: Dict[int, str] = {}
    for s in steps:
        try:
            titles[int(s["num"])] = str(s.get("title") or "").strip()
        except (TypeError, ValueError):
            continue
    for num, body in bodies.items():
        if not titles.get(num):
            titles[num] = _step_title(body)

    def step_cell(num: Optional[int], extra: str = "") -> str:
        if num is None:
            return extra or "—"
        body = bodies.get(num, "")
        name = extra
        if body:
            name = _audit_snip(
                re.sub(r"^(?:параллельно|одновременно)\s*[:,—–-]?\s*", "", body, flags=re.I),
                72,
            )
            name = re.sub(r"\s*\(\d+[^\)]*\)\s*$", "", name).rstrip(" .")
        if not name:
            name = titles.get(num) or _step_title(body)
        return f"{num}. {name}" if name else str(num)

    rows: List[List[str]] = []
    ppe_num: Optional[int] = None
    for s in raw:
        body = s.get("body") or ""
        if _is_prepare_ppe_ground(body):
            ppe_num = int(s["num"])
            quote = ""
            m = re.search(r"(Параллельно\s*[:,—–-]?\s*.{0,80})", body, re.I)
            if m:
                quote = _audit_snip(m.group(1))
            elif re.search(r"параллельно", body, re.I):
                quote = _audit_snip(body)
            if quote:
                rows.append(
                    [
                        "Подготовка СИЗ, инструмента и переносных заземлений параллельно оперативным переключениям",
                        step_cell(ppe_num),
                        "да",
                        f"«{quote}»",
                        "низкий",
                        "высокая",
                    ]
                )
            break

    repair_num = brief_num = None
    for s in raw:
        body = s.get("body") or ""
        num = int(s["num"])
        if repair_num is None and _REPAIR_WORK_RE.search(body) and not re.search(r"закрыва", body, re.I):
            repair_num = num
        if re.search(r"инструктаж", body, re.I):
            brief_num = num
    if repair_num is not None and brief_num is not None and repair_num > brief_num:
        quotes = [_audit_snip(bodies[brief_num], 80), _audit_snip(bodies[repair_num], 80)]
        rows.append(
            [
                "Ремонт остаётся после допуска и целевого инструктажа",
                step_cell(repair_num),
                "да",
                "«" + "; ".join(quotes) + "»",
                "низкий",
                "высокая",
            ]
        )

    zero_nums: List[int] = []
    for act in ctx.get("tobe_actions") or []:
        if not isinstance(act, dict):
            continue
        kind = str(act.get("kind") or "")
        detail = str(act.get("detail") or "")
        found = [int(x) for x in re.findall(r"шаг(?:и)?\s+(\d+)", detail, re.I)]
        if kind == "parallel":
            ppe_hit = ppe_num is not None and (ppe_num in found or any(_is_prepare_ppe_ground(bodies.get(n, "")) for n in found))
            already = False
            if len(found) >= 2 and found[1] in bodies and _is_parallel_body(bodies[found[1]]):
                already = True
            if ppe_hit or already:
                continue
            rows.append(
                [
                    "Параллельное выполнение независимых шагов",
                    " и ".join(step_cell(n) for n in found[:2]) if found else "—",
                    "нет",
                    _OPTIMIZER_MOVE,
                    "средний",
                    "средняя",
                ]
            )
        elif kind == "zero_rework":
            zero_nums.extend(found[:1] or [])
        elif kind == "automation":
            n = found[0] if found else None
            body = bodies.get(n, "") if n else ""
            if n and re.search(r"систем\w+\s+автоматическ|автоматически:", body, re.I):
                rows.append(
                    [
                        "Автоматизация журнальной фиксации",
                        step_cell(n),
                        "да",
                        f"«{_audit_snip(body)}»",
                        "низкий",
                        "высокая",
                    ]
                )
            else:
                rows.append(
                    [
                        "Автоматизация журнальной фиксации",
                        step_cell(n) if n else "—",
                        "нет",
                        _OPTIMIZER_MOVE,
                        "низкий",
                        "средняя",
                    ]
                )

    if zero_nums or (
        facts.get("tobe_ready")
        and facts.get("loops_after") is not None
        and int(facts.get("loops_before") or 0) > int(facts.get("loops_after") or 0)
    ):
        uniq = list(dict.fromkeys(zero_nums))
        step_lbl = ", ".join(step_cell(n) for n in uniq[:3]) if uniq else "циклы возврата"
        rows.append(
            [
                "Замена циклов возврата на эскалацию",
                step_lbl,
                "нет",
                _OPTIMIZER_MOVE,
                "средний",
                "средняя",
            ]
        )

    if not rows:
        rows.append(
            [
                "Сверка As-Is и To-Be",
                "—",
                "нет",
                _OPTIMIZER_MOVE,
                "средний",
                "низкая",
            ]
        )

    head: List[str] = []
    if facts.get("tobe_ready"):
        lb, la = int(facts.get("loops_before") or 0), int(facts.get("loops_after") or 0)
        rw_b, rw_a = float(facts.get("rw_before") or 0), float(facts.get("rw_after") or 0)
        cp_b = float(facts.get("cp_before") or 0)
        cp_a = float(facts.get("cp_after") or 0)
        head.append(
            f"Циклы {lb} → {la}; срок с возвратами {_fh(rw_b)} → {_fh(rw_a)}; "
            f"голый КП {_fh(cp_b)} → {_fh(cp_a)}."
        )
    else:
        head.append("To-Be ещё не посчитан — сверка шагов с исходным регламентом.")
    table = [
        "| Изменение | Шаг (номер и название) | Подтверждено регламентом | Основание | Риск | Уверенность |",
        "|---|---|---|---|---|---|",
    ]
    for row in rows:
        table.append("| " + " | ".join(_audit_cell(c) for c in row) + " |")
    return "\n".join(head + table)


def heuristic_analysis(message: str, ctx: Dict[str, Any]) -> str:
    low = message.lower()
    roles = _roles_in_message(message, _roles_of(ctx))
    if _OPT_AUDIT_Q_RE.search(low):
        return _analysis_opt_audit(ctx)
    if _QUALITY_Q_RE.search(low):
        return _analysis_quality_score(ctx)
    if _WHY_SAVED_RE.search(low) and not re.search(r"одинаков", low):
        return _analysis_why_saved(ctx)
    visible = _visible_answer(message, ctx)
    if visible:
        return visible
    if re.search(r"сравни|as-is|as is|to-be|tobe|до и после|до/после|целев\w+\s+схем|реинжинир", low):
        return _analysis_tobe(ctx)
    if re.search(r"читаем|метро|нотаци|подпроцесс|анти-метро|методолог", low):
        return _analysis_readability(ctx)
    if re.search(r"sla|срок|срыв|задерж|критическ|длительн|долго|почему.*(долго|медлен)", low):
        return _analysis_sla(ctx)
    if re.search(r"цикл|возврат|доработ|rework|замечани", low):
        return _analysis_loops(ctx)
    if re.search(r"ускор|оптимиз|быстрее|повысить эффективн|улучш", low):
        return _analysis_speedup(ctx)
    if re.search(r"систем|it\b|ит-|scada|документ|артефакт|ландшафт|crm|1с|sap", low):
        return _analysis_landscape(ctx)
    if re.search(r"нагрузк|bus|загружен|разгруз|перегруз|исполнител|кто.*(ключев|зависит)", low) or roles:
        return _analysis_load(ctx, roles)
    if re.search(r"риск|узк|проблем|слаб|разбор", low):
        return _analysis_sla(ctx) + "\n\n" + _analysis_load(ctx, [])
    return _analysis_open(message, ctx)


# --------------------------- «Предложить следующий шаг» --------------------------- #
def suggest_next_steps(ctx: Dict[str, Any]) -> str:
    steps = ctx["steps"]
    if not steps:
        return "Процесс пока пуст — сгенерируйте диаграмму."
    last = steps[-1]
    text_all = " ".join(f"{s['title']} {s['role']}" for s in steps).lower()
    ideas: List[Tuple[str, str]] = []
    if not re.search(r"уведом|информир|сообщ", text_all):
        ideas.append((f"уведомить заинтересованные стороны о результате",
                      f"Добавь уведомление заявителя после шага {last['num']}"))
    if not ctx["it_systems"]:
        ideas.append(("зафиксировать результат в информационной системе (сейчас процесс не привязан к системам)",
                      f"Добавь фиксацию результата в информационной системе после шага {last['num']}"))
    if ctx["rework_loops"] and not re.search(r"эскалац|руководител", text_all):
        ideas.append(("добавить эскалацию руководителю после повторного возврата (ограничить циклы доработки)",
                      f"Добавь эскалацию руководителю после шага {max(1, last['num'] - 1)}"))
    if not re.search(r"контрол|проверк|аудит|приёмк|приемк", " ".join(s["title"] for s in steps[-3:]).lower()):
        ideas.append(("финальный контроль результата перед закрытием процесса",
                      f"Добавь контроль результата руководителем после шага {last['num']}"))
    if not re.search(r"архив|хранени", text_all):
        ideas.append(("архивировать документы процесса", f"Добавь архивирование документов после шага {last['num']}"))
    ideas = ideas[:3] or [("процесс выглядит завершённым; можно добавить шаг оценки удовлетворённости заявителя",
                           f"Добавь уведомление заявителя после шага {last['num']}")]
    out = [f"**Последний шаг сейчас — {last['num']}. «{last['title']}» ({last['role']}).** Что логично добавить дальше:"]
    for i, (why, cmd) in enumerate(ideas, 1):
        out.append(f"{i}. {why[:1].upper() + why[1:]}.  \n   Команда: «{cmd}»")
    return "\n".join(out)


# --------------------------- режим 2: правка регламента на лету --------------------------- #
_STEP_LINE_RE = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
_REF_STRICT = re.compile(r"(?<![A-Za-zА-Яа-яЁё])(п(?:ункт\w*|\.|п\.)?\s*|шаг\w*\s*)(\d+)", re.I)
_HEADER_RE = re.compile(
    r"^(?:регламент|название|процесс)\s*[:—–-]|^(?:целев\w+[^:\n]{0,40}|sla)\s*[:—–-]",
    re.I,
)
_NOUN_VERBS: List[Tuple[str, str, str]] = [  # (основа существительного, 3 л. ед.ч., вин. падеж)
    ("согласован", "согласовывает", "согласование"), ("проверк", "проверяет", "проверку"),
    ("утвержден", "утверждает", "утверждение"), ("уведомлен", "уведомляет", "уведомление"),
    ("подписан", "подписывает", "подписание"), ("рассмотрен", "рассматривает", "рассмотрение"),
    ("оценк", "оценивает", "оценку"), ("экспертиз", "проводит экспертизу", "экспертизу"),
    ("контрол", "контролирует", "контроль"), ("регистрац", "регистрирует", "регистрацию"),
    ("информирован", "информирует", "информирование"), ("анализ", "анализирует", "анализ"),
    ("приёмк", "принимает", "приёмку"), ("приемк", "принимает", "приёмку"), ("оплат", "оплачивает", "оплату"),
    ("инструктаж", "проводит инструктаж", "инструктаж"), ("аудит", "проводит аудит", "аудит"),
    ("архивирован", "архивирует", "архивирование"), ("эскалац", "эскалирует", "эскалацию"),
    ("фиксац", "фиксирует", "фиксацию"),
]


def _step_title(body: str) -> str:
    """Короткое название шага для отчёта о правках (без «Параллельно:», роли и срока)."""
    b = re.sub(r"^(?:параллельно|одновременно)\s*[:,—–-]?\s*", "", body.strip(), flags=re.I)
    b = _parse_decision(_DUR_RE.sub(" ", b))[0] or b
    role, rs, re_ = _find_role(b)
    if role and rs == 0:
        b = b[re_:]
    return _task_title(b)[:80]


def _split_steps(text: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Текст регламента → (строки заголовка, [{num, body}]). Проза без нумерации режется на предложения."""
    text = _segment_unnumbered_prose(text or "")
    header: List[str] = []
    steps: List[Dict[str, Any]] = []
    for line in text.replace("\r\n", "\n").split("\n"):
        if not line.strip():
            continue
        m = _STEP_LINE_RE.match(line)
        if m:
            steps.append({"num": int(m.group(1)), "body": m.group(2).strip()})
        elif steps:
            steps[-1]["body"] += " " + line.strip()
        elif _HEADER_RE.match(line.strip()) or not header:
            header.append(line.strip())
        else:
            steps.append({"num": len(steps) + 1, "body": line.strip()})
    if len(steps) < 2:
        body = " ".join(l.strip() for l in text.split("\n") if l.strip() and not _HEADER_RE.match(l.strip()))
        sentences = [s.strip() for s in re.split(r"(?<=[.;!?])\s+(?=[А-ЯЁA-Z«])", body) if len(s.strip()) > 8]
        steps = [{"num": i, "body": s} for i, s in enumerate(sentences, 1)]
        header = [l.strip() for l in text.split("\n") if _HEADER_RE.match(l.strip())]
    return header, steps


def _join_steps(header: List[str], steps: List[Dict[str, Any]]) -> str:
    return "\n".join([*header, "", *[f"{i}. {s['body']}" for i, s in enumerate(steps, 1)]]).strip() + "\n"


def _remap(body: str, fn: Callable[[int], int]) -> str:
    return _REF_STRICT.sub(lambda m: f"{m.group(1)}{fn(int(m.group(2)))}", body)


def _nums_in(fragment: str) -> List[int]:
    nums: List[int] = []
    for m in re.finditer(r"(\d+)\s*(?:[-–—]|по)\s*(\d+)|(\d+)", fragment):
        if m.group(1):
            nums += list(range(int(m.group(1)), int(m.group(2)) + 1))
        else:
            nums.append(int(m.group(3)))
    return list(dict.fromkeys(nums))


def _step_refs(message: str) -> List[int]:
    m = re.search(r"(?:шаг\w*|этап\w*|пункт\w*|п\.)\s*((?:\d+\s*(?:[-–—]|по|,|и)?\s*)+)", message, re.I)
    return _nums_in(m.group(1)) if m else []


def _default_duration(steps: List[Dict[str, Any]]) -> str:
    hours = []
    for s in steps:
        d = _DUR_RE.search(s["body"])
        if d:
            hours.append(_to_hours(d.group(1), d.group(2)))
    hours.sort()
    median = hours[len(hours) // 2] if hours else 1.0
    return "(1 рабочий день)" if median >= 8 else ("(2 часа)" if median >= 1 else "(30 минут)")


def _so_or_s(word: str) -> str:
    """Предлог «со» перед творительным на с, з, ж, ш, щ; иначе «с»."""
    ch = (word or "").lstrip().lower()[:1]
    return "со" if ch in "сзжшщ" else "с"


def _role_instrumental(role: str) -> str:
    """«Служба экологии» → «службой экологии» (творительный для «с …»)."""
    words = role.split()
    first = words[0]
    low = first.lower()
    if low.endswith("а"):
        first = first[:-1] + "ой"
    elif low.endswith("я"):
        first = first[:-1] + "ей"
    elif low.endswith("ь"):
        first = first[:-1] + "ью"
    elif not low.endswith(("ом", "ем", "ой", "ей", "ью", "ами", "ями")):
        first = first + "ом"
    words[0] = first[:1].lower() + first[1:]
    return " ".join([words[0]] + [w[:1].lower() + w[1:] for w in words[1:]])


def _build_new_step(content: str, anchor_role: str, duration: str) -> str:
    content = re.sub(r"^\s*(?:нов\w+\s+)?(?:шаг|этап|пункт)\s*[:—–-]?\s*", "", content.strip(" .:;«»\"'"), flags=re.I).strip()
    role, rs, re_ = _find_role(content)
    has_verb = any(_is_verb(w) for w in re.findall(r"[А-Яа-яЁё]+", content))
    if role and has_verb and rs == 0:
        body = content
    else:
        noun = next(((stem, v3, acc) for stem, v3, acc in _NOUN_VERBS if re.search(r"\b" + stem, content, re.I)), None)
        rest = content
        with_whom = ""
        extra_role = role
        if extra_role:
            # «с экологами» / «для СБ» — объект действия, не исполнитель; предлог перед ролью тоже вырезаем.
            lead = content[:rs]
            m_with = re.search(r"(с|со|у|от|для|силами|через)\s*$", lead, re.I)
            if m_with:
                prep = m_with.group(1).lower()
                whom = _role_instrumental(extra_role) if prep in ("с", "со") else extra_role[:1].lower() + extra_role[1:]
                if prep in ("с", "со"):
                    prep = _so_or_s(whom)
                with_whom = f"{prep} {whom}"
                rest = (lead[: m_with.start()] + " " + content[re_:]).strip()
                extra_role = None
            elif extra_role and extra_role != anchor_role:
                # «проверку охраны труда» — тема шага, исполнитель остаётся якорем
                rest = content
                extra_role = None
            else:
                rest = (lead + " " + content[re_:]).strip()
            rest = re.sub(r"^\s*(?:с|со|у|от|для|силами|через)\s+", "", rest, flags=re.I).strip()
        who = extra_role or anchor_role
        if noun:
            stem, v3, acc = noun
            rest = re.sub(r"\b" + stem + r"\w*", "", rest, flags=re.I).strip(" ,.;")
            if with_whom:
                body = f"{who} {v3} {with_whom}" + (f" {rest}" if rest else "")
            elif not rest:
                body = f"{who} {v3} результаты предыдущего шага"
            else:
                body = f"{who} проводит {acc} {rest}"
        elif has_verb and not role:
            body = f"{who} {content[:1].lower() + content[1:]}"
        else:
            body = f"{who} выполняет: {content[:1].lower() + content[1:]}"
    body = body[:1].upper() + body[1:]
    return f"{body} {duration}."


@dataclass
class EditResult:
    text: str
    changes: List[str]


def heuristic_edit(message: str, text: str) -> Tuple[Optional[EditResult], str]:
    """Детерминированная правка регламента по команде. Возвращает (результат | None, пояснение для пользователя)."""
    header, steps = _split_steps(normalize_regulation(text))
    if len(steps) < 2:
        return None, "В регламенте не удалось выделить шаги — пронумеруйте их («1. Диспетчер принимает …»)."
    low = message.lower()
    n = len(steps)
    changes: List[str] = []
    refs = [r for r in _step_refs(message) if 1 <= r <= n]

    def by_phrase(phrase: str) -> Optional[int]:
        stems = {w[:5].lower() for w in re.findall(r"[А-Яа-яЁё]{4,}", phrase)} - {"шаг", "этап", "пункт", "шага", "этапа"}
        scored = []
        for i, s in enumerate(steps, 1):
            b = s["body"].lower()
            score = sum(1 for st in stems if st in b)
            if score:
                scored.append((score, i))
        scored.sort(reverse=True)
        if scored and (len(scored) == 1 or scored[0][0] > scored[1][0]):
            return scored[0][1]
        return None

    # --- поменять местами ---
    if re.search(r"местами|поменя\w+\s+(?:шаг|этап)\w*\s+\d+\s*(?:и|с)\s*\d+", low) and len(refs) == 2:
        a, b = refs
        swap = lambda k: b if k == a else (a if k == b else k)  # noqa: E731
        bodies = [s["body"] for s in steps]
        bodies[a - 1], bodies[b - 1] = bodies[b - 1], bodies[a - 1]
        steps = [{"num": i + 1, "body": _remap(body, swap)} for i, body in enumerate(bodies)]
        changes.append(f"шаги {a} и {b} поменяны местами")
        return EditResult(_join_steps(header, steps), changes), ""

    # --- параллельность ---
    if re.search(r"параллельн|одновременн", low) and len(refs) >= 2:
        make_seq = re.search(r"(?:не\s+параллельн|последовательн)", low) is not None
        first = min(refs)
        for k in sorted(refs):
            body = steps[k - 1]["body"]
            has = bool(re.match(r"^(?:параллельно|одновременно)\b", body, re.I))
            if make_seq and has:
                steps[k - 1]["body"] = re.sub(r"^(?:параллельно|одновременно)\s*[:,—–-]?\s*", "", body, flags=re.I)
                changes.append(f"шаг {k} теперь выполняется последовательно")
            elif not make_seq and k != first and not has:
                steps[k - 1]["body"] = "Параллельно: " + body
                changes.append(f"шаг {k} выполняется параллельно шагу {k - 1}")
        if not changes:
            return None, "Эти шаги уже в нужном режиме — изменений не потребовалось."
        return EditResult(_join_steps(header, steps), changes), ""

    # --- целевой срок процесса ---
    if re.search(r"целев\w+\s+срок|\bsla\b", low) and not refs:
        d = _DUR_RE.search(message)
        if d:
            line = f"Целевой срок процесса: {d.group(0).strip(' ()')}"
            header = [h for h in header if not re.match(r"^(?:целев\w+\s+(?:срок|sla)|sla)", h, re.I)] + [line]
            changes.append(f"целевой срок: {d.group(0).strip(' ()')}")
            return EditResult(_join_steps(header, steps), changes), ""

    # --- удаление ---
    if re.search(r"\b(?:удал\w+|убер\w+|убрать|исключ\w+)\b", low):
        targets = refs or ([by_phrase(message)] if by_phrase(message) else [])
        targets = [t for t in targets if t]
        if not targets:
            return None, "Не понял, какой шаг удалить: укажите номер («Удали шаг 5») или название шага."
        if len(steps) - len(targets) < 2:
            return None, "Нельзя удалить так много шагов — в процессе должно остаться минимум два."
        removed = set(targets)
        new_index: Dict[int, int] = {}
        counter = 0
        for old in range(1, n + 1):
            if old not in removed:
                counter += 1
                new_index[old] = counter

        def fn(k: int) -> int:
            if k in new_index:
                return new_index[k]
            # ссылка на удалённый шаг: ближайший сохранившийся по направлению «вперёд»
            nxt = next((new_index[j] for j in range(k + 1, n + 1) if j in new_index), None)
            prv = next((new_index[j] for j in range(k - 1, 0, -1) if j in new_index), None)
            return nxt if nxt is not None else (prv if prv is not None else 1)

        kept = [{"num": 0, "body": _remap(s["body"], fn)} for i, s in enumerate(steps, 1) if i not in removed]
        for t in sorted(removed):
            changes.append(f"удалён шаг {t}: «{_step_title(steps[t - 1]['body'])}»")
        return EditResult(_join_steps(header, kept), changes), ""

    # --- срок шага ---
    if refs and _DUR_RE.search(re.sub(r"(?:шаг\w*|этап\w*|пункт\w*)\s*\d+", "", message, flags=re.I)) and re.search(
            r"срок|сократ|увелич|измени|постав|установи|до\s+\d|на\s+\d", low):
        k = refs[0]
        d = _DUR_RE.search(re.sub(r"(?:шаг\w*|этап\w*|пункт\w*)\s*\d+", "", message, flags=re.I))
        new_dur = d.group(0).strip(" ()")
        body = steps[k - 1]["body"]
        existing = _DUR_RE.search(body)
        if existing:
            body = body[: existing.start()] + f"({new_dur})" + body[existing.end():]
        else:
            m = re.search(r"\.\s*(?:Если\b|$)", body)
            cut = m.start() if m else len(body)
            body = f"{body[:cut].rstrip()} ({new_dur}){body[cut:]}"
        steps[k - 1]["body"] = re.sub(r"\s{2,}", " ", body)
        changes.append(f"срок шага {k}: {new_dur}")
        return EditResult(_join_steps(header, steps), changes), ""

    # --- передать шаг другой роли ---
    if refs and re.search(r"передай|назначь|поручи|перенес\w+\s+(?:на|к)", low):
        role = next((r for r in [_find_role(message)[0]] if r), None)
        if role:
            k = refs[0]
            body = steps[k - 1]["body"]
            old_role, rs, re_ = _find_role(body)
            if old_role and rs == 0:
                rest = body[re_:].lstrip()
                body = f"{role} {rest}"
            else:
                body = f"{role} {body[:1].lower() + body[1:]}"
            steps[k - 1]["body"] = body
            changes.append(f"шаг {k} передан роли «{role}»")
            return EditResult(_join_steps(header, steps), changes), ""

    # --- добавление ---
    if re.search(r"\b(?:добав\w+|вставь\w*|вставить|включи\w*|дополни\w*)\b", low):
        m = re.search(r"\b(?:добав\w+|вставь\w*|вставить|включи\w*|дополни\w*)\b\s*(.*)$", message, re.I | re.S)
        content = m.group(1) if m else ""
        pos_after = re.search(r"\bпосле\s+(?:(?:шага|этапа|пункта|п\.)\s*)?(\d+)(?:\s*-?\s*(?:го|ого)\s*(?:шага|этапа)?)?", content, re.I)
        pos_before = re.search(r"\bперед\s+(?:(?:шагом|этапом|пунктом|п\.)\s*)?(\d+)", content, re.I)
        in_start = re.search(r"\bв\s+начал\w+", content, re.I)
        in_end = re.search(r"\bв\s+конец\b|\bв\s+конце\b", content, re.I)
        anchor = n
        if pos_after:
            anchor = int(pos_after.group(1))
        elif pos_before:
            anchor = int(pos_before.group(1)) - 1
        elif in_start:
            anchor = 0
        for rx in (pos_after, pos_before, in_start, in_end):
            if rx:
                content = content.replace(rx.group(0), " ")
        content = content.strip(" ,.;:")
        if not content:
            return None, "Что именно добавить? Например: «Добавь согласование с экологами после шага 3»."
        if not 0 <= anchor <= n:
            return None, f"В процессе {n} шагов — шага {anchor} нет."
        anchor_role = steps[max(anchor, 1) - 1]["body"]
        anchor_role = _find_role(anchor_role)[0] or "Исполнитель"
        new_body = _build_new_step(content, anchor_role, _default_duration(steps))
        pos = anchor + 1  # номер нового шага

        def fn_shift(k: int) -> int:
            return k + 1 if k >= pos else k

        shifted = []
        for i, s in enumerate(steps, 1):
            if i == anchor:  # ветка «да» якорного шага теперь ведёт на новый шаг, а не мимо него
                body = _REF_STRICT.sub(
                    lambda mm: f"{mm.group(1)}{pos if int(mm.group(2)) == pos else fn_shift(int(mm.group(2)))}", s["body"])
            else:
                body = _remap(s["body"], fn_shift)
            shifted.append({"num": 0, "body": body})
        shifted.insert(pos - 1, {"num": pos, "body": new_body})
        where = "в начало" if anchor == 0 else (f"после шага {anchor}" if anchor < n else "в конец")
        changes.append(f"добавлен шаг {pos} ({where}): «{_step_title(new_body)}»")
        return EditResult(_join_steps(header, shifted), changes), ""

    return None, ""


_EDIT_HELP = (
    "Не смог однозначно понять правку. Поддерживаются команды, например:\n"
    "- «Добавь согласование с экологами после шага 3»\n"
    "- «Удали шаг 5»\n"
    "- «Сделай шаги 4 и 5 параллельными»\n"
    "- «Измени срок шага 3 на 2 часа»\n"
    "- «Передай шаг 4 службе безопасности»\n"
    "- «Поменяй шаги 6 и 7 местами»\n"
    "- «Установи целевой срок 10 часов»"
)


def _llm_edit(message: str, text: str, trace: List[str]) -> Optional[Tuple[str, str]]:
    messages = [
        {"role": "system", "content": EDIT_SYSTEM},
        {"role": "user", "content": f"ТЕКУЩИЙ РЕГЛАМЕНТ:\n{text.strip()}\n\nКОМАНДА ПОЛЬЗОВАТЕЛЯ: {message.strip()}"},
    ]
    got = _chat_llm(messages, trace)
    if not got:
        return None
    label, raw = got
    new_text = _strip_markdown(raw).strip()
    _, old_steps = _split_steps(normalize_regulation(text))
    _, new_steps = _split_steps(normalize_regulation(new_text))
    if len(new_steps) < 2 or abs(len(new_steps) - len(old_steps)) > 6 or new_text.strip() == text.strip():
        trace.append(f"{label}: правка отклонена (некорректный регламент в ответе)")
        return None
    return label, new_text + ("\n" if not new_text.endswith("\n") else "")


def _rebuild_process(
    text: str, engine: str, trace: List[str], use_llm: bool = False
) -> Tuple[str, Dict[str, Any], str]:
    """Перестроение диаграммы по обновлённому регламенту.

    Если правка идёт из LLM-сессии — generate_bpmn_from_text(..., use_llm=True).
    Эмулятор только при явном сбое или таймауте.
    """
    started = time.time()
    norm = normalize_regulation(text)
    if use_llm:
        xml, audit, err = generate_bpmn_from_text(norm, use_llm=True)
        if xml and not err:
            gen = audit.setdefault("generation", {})
            gen["rebuild"] = engine
            if trace:
                gen["trace"] = list(trace) + list(gen.get("trace") or [])
            return xml, audit, ""
        (trace or []).append(
            f"LLM-перестроение недоступно ({err or 'пустой XML'}) — fallback на эмулятор"
        )
    try:
        header = parse_regulation(norm)
    except Exception:  # noqa: BLE001
        header = ParsedRegulation(title="Бизнес-процесс по регламенту", sla_hours=None)
    code, info = emulate_generation(norm)
    xml, audit, err = execute_generated_code(code, header.title, header.sla_hours, regulation_text=norm)
    if err:
        return "", {}, err
    audit["artifacts"], audit["it_systems"] = header.artifacts, header.it_systems
    audit["generation"] = {
        "engine": engine, "fallback": bool(use_llm), "attempts": 1, "trace": trace, "rejected": [],
        "code": code, "elapsed_s": round(time.time() - started, 2), "parsed": info,
    }
    return xml, audit, ""


def _audit_delta(old: Dict[str, Any], new: Dict[str, Any]) -> str:
    try:
        o, n = old.get("sla") or {}, new.get("sla") or {}
        lines = []
        facts_o = process_facts(old)
        facts_n = process_facts(new)
        if o and n:
            lines.append(
                f"- срок с возвратами: {_fh(facts_o['with_rework_hours'])} → **{_fh(facts_n['with_rework_hours'])}**"
            )
            lines.append(
                f"- критический путь: {_fh(float(o['critical_path_hours']))} → **{_fh(float(n['critical_path_hours']))}**"
            )
            lines.append(
                f"- циклы возврата: {facts_o['rework_loops_n']} → **{facts_n['rework_loops_n']}**"
            )
            if bool(o.get("breach")) != bool(n.get("breach")):
                lines.append(f"- срыв SLA: {'да' if o.get('breach') else 'нет'} → **{'да' if n.get('breach') else 'нет'}**")
        ob, nb = old.get("bus_factor") or {}, new.get("bus_factor") or {}
        if ob and nb:
            lines.append(f"- bus-factor: {float(ob['max_share']):.0%} → **{float(nb['max_share']):.0%}** ({nb['top_role']})")
        os_, ns_ = (old.get("stats") or {}), (new.get("stats") or {})
        if os_ and ns_:
            lines.append(f"- узлов в диаграмме: {os_.get('nodes')} → **{ns_.get('nodes')}**")
        return "\n".join(lines)
    except Exception:  # noqa: BLE001
        return ""


# --------------------------- режим 3: реверс-генерация (инструкция по схеме) --------------------------- #
def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def parse_bpmn_structure(xml: str) -> Dict[str, Any]:
    """BPMN XML → узлы, потоки, дорожки и порядок (по координатам DI): основа реверс-генерации."""
    root = _ET.fromstring(xml.encode("utf-8") if isinstance(xml, str) else xml)
    nodes: Dict[str, Dict[str, Any]] = {}
    flows: List[Dict[str, str]] = []

    def walk(el: "_ET.Element", parent_sub: Optional[str]) -> None:
        for child in el:
            tag = _local(child.tag)
            cid = child.get("id", "")
            if tag in _TASK_TAGS or tag in _GATE_TAGS or tag in ("startEvent", "endEvent", "subProcess"):
                nodes[cid] = {"id": cid, "type": tag, "name": child.get("name", ""), "sub": parent_sub}
                if tag == "subProcess":
                    walk(child, cid)
            elif tag == "sequenceFlow":
                flows.append({"id": cid, "src": child.get("sourceRef", ""), "dst": child.get("targetRef", ""), "name": child.get("name", "")})

    for proc in root.iter(f"{{{_NS_BPMN}}}process"):
        walk(proc, None)
    lane_of: Dict[str, str] = {}
    for lane in root.iter(f"{{{_NS_BPMN}}}lane"):
        for ref in lane.findall(f"{{{_NS_BPMN}}}flowNodeRef"):
            if ref.text:
                lane_of[ref.text.strip()] = lane.get("name", "")
    pos: Dict[str, Tuple[float, float]] = {}
    for shape in root.iter(f"{{{_NS_DI}}}BPMNShape"):
        b = shape.find(f"{{{_NS_DC}}}Bounds")
        if b is not None:
            pos[shape.get("bpmnElement", "")] = (float(b.get("x", 0)), float(b.get("y", 0)))
    for nid, node in nodes.items():
        owner = nid
        while owner in nodes and owner not in lane_of and nodes[owner]["sub"]:
            owner = nodes[owner]["sub"]  # внутренние узлы подпроцесса принадлежат дорожке подпроцесса
        node["lane"] = lane_of.get(owner, "")
        node["x"], node["y"] = pos.get(nid, (0.0, 0.0))
    return {"nodes": nodes, "flows": flows, "lanes": list(dict.fromkeys(lane_of.values()))}


def _hours_phrase(hours: Optional[float]) -> str:
    if hours is None:
        return ""
    minutes = hours * 60.0
    if hours < 1 and abs(minutes - round(minutes)) < 0.05:
        return f" ({int(round(minutes))} минут)"
    if abs(hours - round(hours)) < 0.05:
        return f" ({int(round(hours))} ч)"
    return f" ({hours:.1f} ч)".replace(".0 ч", " ч")


def render_result_regulation(source_text: str, xml: str) -> str:
    """Текст результата: имена «глагол + объект», подписи веток и границы подпроцесса со схемы.

    Без этих трёх кусков повторная сборка того, что видит пользователь, теряет балл.
    Часы и фразы возврата берутся из исходного регламента, если шаг узнаётся.
    """
    source = (source_text or "").strip()
    if not xml or not source:
        return source
    try:
        struct = parse_bpmn_structure(xml)
    except Exception:  # noqa: BLE001
        return source
    nodes: Dict[str, Dict[str, Any]] = struct["nodes"]
    flows: List[Dict[str, str]] = struct["flows"]
    if not nodes:
        return source
    try:
        parsed = parse_regulation(normalize_regulation(source))
    except Exception:  # noqa: BLE001
        parsed = ParsedRegulation(title="Бизнес-процесс", sla_hours=None)

    outgoing: Dict[str, List[Dict[str, str]]] = defaultdict(list)
    for flow in flows:
        outgoing[flow["src"]].append(flow)

    def _expand(owner: Optional[str]) -> List[Dict[str, Any]]:
        kids = [
            n for n in nodes.values()
            if n.get("sub") == owner and (n["type"] in _TASK_TAGS or n["type"] == "subProcess")
        ]
        kids.sort(key=lambda n: (float(n.get("x") or 0), float(n.get("y") or 0), n.get("id") or ""))
        ordered: List[Dict[str, Any]] = []
        for node in kids:
            if node["type"] == "subProcess":
                ordered.extend(_expand(node["id"]))
            else:
                ordered.append(node)
        return ordered

    tasks = _expand(None)
    if len(tasks) < 2:
        return source

    pool = list(tasks)
    bound: Dict[int, Dict[str, Any]] = {}
    for step in parsed.steps:
        title = (step.title or "").strip().lower()
        if not title:
            continue
        hit = next((t for t in pool if (t.get("name") or "").strip().lower() == title), None)
        if hit is None:
            continue
        bound[step.num] = hit
        pool.remove(hit)
    # Похожее имя чужой срок не получает: неузнанный шаг остаётся без часов.
    next_num = max((s.num for s in parsed.steps), default=0)
    for task in pool:
        next_num += 1
        bound_extra = task.setdefault("_out_num", next_num)
        _ = bound_extra
    num_by_task = {task["id"]: num for num, task in bound.items()}
    num_by_task.update({task["id"]: int(task["_out_num"]) for task in pool})

    def _land(node_id: str, guard: int = 0) -> Optional[str]:
        if guard > 8:
            return None
        node = nodes.get(node_id) or {}
        if node.get("type") in _TASK_TAGS and node_id in num_by_task:
            return node_id
        if node.get("type") == "subProcess":
            inner = _expand(node_id)
            return inner[0]["id"] if inner else None
        if node.get("type") == "endEvent" and not node.get("sub"):
            return "END"
        for flow in outgoing.get(node_id, []):
            found = _land(flow["dst"], guard + 1)
            if found:
                return found
        return None

    def _stage_phrase(task: Optional[Dict[str, Any]], step: Optional[Step]) -> str:
        if task is not None:
            parent = nodes.get(task.get("sub") or "")
            if parent and parent.get("type") == "subProcess" and parent.get("name"):
                return f" (этап «{parent['name']}»)"
        if step is not None and step.stage:
            return f" (этап «{step.stage}»)"
        return ""

    def _gateway_sentence(task: Optional[Dict[str, Any]], step: Optional[Step]) -> str:
        step_num = step.num if step is not None else 0
        back_num = None
        if step is not None and step.decision is not None and step.decision.no_ref:
            if step.decision.no_back or step.decision.no_ref < step.num:
                back_num = step.decision.no_ref
        elif step is not None and step.back_ref:
            back_num = step.back_ref
        gateway = None
        if task is not None:
            for flow in outgoing.get(task["id"], []):
                nxt = nodes.get(flow["dst"]) or {}
                if nxt.get("type") in ("exclusiveGateway", "inclusiveGateway"):
                    gateway = nxt
                    break
        if gateway is not None:
            outs = outgoing.get(gateway["id"], [])
            labeled = [(f, _land(f["dst"])) for f in outs]
            labeled = [(f, dest) for f, dest in labeled if dest]
            if len(labeled) >= 2:
                question = str(gateway.get("name") or "Условие").rstrip("?").strip() or "Условие"

                def _dest(dest: str, label: str, *, returning: bool) -> str:
                    if dest == "END":
                        return "завершить процесс"
                    num = num_by_task.get(dest)
                    quote = f"«{label}» — " if label else ""
                    # Иначе-возврат из исходника держит свой номер, даже если ребро село на другой шаг.
                    if returning and back_num is not None:
                        return f"{quote}вернуть на п.{back_num}"
                    if num is not None and num < step_num:
                        return f"{quote}вернуть на п.{num}"
                    if num is None:
                        return "перейти дальше"
                    return f"{quote}перейти к п.{num}"

                yes_f, yes_to = labeled[0]
                no_f, no_to = labeled[1]
                yes_label = str(yes_f.get("name") or "Да").strip() or "Да"
                no_label = str(no_f.get("name") or "Иначе").strip() or "Иначе"
                return (
                    f". Если {question} — {_dest(yes_to, yes_label, returning=False)}, "
                    f"иначе {_dest(no_to, no_label, returning=back_num is not None)}"
                )
        if step is not None and step.decision is not None and step.decision.no_back and step.decision.no_ref:
            return f". иначе «{step.decision.no_label}» — вернуть на п.{step.decision.no_ref}"
        if step is not None and step.back_ref:
            return f". вернуть на п.{step.back_ref}"
        return ""

    lines: List[str] = []
    title = parsed.title or "Бизнес-процесс"
    lines.append(f"Регламент: {title}")
    if parsed.sla_hours:
        lines.append(f"Целевой срок: {_hours_phrase(parsed.sla_hours).strip(' ()')}")
    # Порядок и номер — из регламента, не из координат картинки.
    for step in sorted(parsed.steps, key=lambda s: s.num):
        task = bound.get(step.num)
        name = (str(task.get("name") or "").strip() if task is not None else "") or (step.title or "Выполнить действие")
        role = (str(task.get("lane") or "").strip() if task is not None else "") or step.role or "Исполнитель"
        parallel = "Параллельно: " if step.parallel else ""
        body = f"{parallel}{role}: {name}{_stage_phrase(task, step)}{_hours_phrase(step.hours)}{_gateway_sentence(task, step)}"
        lines.append(f"{step.num}. {body}.")
    for task in pool:
        name = str(task.get("name") or "").strip() or "Выполнить действие"
        role = str(task.get("lane") or "").strip() or "Исполнитель"
        body = f"{role}: {name}{_stage_phrase(task, None)}{_gateway_sentence(task, None)}"
        lines.append(f"{int(task['_out_num'])}. {body}.")
    rendered = "\n".join(lines).strip()
    return rendered or source


def align_result_quality(
    text: str,
    xml: str,
    audit: Dict[str, Any],
    use_llm: bool = False,
) -> Tuple[str, Dict[str, Any]]:
    """Текст результата и аудит карточки.

    Карточка равна местной повторной сборке текста на экране, даже если она ниже рисунка в памяти.
    Второй вызов Groq отсюда не уходит. Схема To-Be — эта сборка, не рисунок с большим баллом.
    """
    _ = use_llm
    data = dict(audit or {})
    memory = data.get("methodology") if isinstance(data.get("methodology"), dict) else {}
    shown = render_result_regulation(text, xml) if xml else (text or "")
    if not shown.strip():
        shown = text or ""
    same_text = normalize_regulation(shown) == normalize_regulation(text or "")
    if same_text and not use_llm:
        if memory:
            data["methodology"] = dict(memory)
        data["result_text"] = shown
        return shown, data
    rebuilt_xml, rebuilt, rebuild_err = generate_bpmn_from_text(shown, use_llm=False)
    rebuilt_meth = (rebuilt or {}).get("methodology") if isinstance((rebuilt or {}).get("methodology"), dict) else None
    if rebuilt and not rebuild_err and rebuilt_meth:
        data["methodology"] = dict(rebuilt_meth)
        data["result_xml"] = rebuilt_xml or ""
        data["result_audit"] = rebuilt
    data["result_text"] = shown
    return shown, data


def card_matches_paste(text: str, use_llm: bool = False) -> Dict[str, Any]:
    """Балл карточки — повторная сборка текста, который показан как результат."""
    xml, audit, err = generate_bpmn_from_text(text, use_llm=use_llm)
    memory = int(((audit or {}).get("methodology") or {}).get("score") or 0)
    shown, audit = align_result_quality(text, xml or "", audit or {}, use_llm=use_llm)
    score = int(((audit.get("methodology") or {}).get("score") or 0))
    return {
        "shown": shown,
        "memory": memory,
        "card": score,
        "xml": xml,
        "audit": audit,
        "error": err or "",
    }


def _work_neighbors(struct: Dict[str, Any], node_id: str, forward: bool) -> List[Tuple[Dict[str, Any], List[str]]]:
    """Ближайшие рабочие узлы выше/ниже по потоку (сквозь шлюзы и события) + условия на пути."""
    nodes, flows = struct["nodes"], struct["flows"]
    result: List[Tuple[Dict[str, Any], List[str]]] = []
    seen = set()
    stack: List[Tuple[str, List[str]]] = [(node_id, [])]
    while stack:
        cur, labels = stack.pop()
        for f in flows:
            a, b = (f["src"], f["dst"]) if forward else (f["dst"], f["src"])
            if a != cur:
                continue
            nxt = nodes.get(b)
            if not nxt or (b, tuple(labels)) in seen:
                continue
            seen.add((b, tuple(labels)))
            new_labels = labels + ([f["name"]] if f["name"] else [])
            if nxt["type"] in _TASK_TAGS or nxt["type"] == "subProcess":
                result.append((nxt, new_labels))
            elif nxt["type"] == "endEvent":
                if not nxt["sub"]:  # «Завершено» внутри подпроцесса — служебное событие
                    result.append((nxt, new_labels))
            elif nxt["type"] in _GATE_TAGS:
                stack.append((b, new_labels))
    return result


def build_role_instruction(role: str, current_xml: str, current_text: str, audit: Dict[str, Any]) -> str:
    """Пошаговая должностная инструкция для роли по схеме BPMN (детерминированный скелет из диаграммы)."""
    struct = parse_bpmn_structure(current_xml)
    nodes = struct["nodes"]
    ctx = build_process_context(current_text, current_xml, audit)
    text_steps = {s["title"]: s for s in ctx["steps"]}
    mine = sorted(
        (n for n in nodes.values() if n["lane"] == role and (n["type"] in _TASK_TAGS or n["type"] == "subProcess")),
        key=lambda n: (n["x"], n["y"]),
    )
    inner = {n["id"]: sorted((c for c in nodes.values() if c["sub"] == n["id"] and c["type"] in _TASK_TAGS), key=lambda c: (c["x"], c["y"]))
             for n in mine if n["type"] == "subProcess"}
    if not mine:
        return f"В схеме нет задач для роли «{role}»."
    load = {r["role"]: r for r in ctx["lane_load"]}
    crit_names = {c["name"]: c for c in ctx["critical_path"] if c["role"] == role}
    lines = [f"# Должностная инструкция по процессу: {role}", f"**Процесс:** {ctx['title']}", ""]
    lines.append("## 1. Общие положения")
    share = load.get(role)
    tasks_total = sum(len(inner[n["id"]]) if n["type"] == "subProcess" else 1 for n in mine)
    lines.append(f"- Роль «{role}» выполняет {tasks_total} задач процесса" + (f" ({float(share['share']):.0%} от всех шагов)." if share else "."))
    first = mine[0]
    preds = [p for p, lab in _work_neighbors(struct, first["id"], forward=False)]
    lines.append("")
    lines.append("## 2. Входные данные и триггеры")
    if preds:
        lines += [f"- Результат шага «{p['name']}» (роль «{p['lane']}»)" for p in preds if p["type"] != "endEvent"]
    else:
        lines.append("- Процесс начинается с действий этой роли: поступление заявки / сообщения / запроса.")
    lines.append("")
    lines.append("## 3. Порядок действий")
    n = 0

    def step_lines(task: Dict[str, Any], indent: str = "") -> None:
        nonlocal n
        n += 1
        meta = text_steps.get(task["name"])
        extra = []
        if meta and meta["hours"]:
            extra.append(f"срок: {_fh(float(meta['hours']))}")
        elif task["name"] in crit_names:
            extra.append(f"срок: {_fh(float(crit_names[task['name']]['hours']))}")
        if meta and meta["systems"]:
            extra.append("системы: " + ", ".join(meta["systems"]))
        if meta and meta["artifacts"]:
            extra.append("документы: " + ", ".join(meta["artifacts"]))
        lines.append(f"{indent}{n}. **{task['name']}**" + (f" — {'; '.join(extra)}" if extra else "") + ".")
        for nxt, labels in _work_neighbors(struct, task["id"], forward=True):
            if labels:
                target = "завершение процесса" if nxt["type"] == "endEvent" else f"«{nxt['name']}» ({nxt['lane']})"
                lines.append(f"{indent}   - Если «{labels[-1]}» — перейти к {target}.")

    for task in mine:
        if task["type"] == "subProcess":
            lines.append(f"**Этап «{task['name']}»** (выполняется по порядку):")
            for child in inner.get(task["id"], []):
                step_lines(child, "   ")
        else:
            step_lines(task)
    lines.append("")
    lines.append("## 4. Передача результата")
    outs: List[str] = []
    for task in mine:
        ids = [task["id"]] + ([c["id"] for c in inner.get(task["id"], [])] if task["type"] == "subProcess" else [])
        for tid in ids:
            for nxt, labels in _work_neighbors(struct, tid, forward=True):
                if nxt["lane"] != role:
                    outs.append(f"- «{nxt['name']}»" if nxt["type"] == "endEvent" else f"- Роли «{nxt['lane']}»: «{nxt['name']}»")
    lines += list(dict.fromkeys(outs)) or ["- Результат остаётся внутри роли; процесс завершается её действиями."]
    lines.append("")
    lines.append("## 5. Контроль сроков и риски")
    if crit_names:
        lines.append("- Шаги на критическом пути SLA: " + "; ".join(f"«{k}» ({_fh(float(v['hours']))})" for k, v in crit_names.items()))
    loops = [l for l in ctx["rework_loops"] if l.get("lane") == role]
    for l in loops:
        lines.append(f"- Возврат «{l['label']}»: {l['from']} → {l['to']}, стоимость цикла {_fh(float(l['cycle_hours']))} — снижать входным контролем.")
    bus = ctx["bus_factor"]
    if bus and bus.get("top_role") == role and bus.get("status") == "risk":
        lines.append(f"- ⚠️ Bus-factor: на роли {float(bus['max_share']):.0%} шагов — необходим назначенный заместитель.")
    if len(lines) and lines[-1].startswith("## 5"):
        lines.append("- Существенных рисков для роли не выявлено.")
    return "\n".join(lines)


# --------------------------- единая точка входа --------------------------- #
def _normalize_history(history: Any) -> List[Dict[str, str]]:
    out: List[Dict[str, str]] = []
    for item in history or []:
        if isinstance(item, dict) and item.get("content"):
            role = "assistant" if item.get("role") == "assistant" else "user"
            out.append({"role": role, "content": str(item["content"])})
        elif isinstance(item, (tuple, list)) and len(item) == 2:
            out.append({"role": "assistant" if item[0] in ("assistant", "ai") else "user", "content": str(item[1])})
    return out


def _source_note(label: Optional[str], trace: List[str]) -> str:
    if label:
        return f"\n\n<sub>Ответ: {label}</sub>"
    reason = f" ({'; '.join(trace[:2])})" if trace else ""
    return f"\n\n<sub>Ответ: локальный анализ процесса — облачная модель недоступна{reason}</sub>"


def assistant_chat(
    message: str,
    history: List[Dict[str, str]],
    current_xml: str,
    current_audit: Dict[str, Any],
    current_text: str,
    use_llm: bool = True,
    tobe_delta: Optional[dict] = None,
) -> Tuple[str, Optional[str], Optional[str], Optional[Dict[str, Any]]]:
    """Диалог с AI-ассистентом над активным процессом.

    Возвращает (reply_text, updated_text, updated_xml, updated_audit). Поля updated_* заполнены, только если
    ассистент изменил процесс: новый текст регламента, новая BPMN-диаграмма и её аудит.
    Исключения наружу не выходят: при любом сбое возвращается понятный ответ, процесс остаётся без изменений.
    """
    try:
        message = (message or "").strip()
        if not message:
            return "Напишите вопрос о процессе или команду на изменение — например, «Добавь согласование с экологами после шага 3».", None, None, None
        audit = current_audit or {}
        ctx = build_process_context(current_text, current_xml, audit, tobe_delta=tobe_delta)
        if not ctx["steps"] and not current_xml:
            return "Активного процесса пока нет: выберите регламент и нажмите «Сгенерировать BPMN 2.0».", None, None, None
        intent = classify_intent(message)
        if intent == "instruction" and not re.search(r"инструкц|памятк|должностн", message, re.I):
            if _visible_answer(message, ctx):
                intent = "analysis"
        trace: List[str] = []
        hist = _normalize_history(history)[-8:]

        # ---- режим 2: правка на лету ----
        if intent == "edit":
            new_text, label, changes = None, None, []
            result, hint = heuristic_edit(message, current_text)
            if result:
                new_text, changes, label = result.text, result.changes, "детерминированная правка регламента"
            elif use_llm:
                llm = _llm_edit(message, current_text, trace)
                if llm:
                    label, new_text = llm
                    changes = ["правка сформулирована моделью по вашей команде"]
            if not new_text:
                return (hint or _EDIT_HELP) + _source_note(None, trace), None, None, None
            xml, new_audit, err = _rebuild_process(new_text, f"assistant · {label}", trace, use_llm=use_llm)
            if err:
                return f"Правка сформирована, но диаграмму построить не удалось: {err}. Процесс оставлен без изменений.", None, None, None
            delta = _audit_delta(audit, new_audit)
            reply = "✨ **Диаграмма обновлена ассистентом в диалоге.**\n\n" + "\n".join(f"- {c}" for c in changes)
            if delta:
                reply += "\n\n**Влияние на метрики:**\n" + delta
            return _limit_lines(reply, 12), new_text, xml, new_audit

        # ---- режим 3: реверс-генерация инструкции ----
        if intent == "instruction":
            roles = _roles_of(ctx)
            named = _roles_in_message(message, roles)
            role = named[0] if named else (ctx["bus_factor"].get("top_role") if ctx["bus_factor"] else None) or (roles[0] if roles else "")
            if not role:
                return "Не нашёл ролей в схеме: сначала сгенерируйте диаграмму.", None, None, None
            if not current_xml:
                return "Для реверс-генерации нужна диаграмма — сгенерируйте BPMN.", None, None, None
            skeleton = build_role_instruction(role, current_xml, current_text, audit)
            label: Optional[str] = None
            final = skeleton
            if use_llm:
                got = _chat_llm([{"role": "system", "content": POLISH_SYSTEM}, {"role": "user", "content": skeleton}], trace)
                if got and len(got[1]) > 200 and role.lower()[:5] in got[1].lower():
                    label, final = got[0], got[1]
            note = "" if named else f"\n\n_Роль не указана — взята самая загруженная. Другие роли: {', '.join(r for r in roles if r != role)}._"
            return final + note + _source_note(label, trace), None, None, None

        # ---- режим 1: аналитика ----
        if intent == "next_step" and not use_llm:
            return _sidebar_shape(suggest_next_steps(ctx), ctx, intent), None, None, None
        low = message.lower()
        force_local = bool(
            _QUALITY_Q_RE.search(low) or _WHY_SAVED_RE.search(low) or _FACT_OVERRIDE_RE.search(low)
            or _OPT_AUDIT_Q_RE.search(low)
        )
        if use_llm and not force_local:
            system = CHAT_SYSTEM + format_context_for_prompt(ctx)
            got = _chat_llm([{"role": "system", "content": system}, *hist, {"role": "user", "content": message}], trace)
            if got:
                return _sidebar_shape(got[1].strip(), ctx, intent), None, None, None
        body = suggest_next_steps(ctx) if intent == "next_step" else heuristic_analysis(message, ctx)
        return _sidebar_shape(body, ctx, intent), None, None, None
    except Exception as exc:  # noqa: BLE001 — диалог не должен ронять приложение
        return f"Не удалось обработать запрос: {type(exc).__name__}: {exc}. Процесс оставлен без изменений.", None, None, None


# --------------------------- Паспорт процесса (обратная выгрузка) --------------------------- #
def _md_cell(value: Any) -> str:
    text = "—" if value is None or value == "" else str(value)
    return text.replace("|", "\\|").replace("\n", " ").strip() or "—"


def _md_join(items: Any) -> str:
    if not items:
        return "—"
    if isinstance(items, str):
        return _md_cell(items)
    names: List[str] = []
    for it in items:
        if isinstance(it, dict):
            name = str(it.get("name") or "").strip()
            if name:
                names.append(name)
        else:
            s = str(it).strip()
            if s:
                names.append(s)
    return _md_cell(", ".join(dict.fromkeys(names))) if names else "—"


def _step_transitions(step: Step) -> str:
    parts: List[str] = []
    if step.parallel:
        parts.append("параллельно с предыдущим шагом")
    d = step.decision
    if d is None:
        return "; ".join(parts) if parts else "—"
    if d.yes_end:
        parts.append(f"«{d.yes_label}» → завершение")
    elif d.yes_ref:
        parts.append(f"«{d.yes_label}» → п.{d.yes_ref}")
    else:
        parts.append(f"«{d.yes_label}» → далее")
    if d.has_else or d.no_ref or d.no_end or d.no_back:
        if d.no_end:
            parts.append(f"«{d.no_label}» → завершение")
        elif d.no_ref:
            dest = f"п.{d.no_ref}"
            if d.no_back:
                dest += " (возврат)"
            parts.append(f"«{d.no_label}» → {dest}")
        else:
            parts.append(f"«{d.no_label}»")
    return "; ".join(parts) if parts else "—"


def _process_owner(regulation_text: str, parsed: ParsedRegulation, audit: Dict[str, Any]) -> str:
    m = re.search(r"владелец(?:\s+процесса)?\s*[:—–-]\s*(.+)$", regulation_text or "", re.I | re.M)
    if m:
        return m.group(1).strip().rstrip(".") or "—"
    top = str((audit.get("bus_factor") or {}).get("top_role") or "").strip()
    if top:
        return top
    if parsed.roles:
        return parsed.roles[0]
    lanes = audit.get("lane_load") or []
    if lanes:
        return str(lanes[0].get("role") or "—")
    return "—"


def generate_process_passport(xml_str: str, audit_data: dict, regulation_text: str) -> str:
    """Официальный Markdown: паспорт процесса, RACI, операционный регламент, ИТ-ландшафт, риски.

    Все поля аудита читаются через .get() с значениями по умолчанию — KeyError исключён.
    """
    audit = audit_data if isinstance(audit_data, dict) else {}
    text = regulation_text or ""
    try:
        parsed = parse_regulation(normalize_regulation(text)) if text.strip() else ParsedRegulation(
            title="Бизнес-процесс", sla_hours=None
        )
    except Exception:  # noqa: BLE001
        parsed = ParsedRegulation(title="Бизнес-процесс", sla_hours=None)

    sla = audit.get("sla") or {}
    stats = audit.get("stats") or {}
    bus = audit.get("bus_factor") or {}
    load = audit.get("lane_load") or []
    path = audit.get("critical_path") or []
    loops = audit.get("rework_loops") or []
    recs = audit.get("recommendations") or []
    systems = audit.get("it_systems") or parsed.it_systems
    artifacts = audit.get("artifacts") or parsed.artifacts
    threshold = float(bus.get("threshold") or 0.45)
    target = sla.get("target_hours")
    if target is None:
        target = parsed.sla_hours
    owner = _process_owner(text, parsed, audit)
    title = parsed.title or "Бизнес-процесс по регламенту"
    nodes = stats.get("nodes") if stats.get("nodes") is not None else "—"
    subs = stats.get("subprocesses") if stats.get("subprocesses") is not None else "—"
    roles_n = stats.get("lanes") if stats.get("lanes") is not None else (len(parsed.roles) or "—")

    lines: List[str] = [
        "# Паспорт процесса и операционный регламент ПАО «Интер РАО»",
        "",
        "*Документ сформирован автоматически модулем «Архитектор BPMN-диаграмм». "
        "Подлежит согласованию Дирекцией бизнес-архитектуры.*",
        "",
        "## 1. Паспорт процесса",
        "",
        f"| Параметр | Значение |",
        f"| --- | --- |",
        f"| Наименование | {_md_cell(title)} |",
        f"| Целевой SLA | {_md_cell(_fh(float(target)) if target else 'не задан')} |",
        f"| Владелец процесса | {_md_cell(owner)} |",
        f"| Число узлов диаграммы | {_md_cell(nodes)} |",
        f"| Подпроцессов | {_md_cell(subs)} |",
        f"| Ролей (дорожек) | {_md_cell(roles_n)} |",
        f"| Рабочих шагов | {_md_cell(stats.get('work_items') if stats.get('work_items') is not None else len(parsed.steps))} |",
        "",
        "## 2. Матрица ролей и ответственности (RACI)",
        "",
        "| Роль | Количество задач | Доля нагрузки, % | Статус bus-factor |",
        "| --- | ---: | ---: | --- |",
    ]
    if load:
        for item in sorted(load, key=lambda it: -float((it or {}).get("share") or 0)):
            role = (item or {}).get("role") or "—"
            tasks = (item or {}).get("tasks") or 0
            share = float((item or {}).get("share") or 0)
            status = "риск bus-factor" if share > threshold else "норма"
            if role == (bus.get("top_role") or "") and str(bus.get("status") or "") == "risk":
                status = "риск bus-factor (ключевой исполнитель)"
            lines.append(f"| {_md_cell(role)} | {int(tasks)} | {share * 100:.0f} | {status} |")
    elif parsed.roles:
        n = len(parsed.steps) or 1
        counts: Dict[str, int] = {}
        for st in parsed.steps:
            counts[st.role] = counts.get(st.role, 0) + 1
        for role, cnt in sorted(counts.items(), key=lambda kv: -kv[1]):
            share = cnt / n
            status = "риск bus-factor" if share > threshold else "норма"
            lines.append(f"| {_md_cell(role)} | {cnt} | {share * 100:.0f} | {status} |")
    else:
        lines.append("| — | 0 | 0 | данных нет |")
    lines += [
        "",
        f"_Порог bus-factor: {threshold:.0%}. Роль выше порога — процесс зависит от одного подразделения._",
        "",
        "## 3. Пошаговый операционный регламент",
        "",
        "| № | Роль-исполнитель | Наименование действия | Нормативный срок | Входные документы / артефакты | Используемые ИТ-системы | Условия переходов |",
        "| ---: | --- | --- | --- | --- | --- | --- |",
    ]
    if parsed.steps:
        for st in parsed.steps:
            dur = _fh(float(st.hours)) if st.hours else "—"
            name = st.title or "(шлюз решения)"
            if st.stage:
                name = f"{name} (этап «{st.stage}»)"
            lines.append(
                f"| {st.num} | {_md_cell(st.role)} | {_md_cell(name)} | {_md_cell(dur)} | "
                f"{_md_join(st.artifacts)} | {_md_join(st.systems)} | {_md_cell(_step_transitions(st))} |"
            )
    elif xml_str:
        try:
            struct = parse_bpmn_structure(xml_str)
            work = sorted(
                (n for n in (struct.get("nodes") or {}).values()
                 if n.get("type") in _TASK_TAGS or n.get("type") == "subProcess"),
                key=lambda n: (float(n.get("x") or 0), float(n.get("y") or 0)),
            )
            if work:
                for i, n in enumerate(work, 1):
                    lines.append(
                        f"| {i} | {_md_cell(n.get('lane'))} | {_md_cell(n.get('name'))} | — | — | — | — |"
                    )
            else:
                lines.append("| — | — | Шаги не распознаны | — | — | — | — |")
        except Exception:  # noqa: BLE001
            lines.append("| — | — | Шаги не распознаны | — | — | — | — |")
    else:
        lines.append("| — | — | Шаги не распознаны | — | — | — | — |")

    lines += [
        "",
        "## 4. ИТ-ландшафт и документооборот",
        "",
        "### 4.1. Реестр ИТ-систем",
        "",
        "| Система | Упоминаний | Шаги | Роли |",
        "| --- | ---: | --- | --- |",
    ]
    if systems:
        for it in systems:
            item = it if isinstance(it, dict) else {"name": it}
            steps = item.get("steps") or []
            roles = item.get("roles") or []
            lines.append(
                f"| {_md_cell(item.get('name'))} | {int(item.get('mentions') or 0)} | "
                f"{_md_cell(', '.join(str(s) for s in steps) if steps else '—')} | {_md_join(roles)} |"
            )
    else:
        lines.append("| — | 0 | — | системы в регламенте не обнаружены |")

    lines += [
        "",
        "### 4.2. Реестр документов и артефактов",
        "",
        "| Документ | Упоминаний | Шаги | Роли |",
        "| --- | ---: | --- | --- |",
    ]
    if artifacts:
        for it in artifacts:
            item = it if isinstance(it, dict) else {"name": it}
            steps = item.get("steps") or []
            roles = item.get("roles") or []
            lines.append(
                f"| {_md_cell(item.get('name'))} | {int(item.get('mentions') or 0)} | "
                f"{_md_cell(', '.join(str(s) for s in steps) if steps else '—')} | {_md_join(roles)} |"
            )
    else:
        lines.append("| — | 0 | — | документы в регламенте не обнаружены |")

    crit_h = sla.get("critical_path_hours")
    rework_h = sla.get("rework_hours")
    with_rw = sla.get("with_rework_hours")
    breach = sla.get("breach")
    path_txt = " → ".join(
        f"«{(p or {}).get('name') or '—'}» ({_fh(float((p or {}).get('hours') or 0))})"
        for p in path
        if (p or {}).get("name")
    ) or "не рассчитан"
    lines += [
        "",
        "## 5. Карта рисков и рекомендации",
        "",
        "### 5.1. Критический путь (алгоритм Беллмана — Форда)",
        "",
        f"- Длительность критического пути: **{_md_cell(_fh(float(crit_h)) if crit_h is not None else '—')}**",
        f"- Целевой SLA: **{_md_cell(_fh(float(target)) if target else 'не задан')}**",
        f"- Срыв SLA: **{'да' if breach else 'нет'}**",
        f"- Маршрут: {path_txt}",
        "",
        "### 5.2. Стоимость циклов возврата",
        "",
        f"- Худший одиночный возврат: **{_md_cell(_fh(float(rework_h)) if rework_h is not None else '0 мин')}**",
        f"- Срок с учётом худшего возврата: **{_md_cell(_fh(float(with_rw)) if with_rw is not None else '—')}**",
        f"- Сумма всех циклов (по одному разу): **{_md_cell(_fh(float(sla.get('rework_total_hours') or 0)))}**",
        "",
    ]
    if loops:
        lines += [
            "| Цикл | Откуда | Куда | Стоимость, ч |",
            "| --- | --- | --- | ---: |",
        ]
        for loop in loops:
            item = loop or {}
            hours = item.get("cycle_hours")
            cost = f"{float(hours):.1f}" if hours is not None else "—"
            lines.append(
                f"| {_md_cell(item.get('label'))} | {_md_cell(item.get('from'))} | "
                f"{_md_cell(item.get('to'))} | {cost} |"
            )
        lines.append("")
    else:
        lines.append("_Циклы возврата не обнаружены._")
        lines.append("")

    lines += ["### 5.3. Меры по оптимизации", ""]
    if recs:
        for rec in recs:
            lines.append(f"- {_md_cell(rec)}")
    else:
        lines.append("- Существенных узких мест не выявлено: процесс сбалансирован по ролям.")
    if xml_str:
        lines += ["", f"_Исходная BPMN-модель: {len(xml_str)} символов XML._"]
    lines.append("")
    return "\n".join(lines)


# --------------------------- инспектор задачи и каталог кликов по холсту --------------------------- #
_KIND_RU = {
    "userTask": "Пользовательская задача",
    "scriptTask": "Скриптовая / сервисная задача",
    "task": "Задача",
    "exclusiveGateway": "Исключающий шлюз (XOR)",
    "parallelGateway": "Параллельный шлюз (AND)",
    "inclusiveGateway": "Включающий шлюз (OR)",
    "startEvent": "Стартовое событие",
    "endEvent": "Завершающее событие",
    "subProcess": "Подпроцесс",
}

_ROLE_QUAL = (
    ("диспетчер", "IV–V группа по электробезопасности, право ведения оперативных переговоров и переключений."),
    ("начальник смены", "V группа по электробезопасности, ответственный руководитель работ в электроустановке."),
    ("служб безопасн", "V группа, контроль нарядов-допусков, проверка состава бригады и удостоверений."),
    ("ремонтн", "III–IV группа по электробезопасности, допуск к работам в электроустановках до и выше 1000 В."),
    ("эколог", "Профильная подготовка по охране окружающей среды; допуск к площадке — по наряду."),
    ("охран труд", "Специалист по ОТ, группа не ниже III, контроль СИЗ и инструктажа."),
    ("закуп", "Квалификация закупщика / договорной работы; ЭЦП для электронной площадки."),
    ("бухгалт", "Квалификация бухгалтера энергосбыта / генерирующей компании, доступ к 1С / ERP."),
)

INSPECT_SYSTEM = """Ты — старший мастер оперативного персонала ПАО «Интер РАО».
По задаче BPMN составь операционную карточку исполнителя. Ответь ТОЛЬКО JSON без markdown:
{
  "role_requirements": "квалификация, разряд, группа допуска по электробезопасности (II–V)",
  "procedure_steps": ["шаг 1", "шаг 2", "шаг 3", "шаг 4"],
  "safety_and_tools": "приборы, СИЗ, заземления, ИТ-системы",
  "input_trigger": "основание начать работы и входные документы",
  "output_artifact": "результат: акт, журнал, статус"
}
Только факты из контекста процесса; 3–4 конкретных технических действия; по-русски; ничего не выдумывай сверх роли и названия шага.
СИЗ выдавай только для полевых ролей (бригада, монтёр, обходчик, электромонтёр).
Для офисных/пультовых (диспетчер, юрист, начальник смены, бухгалтер, закупки, СБ) пиши: «СИЗ: Не требуются (офисный/пультовой режим)»."""


_FIELD_ROLE_RE = re.compile(
    r"бригад|монт[её]р|обходчик|электромонт|линейщ|производитель работ|ремонтн\w+\s+персонал",
    re.I,
)
_OFFICE_ROLE_RE = re.compile(
    r"диспетчер|юрист|юридическ|начальник смены|бухгалтер|закуп|тендер|"
    r"служб\w+\s+безопасност|\bсб\b|техническ\w+\s+департамент|"
    r"центр обслуживания|заявител|секретар|экономист",
    re.I,
)
_PPE_OFFICE = "СИЗ: Не требуются (офисный/пультовой режим)"
_PPE_FIELD = (
    "СИЗ: каска, термостойкий комбинезон, диэлектрические перчатки и боты (по наряду). "
    "При работах в электроустановке — переносные заземления и запирание коммутационных аппаратов."
)


def _ppe_for_role(role: str, sys_txt: str = "") -> str:
    """СИЗ только для полевых ролей; офис/пульт — явный отказ от касок."""
    low = (role or "").lower()
    extra = f" Системы: {sys_txt}." if sys_txt else ""
    if _OFFICE_ROLE_RE.search(low) or not _FIELD_ROLE_RE.search(low):
        return _PPE_OFFICE + extra
    return _PPE_FIELD + extra


def _role_requirements(role: str) -> str:
    low = (role or "").lower()
    for stem, text in _ROLE_QUAL:
        if stem in low:
            return text
    return (
        f"Исполнитель роли «{role or 'специалист'}»: профильная подготовка, "
        "группа по электробезопасности не ниже III, допуск к работам по действующему наряду."
    )


def _match_step(task_name: str, ctx: Dict[str, Any]) -> Dict[str, Any]:
    name = (task_name or "").lower()
    best: Dict[str, Any] = {}
    score = 0
    for st in ctx.get("steps") or []:
        title = str(st.get("title") or "").lower()
        s = 0
        if title and (title in name or name in title):
            s = 3
        else:
            stems = {w[:6] for w in re.findall(r"[А-Яа-яЁё]{5,}", name)}
            s = sum(1 for w in stems if w in title)
        if s > score:
            score, best = s, st
    return best


def _heuristic_inspect(task_name: str, task_role: str, ctx: Dict[str, Any]) -> Dict[str, Any]:
    name = (task_name or "Выполнить действие").strip()
    role = (task_role or "Исполнитель").strip()
    low = name.lower()
    step = _match_step(name, ctx)
    systems = list(step.get("systems") or [])
    artifacts = list(step.get("artifacts") or [])
    if not systems:
        systems = [i.get("name") for i in (ctx.get("it_systems") or []) if isinstance(i, dict) and i.get("name")]
    if not artifacts:
        artifacts = [i.get("name") for i in (ctx.get("artifacts") or []) if isinstance(i, dict) and i.get("name")]
    sys_txt = ", ".join(str(s) for s in systems[:4]) or "оперативный журнал, ИС/АС предприятия"
    doc_txt = ", ".join(str(a) for a in artifacts[:4]) or "задание смены / наряд-допуск"

    steps = [f"Получить задание и подтвердить полномочия роли «{role}»."]
    action = name[:1].lower() + name[1:] if name else "выполнить операцию по регламенту"
    steps.append(f"Выполнить: {action}.")
    if any(k in low for k in ("наряд", "допуск", "безопасн", "инструктаж")):
        steps.append("Проверить наряд-допуск, состав бригады, удостоверения и наличие СИЗ.")
        steps.append("Допустить персонал к работе или вернуть пакет документов на устранение замечаний.")
    elif any(k in low for k in ("журнал", "фиксир", "регистр", "заявк")):
        steps.append(f"Внести запись в {sys_txt}.")
        steps.append("Сверить статус заявки со смежной ролью и закрыть шаг в системе.")
    elif any(k in low for k in ("ремонт", "переключ", "вывест", "восстанов", "схем")):
        steps.append("Выполнить оперативные переключения / работы по бланку, установить переносные заземления.")
        steps.append("Подтвердить восстановление схемы и снять наряд после осмотра.")
    elif any(k in low for k in ("осмотр", "диагност", "измер", "тепловиз", "изоляц")):
        steps.append("Провести измерения приборами (тепловизор, мегаомметр) и зафиксировать дефекты.")
        steps.append("Передать дефектную ведомость следующей роли по схеме.")
    else:
        steps.append("Сверить результат с нормами ПТЭ и локальными инструкциями предприятия.")
        steps.append("Передать результат следующей роли и зафиксировать статус в учётной системе.")

    if "script" in low or "автомат" in low:
        safety = f"Автоматизированная операция. Системы: {sys_txt}. Контроль журнала событий, без работ в электроустановке."
    else:
        safety = _ppe_for_role(role, sys_txt)
    trigger = (
        f"Основание: поступление шага «{name}» роли «{role}» по схеме процесса"
        + (f"; входные документы: {doc_txt}." if artifacts or doc_txt else ".")
    )
    output = (
        f"Результат шага «{name}»: статус выполнен; "
        + (f"оформлены {doc_txt}; " if artifacts else "запись в оперативном журнале; ")
        + "передача следующей роли по BPMN."
    )
    return {
        "role_requirements": _role_requirements(role),
        "procedure_steps": steps[:4],
        "safety_and_tools": safety,
        "input_trigger": trigger,
        "output_artifact": output,
        "source": "heuristic",
    }


def inspect_task_details(task_name: str, task_role: str, process_context: dict) -> dict:
    """Развёрнутая операционная инструкция исполнителя по шагу BPMN.

    Облачная модель (Groq/OpenAI) → эвристики процесса. Всегда возвращает словарь с ключами
    role_requirements, procedure_steps, safety_and_tools, input_trigger, output_artifact.
    """
    ctx = process_context if isinstance(process_context, dict) else {}
    fallback = _heuristic_inspect(task_name or "", task_role or "", ctx)
    if os.getenv("BPMN_AI_MODE", "auto").lower() == "emulator":
        return fallback
    payload = {
        "task": task_name,
        "role": task_role,
        "process": ctx.get("title"),
        "sla": ctx.get("sla"),
        "step_match": _match_step(task_name or "", ctx),
        "it_systems": ctx.get("it_systems"),
        "artifacts": ctx.get("artifacts"),
        "critical_path": [c.get("name") for c in (ctx.get("critical_path") or [])][:12],
    }
    trace: List[str] = []
    got = _chat_llm(
        [
            {"role": "system", "content": INSPECT_SYSTEM},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)[:6000]},
        ],
        trace,
    )
    if not got:
        return fallback
    raw = got[1]
    try:
        m = re.search(r"\{[\s\S]*\}", raw)
        data = json.loads(m.group(0) if m else raw)
    except Exception:  # noqa: BLE001
        return fallback
    steps = data.get("procedure_steps") if isinstance(data, dict) else None
    if isinstance(steps, str):
        steps = [s.strip(" -•") for s in re.split(r"[\n;]+", steps) if s.strip()]
    if not isinstance(steps, list) or len(steps) < 2:
        steps = fallback["procedure_steps"]
    return {
        "role_requirements": str(data.get("role_requirements") or fallback["role_requirements"]),
        "procedure_steps": [str(s).strip() for s in steps[:6] if str(s).strip()],
        "safety_and_tools": _ppe_for_role(task_role or "", ""),
        "input_trigger": str(data.get("input_trigger") or fallback["input_trigger"]),
        "output_artifact": str(data.get("output_artifact") or fallback["output_artifact"]),
        "source": got[0],
    }


def live_node_comment(name: str, role: str, kind: str, ctx: Dict[str, Any]) -> str:
    """Короткий комментарий ИИ для карточки клика: зачем шаг в энергосистеме и где риск задержки."""
    low = (name or "").lower()
    crit = {str(c.get("name") or "") for c in (ctx.get("critical_path") or [])}
    on_crit = name in crit or any(name and name[:18] in (c or "") for c in crit)
    bus = ctx.get("bus_factor") or {}
    top = str(bus.get("top_role") or "")
    loops = ctx.get("rework_loops") or []
    loop_hit = any(name and name[:12] in str(l.get("from", "")) + str(l.get("to", "")) for l in loops)

    if kind in _GATE_TAGS:
        why = "Развилка фиксирует решение, от которого зависит допуск, схема или возврат на доработку."
        risk = "Неподписанная или спорная ветка здесь даёт ложный маршрут и простой бригады."
    elif kind == "startEvent":
        why, risk = "Точка входа аварии или заявки в оперативный контур.", "Поздняя фиксация старта съедает весь запас SLA."
    elif kind == "endEvent":
        why, risk = "Подтверждение, что схема и документы закрыты.", "Незакрытый конец оставляет «висящий» наряд и открытый SLA."
    elif kind == "subProcess":
        why = "Декомпозиция однотипных шагов одной роли — защита от «карты метро»."
        risk = "Если внутри длинная цепочка без контроля, узкое место прячется в подпроцессе."
    elif any(k in low for k in ("наряд", "допуск")):
        why = "Наряд-допуск — юридический и технический барьер перед работой в электроустановке."
        risk = "Замечания СБ и возврат наряда — типичный цикл, который раздувает SLA."
    elif any(k in low for k in ("ремонт", "восстанов", "переключ")):
        why = "Оперативные переключения и ремонт возвращают оборудование в нормальную схему."
        risk = "Шаг на критическом пути: каждый час простоя — недоотпуск и риск каскада."
    elif any(k in low for k in ("журнал", "фиксир")):
        why = "Оперативный журнал — единый источник правды для диспетчера и смены."
        risk = "Ошибка или задержка записи ломает трассировку аварии и расследование."
    else:
        why = f"Шаг роли «{role or 'исполнитель'}» продвигает процесс по регламенту энергосистемы."
        risk = "Задержка на исполнителе с высокой долей нагрузки повышает bus-factor."
    if on_crit:
        risk = "Узел на критическом пути SLA (Беллман — Форд): задержка сдвигает весь срок восстановления."
    elif loop_hit:
        risk = "Узел связан с циклом возврата — повторный проход умножает трудозатраты."
    elif role and role == top and float(bus.get("max_share") or 0) > 0.4:
        risk = f"Роль «{role}» держит {float(bus.get('max_share') or 0):.0%} шагов — простой ключевого исполнителя останавливает процесс."
    return f"{why} {risk}"


def build_diagram_catalog(xml_str: str, audit_data: dict, regulation_text: str) -> Dict[str, Dict[str, Any]]:
    """id BPMN-элемента → карточка для клика на холсте (роль, критический путь, комментарий ИИ)."""
    ctx = build_process_context(regulation_text or "", xml_str or "", audit_data or {})
    catalog: Dict[str, Dict[str, Any]] = {}
    if not xml_str:
        return catalog
    try:
        struct = parse_bpmn_structure(xml_str)
    except Exception:  # noqa: BLE001
        return catalog
    clickable = _TASK_TAGS | _GATE_TAGS | {"startEvent", "endEvent", "subProcess"}
    for nid, node in (struct.get("nodes") or {}).items():
        kind = str(node.get("type") or "")
        if kind not in clickable:
            continue
        name = str(node.get("name") or "").strip() or kind
        role = str(node.get("lane") or "").strip()
        crit = any(
            name == str(c.get("name") or "") or nid == str(c.get("id") or "")
            for c in (ctx.get("critical_path") or [])
        )
        catalog[nid] = {
            "id": nid,
            "name": name,
            "type": _KIND_RU.get(kind, kind),
            "kind": kind,
            "role": role or "—",
            "critical": bool(crit),
            "comment": live_node_comment(name, role, kind, ctx),
            "copilot": _copilot_node_brief(name, role, kind, ctx),
        }
    return catalog


def _copilot_node_brief(name: str, role: str, kind: str, ctx: Dict[str, Any]) -> str:
    """До 6 строк: вывод, блок/роль на схеме, одна причина (критический путь или цикл)."""
    f = ctx.get("facts") or {}
    crit_names = {str(c.get("name") or "") for c in (f.get("critical_path") or ctx.get("critical_path") or [])}
    on_crit = bool(name) and (name in crit_names or any(name[:16] in (c or "") for c in crit_names))
    loops = f.get("rework_loops") or ctx.get("rework_loops") or []
    hit = next(
        (l for l in loops if name and name[:12] in str(l.get("from", "")) + str(l.get("to", ""))),
        None,
    )
    lines: List[str] = []
    if on_crit:
        lines.append(f"**На критическом пути:** «{name}» ({role or '—'}).")
        lines.append(f"Голый КП {_fh(float(f.get('critical_path_hours') or 0))} без этого шага не сходится.")
    elif hit:
        lines.append(f"**В цикле возврата** «{hit.get('label')}»: «{name}» ({role or '—'}).")
        lines.append(f"Повтор добавляет {_fh(float(hit.get('cycle_hours') or 0))} к сроку с возвратами.")
    else:
        lines.append(f"**«{name}»** — {role or 'исполнитель'}, не на узком месте SLA.")
        lines.append(f"Срок с возвратами {_fh(float(f.get('with_rework_hours') or 0))}.")
    if f.get("speedup_via_rework"):
        lines.append(
            f"To-Be снимает возвраты: {_fh(float(f['rw_before']))} → {_fh(float(f['rw_after']))}, "
            f"циклы {f.get('loops_before')} → {f.get('loops_after')}."
        )
    return _limit_lines("\n".join(lines), 6)


def _copilot_tobe(ctx: Dict[str, Any]) -> str:
    return _limit_lines(_facts_time_reply(ctx), 6)


def canvas_copilot_reply(
    message: str,
    xml_str: str = "",
    audit_data: Optional[Dict[str, Any]] = None,
    regulation_text: str = "",
    tobe_delta: Optional[dict] = None,
    selected_id: Optional[str] = None,
    catalog: Optional[Dict[str, Any]] = None,
) -> str:
    """Локальный копайлот холста: не вызывает облако, не меняет XML/регламент."""
    ctx = build_process_context(regulation_text or "", xml_str or "", audit_data or {}, tobe_delta=tobe_delta)
    msg = (message or "").strip()
    if not msg:
        return "Спросите про блок на схеме, SLA или As-Is/To-Be."
    if classify_intent(msg) == "edit":
        cmd = re.sub(r"\s+", " ", msg).strip(" .")
        return f"Команду в сайдбар: «{cmd}»"
    visible = _visible_answer(msg, ctx)
    if visible:
        return _limit_lines(visible, 12)
    cat = catalog or {}
    if selected_id and selected_id in cat:
        packed = cat[selected_id].get("copilot")
        if packed and re.search(r"этот|выбран|блок|почему|критич|цикл|путь|возврат", msg, re.I):
            return packed
    low = msg.lower()
    f = ctx.get("facts") or {}
    if re.search(r"сравни|as-is|as is|to-be|tobe|до и после|ускор|экономи", low):
        return _copilot_tobe(ctx)
    if re.search(r"sla|срок|срыв|критич|длительн", low):
        return _limit_lines(
            f"**Срок с возвратами {_fh(float(f.get('with_rework_hours') or 0))}**, "
            f"голый КП {_fh(float(f.get('critical_path_hours') or 0))}.\n"
            f"Циклов {int(f.get('rework_loops_n') or 0)}.",
            6,
        )
    if re.search(r"цикл|возврат|доработ|rework", low):
        n = int(f.get("rework_loops_n") or 0)
        if f.get("tobe_ready"):
            return _limit_lines(
                f"**Циклы {f.get('loops_before')} → {f.get('loops_after')}.**\n"
                f"Путь с возвратами {_fh(float(f['rw_before']))} → {_fh(float(f['rw_after']))}.",
                6,
            )
        return _limit_lines(f"**Циклов возврата: {n}.** " + (f.get("rework_loops") or [{}])[0].get("label", ""), 6)
    if re.search(r"роль|нагруз|bus|исполнител", low):
        return _limit_lines(
            f"**Bus-factor: «{f.get('bus_role') or '—'}» {float(f.get('bus_share') or 0):.0%} шагов.**\n"
            "На схеме эта дорожка держит процесс.",
            6,
        )
    if selected_id and selected_id in cat:
        return cat[selected_id].get("copilot") or _copilot_tobe(ctx)
    return _copilot_tobe(ctx) if f.get("tobe_ready") else _limit_lines(
        f"**«{ctx.get('title')}».** Срок с возвратами {_fh(float(f.get('with_rework_hours') or 0))}.",
        6,
    )


def build_canvas_copilot(
    xml_str: str,
    audit_data: dict,
    regulation_text: str,
    tobe_delta: Optional[dict] = None,
) -> Dict[str, Any]:
    """Пакет для плавающего ассистента на холсте: чипы и локальные ответы. Облако не вызывается."""
    ctx = build_process_context(regulation_text or "", xml_str or "", audit_data or {}, tobe_delta=tobe_delta)
    title = str(ctx.get("title") or "Бизнес-процесс")
    sla_a = canvas_copilot_reply("В чём причина срыва SLA?", xml_str, audit_data, regulation_text, tobe_delta)
    speed_a = canvas_copilot_reply("Как ускорить процесс?", xml_str, audit_data, regulation_text, tobe_delta)
    roles_a = canvas_copilot_reply("Как оптимизировать нагрузку ролей?", xml_str, audit_data, regulation_text, tobe_delta)
    compare_a = canvas_copilot_reply("Сравни As-Is и To-Be", xml_str, audit_data, regulation_text, tobe_delta)
    loops_a = canvas_copilot_reply("Циклы возврата на доработку", xml_str, audit_data, regulation_text, tobe_delta)
    chips = [
        {"id": "speed", "label": "⚡ Как ускорить?", "q": "Как ускорить процесс?", "a": speed_a},
        {"id": "sla", "label": "🔍 Анализ SLA", "q": "В чём причина срыва SLA?", "a": sla_a},
        {"id": "roles", "label": "👤 Роли и риски", "q": "Как оптимизировать нагрузку ролей?", "a": roles_a},
        {"id": "tobe", "label": "🔀 As-Is → To-Be", "q": "Сравни As-Is и To-Be", "a": compare_a},
    ]
    return {
        "title": title,
        "greeting": (
            f"Копайлот схемы «{title}»: объясняю то, что на холсте. Правки — в сайдбар."
        ),
        "fallback": sla_a,
        "compare": compare_a,
        "loops": loops_a,
        "why": speed_a,
        "quality": roles_a,
        "readability": sla_a,
        "landscape": roles_a,
        "chips": chips,
    }


# --------------------------- RACI, To-Be, официальный DOCX --------------------------- #
_ACCOUNTABLE_RE = re.compile(
    r"начальник|руководител|директор|владелец|главн\w+\s+инженер|начальник смены|зам[\.\s]",
    re.I,
)
_CONSULTED_RE = re.compile(
    r"безопасност|юрист|правов|эколог|эксперт|согласующ|охрана труда|служба охраны|нормативн",
    re.I,
)
_INFORMED_RE = re.compile(r"диспетчер|заявител|заказчик|потребител|инициатор", re.I)
_JOURNAL_RE = re.compile(
    r"фикс\w+|журнал|реестр|регистр\w+|вносит\s+запис|учётн\w+\s+систем|оперативн\w+\s+журнал",
    re.I,
)
_CONTROL_HINT_RE = re.compile(r"входн\w+\s+контрол|комплектност\w+\s+документ", re.I)
_OT_STOP_RE = re.compile(
    r"допуск|наряд[\s-]*допуск|инструктаж|"
    r"проверк\w+\s+отсутств\w+\s+напряжен|"
    r"установ\w+\s+заземлен|налож\w+\s+заземлен|включ\w+\s+заземляющ|"
    r"(?<!зон[ауиеы]\s)отключен|"
    r"разрешен|согласован|утвержден",
    re.I,
)
_PPE_PREPARE_RE = re.compile(
    r"(?:подготов|готов\w*).{0,80}(?:сиз|инструмент|переносн\w+\s+заземлен)",
    re.I,
)
_INSTALL_GROUND_RE = re.compile(
    r"установ\w+\s+заземлен|налож\w+\s+заземлен|включ\w+\s+заземляющ",
    re.I,
)
_TOBE_MINUTE = 1.0 / 60.0
_REPAIR_WORK_RE = re.compile(
    r"аварийн\w+\s+ремонт|выполн\w+\s+.{0,40}ремонт|ремонт\s+оборудован|"
    r"строительно-монтаж|производств\w+\s+работ|выполн\w+\s+работ",
    re.I,
)
_RACI_STOP = {"выполн", "провод", "оформ", "провер", "принят", "переда", "состав", "оценк"}
_CAUSAL_RE = re.compile(
    r"на\s+основан|по\s+результат|после\s+(?:чего|этого|шага)|затем|"
    r"переда[её]т|входн\w+\s+данн|полученн\w+|исходн\w+\s+данн|"
    r"согласованн\w+\s+документ",
    re.I,
)
TOBE_SYSTEM = """Ты — ведущий бизнес-архитектор ПАО «Интер РАО».
Перепиши регламент, сохранив заголовок и целевой SLA. Правила:
1) Параллелизация: независимые шаги РАЗНЫХ ролей начинай с «Параллельно:».
   СТОП-ЛИСТ охраны труда — НЕ ставь «Параллельно:», если шаг — допуск, наряд-допуск, целевой инструктаж,
   отключение, проверка отсутствия напряжения, УСТАНОВКА заземлений (установить/наложить, ножи).
   Подготовка СИЗ, инструмента и переносных заземлений («подготовить … заземления») — НЕ стоп-лист, её можно
   параллелить с оперативными переключениями диспетчера.
   Фактический ремонт / выполнение работ — СТРОГО ПОСЛЕ допуска и инструктажа, никогда параллельно с ними.
2) Zero-Rework: цикл замени эскалацией на ИСКЛЮЧИТЕЛЬНОЙ ветке, не на счастливом пути.
   Входной контроль — параллельно независимому шагу ИЛИ не длиннее 5 минут (не 15 минут на критическом пути).
   Формулировки «вернуть на п.N / на доработку» замени эскалацией руководителю без повторного цикла.
   Не начинай шаг с существительного («Входной контроль…») — только роль + глагол.
3) Автоматизация: фиксацию в журналах пиши БЕЗ двоеточия после системы:
   «Информационная система автоматически регистрирует … (5 минут)».
Верни ТОЛЬКО текст регламента с нумерованными шагами, без комментариев и markdown."""


def _accountable_role(roles: List[str]) -> str:
    for role in roles or []:
        if _ACCOUNTABLE_RE.search(role or ""):
            return role
    return (roles[0] if roles else "Руководитель процесса") or "Руководитель процесса"


def _consulted_roles(roles: List[str]) -> List[str]:
    return [r for r in (roles or []) if _CONSULTED_RE.search(r or "")]


def _informed_roles(roles: List[str]) -> List[str]:
    return [r for r in (roles or []) if _INFORMED_RE.search(r or "")]


def _step_as_row(item: Any, idx: int) -> Dict[str, Any]:
    if isinstance(item, Step):
        return {
            "num": item.num,
            "title": item.title or "(шлюз решения)",
            "role": item.role or "",
            "decision": bool(item.decision),
            "hours": item.hours,
            "artifacts": list(item.artifacts or []),
            "systems": list(item.systems or []),
        }
    if isinstance(item, dict):
        title = item.get("title") or item.get("name") or item.get("action") or ""
        decision = item.get("decision")
        return {
            "num": item.get("num") if item.get("num") is not None else idx,
            "title": title or "(шаг)",
            "role": item.get("role") or item.get("lane") or item.get("executor") or "",
            "decision": bool(decision) if not isinstance(decision, dict) else True,
            "hours": item.get("hours"),
            "artifacts": list(item.get("artifacts") or []),
            "systems": list(item.get("systems") or []),
        }
    return {"num": idx, "title": str(item), "role": "", "decision": False, "hours": None, "artifacts": [], "systems": []}


def generate_raci_matrix(steps: list, roles: list) -> List[Dict[str, Any]]:
    """Матрица RACI: строки — шаги, столбцы — роли, на пересечении буквы R/A/C/I."""
    role_list = [str(r).strip() for r in (roles or []) if str(r).strip()]
    rows_in = [_step_as_row(s, i + 1) for i, s in enumerate(steps or [])]
    for row in rows_in:
        if row["role"] and row["role"] not in role_list:
            role_list.append(row["role"])
    if not role_list:
        role_list = ["Исполнитель"]
    owner = _accountable_role(role_list)
    consulted = _consulted_roles(role_list)
    informed = _informed_roles(role_list)
    matrix: List[Dict[str, Any]] = []
    last_idx = len(rows_in) - 1
    for i, row in enumerate(rows_in):
        executor = row["role"] or role_list[0]
        title = str(row["title"] or "")
        decision = bool(row["decision"]) or bool(re.search(r"соглас|допуск|утверд|экспертиз", title, re.I))
        milestone = decision or i >= last_idx - 1
        accountable = executor if _ACCOUNTABLE_RE.search(executor) else owner
        if decision:
            accountable = executor
        cells: Dict[str, List[str]] = {role: [] for role in role_list}

        def _add(role: str, letter: str) -> None:
            if role not in cells:
                cells[role] = []
            if letter not in cells[role]:
                cells[role].append(letter)

        _add(executor, "R")
        _add(accountable, "A")
        if decision or bool(re.search(r"проверк|допуск|наряд|экспертиз|правов", title, re.I)):
            for role in consulted:
                if role not in (executor, accountable):
                    _add(role, "C")
        if milestone:
            for role in informed:
                if role not in (executor, accountable):
                    _add(role, "I")
        matrix.append(
            {
                "num": row["num"],
                "title": title,
                "executor": executor,
                "assignments": cells,
            }
        )
    return matrix


def _info_or_causal(a: Step, b: Step, body_a: str = "", body_b: str = "") -> bool:
    """Информационная или причинная зависимость: общий артефакт, ссылка на выход, каузальные маркеры."""
    arts_a = {str(x).strip().lower() for x in (a.artifacts or []) if str(x).strip()}
    arts_b = {str(x).strip().lower() for x in (b.artifacts or []) if str(x).strip()}
    if arts_a and arts_b and (arts_a & arts_b):
        return True
    blob_b = (body_b or "").lower()
    for art in arts_a:
        if len(art) >= 4 and art in blob_b:
            return True
    if _CAUSAL_RE.search(body_b or "") or _CAUSAL_RE.search(b.title or ""):
        return True
    title_a = (a.title or "").strip().lower()
    if len(title_a) >= 12 and title_a in blob_b:
        return True
    _ = body_a
    return False


def _independent_steps(a: Step, b: Step, body_a: str = "", body_b: str = "") -> bool:
    """Параллелить можно только разные роли без информационной/причинной связи и без стоп-листа ОТ."""
    if not a.role or not b.role or a.role == b.role:
        return False
    if a.decision or b.decision or b.parallel:
        return False
    if _info_or_causal(a, b, body_a, body_b):
        return False
    stems_a = {w[:6].lower() for w in re.findall(r"[А-Яа-яЁё]{5,}", a.title or "")}
    stems_b = {w[:6].lower() for w in re.findall(r"[А-Яа-яЁё]{5,}", b.title or "")}
    return not ((stems_a & stems_b) - _RACI_STOP)


def _ot_blob(step: Step, body: str = "") -> str:
    return f"{step.title or ''} {body or ''}"


def _is_prepare_ppe_ground(blob: str) -> bool:
    """«Готовит СИЗ, инструмент и переносные заземления» — подготовка, не установка заземлений."""
    return bool(_PPE_PREPARE_RE.search(blob or "")) and not _INSTALL_GROUND_RE.search(blob or "")


def _ot_sensitive(step: Step, body: str = "") -> bool:
    """Стоп-лист ОТ. Существительное «заземления» в шаге «подготовить» стоп-лист не включает."""
    blob = _ot_blob(step, body)
    if _is_prepare_ppe_ground(blob):
        return False
    if _INSTALL_GROUND_RE.search(blob):
        return True
    return bool(_OT_STOP_RE.search(blob))


def _is_repair_work(step: Step, body: str = "") -> bool:
    return bool(_REPAIR_WORK_RE.search(_ot_blob(step, body)))


def _forbid_parallel(a: Step, b: Step, body_a: str = "", body_b: str = "") -> bool:
    """AND запрещён между шагами стоп-листа ОТ; ремонт не параллелен допуску/инструктажу.

    Подготовка СИЗ/инструмента/переносных заземлений может идти параллельно переключениям диспетчера.
    """
    a_ppe = _is_prepare_ppe_ground(_ot_blob(a, body_a))
    b_ppe = _is_prepare_ppe_ground(_ot_blob(b, body_b))
    a_ot, b_ot = _ot_sensitive(a, body_a), _ot_sensitive(b, body_b)
    a_repair, b_repair = _is_repair_work(a, body_a), _is_repair_work(b, body_b)
    a_permit = bool(re.search(r"допуск|наряд|инструктаж", _ot_blob(a, body_a), re.I)) and not a_ppe
    b_permit = bool(re.search(r"допуск|наряд|инструктаж", _ot_blob(b, body_b), re.I)) and not b_ppe
    if (a_ppe and (b_repair or b_permit)) or (b_ppe and (a_repair or a_permit)):
        return True
    if a_ppe or b_ppe:
        return False
    if a_ot and b_ot:
        return True
    if a_ot or b_ot:
        if a_repair or b_repair or a_permit or b_permit:
            return True
    if (a_repair and b_permit) or (b_repair and a_permit):
        return True
    return False


def _automate_journal_body(body: str) -> str:
    """scriptTask: система сама выполняет действие. Заголовок должен начинаться с глагола, не с роли."""
    text = body.strip()
    if _DUR_RE.search(text):
        text = _DUR_RE.sub("(5 минут)", text, count=1)
    else:
        text = text.rstrip(" .") + " (5 минут)."

    def _action_after_role(src: str) -> str:
        role, start, end = _find_role(src)
        if role and start == 0 and end > 0:
            src = src[end:].lstrip(" :;—–-")
        return src

    action = _action_after_role(text)
    probe = (action[:1].upper() + action[1:]) if action else ""
    action = _action_after_role(probe) or action
    if not action:
        action = "фиксирует запись в журнале (5 минут)."
    action = action[:1].lower() + action[1:]
    return f"Информационная система автоматически {action}"


def _cut_rework_loop(body: str) -> str:
    """Убирает петлю «вернуть на п.N / на доработку» без слов верну/возврат/доработ/повтор."""
    text = re.sub(
        r"(?:иначе\s+)?(?:«[^»]+»\s*[—–-]\s*)?(?:вернуть|возврат(?:ить)?)\s+на\s+(?:п(?:ункт)?\.?\s*)?\d+",
        "иначе эскалация руководителю процесса",
        body,
        flags=re.I,
    )
    text = re.sub(
        r"«[^»]*(?:доработ|замечан|повторн)[^»]*»\s*[—–-]\s*(?:вернуть|возврат)\s+на\s+(?:п(?:ункт)?\.?\s*)?\d+",
        "иначе эскалация руководителю процесса",
        text,
        flags=re.I,
    )
    return re.sub(
        r"(?:если|иначе|при)\b[^.]{0,80}\bназад\b",
        "иначе эскалация руководителю процесса",
        text,
        flags=re.I,
    )


def _is_parallel_body(body: str) -> bool:
    return bool(re.match(r"^(?:параллельно|одновременно)\s*[:,—–-]?", body or "", re.I))


def _tobe_metrics(audit: Optional[dict], text: str) -> Dict[str, float]:
    data = audit or {}
    sla = data.get("sla") or {}
    cp = float(sla.get("critical_path_hours") or 0)
    rw = float(sla.get("with_rework_hours") or cp)
    _, steps = _split_steps(normalize_regulation(text or ""))
    return {
        "cp": cp,
        "rw": rw,
        "loops": float(len(data.get("rework_loops") or [])),
        "q": float(int((data.get("methodology") or {}).get("score") or 0)),
        "n": float(len(steps)),
    }


def _ppe_parallel_preserved(asis_text: str, tobe_text: str) -> bool:
    """Подготовка СИЗ/инструмента/переносных заземлений остаётся параллельной, как в As-Is."""
    _, asis_s = _split_steps(normalize_regulation(asis_text or ""))
    asis_ppe = [s["body"] for s in asis_s if _is_prepare_ppe_ground(s["body"])]
    if not asis_ppe or not any(_is_parallel_body(b) for b in asis_ppe):
        return True
    _, tobe_s = _split_steps(normalize_regulation(tobe_text or ""))
    tobe_ppe = [s["body"] for s in tobe_s if _is_prepare_ppe_ground(s["body"])]
    return bool(tobe_ppe) and any(_is_parallel_body(b) for b in tobe_ppe)


def _ot_order_ok(text: str) -> bool:
    """Ремонт после допуска/инструктажа; установка заземлений и допуск не параллельны; СИЗ-подготовка может быть AND."""
    raw = normalize_regulation(text or "")
    try:
        parsed = parse_regulation(raw)
        _, steps = _split_steps(raw)
    except Exception:  # noqa: BLE001
        return False
    n = min(len(parsed.steps), len(steps))
    if n == 0:
        return True
    bodies = [steps[i]["body"] for i in range(n)]
    idx_repair = next((i for i in range(n) if _is_repair_work(parsed.steps[i], bodies[i])), None)
    idx_brief = next(
        (i for i in range(n) if re.search(r"инструктаж", _ot_blob(parsed.steps[i], bodies[i]), re.I)),
        None,
    )
    idx_permit = next(
        (
            i
            for i in range(n)
            if re.search(r"наряд[\s-]*допуск|\bдопуск", _ot_blob(parsed.steps[i], bodies[i]), re.I)
            and not _is_prepare_ppe_ground(_ot_blob(parsed.steps[i], bodies[i]))
        ),
        None,
    )
    if idx_repair is not None:
        if _is_parallel_body(bodies[idx_repair]):
            return False
        if idx_permit is not None and idx_repair < idx_permit:
            return False
        if idx_brief is not None and idx_repair < idx_brief:
            return False
    for i in range(n):
        if not _is_parallel_body(bodies[i]):
            continue
        blob = _ot_blob(parsed.steps[i], bodies[i])
        if _is_prepare_ppe_ground(blob):
            if i and _forbid_parallel(parsed.steps[i - 1], parsed.steps[i], bodies[i - 1], bodies[i]):
                return False
            continue
        if _ot_sensitive(parsed.steps[i], bodies[i]) or _is_repair_work(parsed.steps[i], bodies[i]):
            return False
        if i and _forbid_parallel(parsed.steps[i - 1], parsed.steps[i], bodies[i - 1], bodies[i]):
            return False
    return True


def _tobe_feasible(asis: Dict[str, float], cand: Dict[str, float], asis_text: str, cand_text: str) -> bool:
    """Жёсткие ограничения + лексикографические откаты (голый путь, возвраты, циклы, quality, ОТ)."""
    if cand["cp"] > asis["cp"] + _TOBE_MINUTE:
        return False
    if asis["loops"] > 0 and cand["rw"] >= asis["rw"] - 1e-9:
        return False
    if cand["loops"] > asis["loops"]:
        return False
    if cand["q"] + 1e-9 < asis["q"]:
        return False
    if not _ot_order_ok(cand_text):
        return False
    if not _ppe_parallel_preserved(asis_text, cand_text):
        return False
    return True


def _heuristic_optimize_to_be(
    regulation_text: str,
    *,
    do_auto: bool = True,
    do_parallel: bool = True,
    do_safety: bool = True,
    do_loops: bool = True,
    do_control: bool = False,
    only_auto: Optional[Set[int]] = None,
    only_parallel: Optional[Set[int]] = None,
    only_loops: Optional[Set[int]] = None,
    only_control: Optional[Set[int]] = None,
) -> Tuple[str, List[Dict[str, str]]]:
    raw_text = normalize_regulation(regulation_text or "")
    header, raw_steps = _split_steps(raw_text)
    try:
        parsed = parse_regulation(raw_text)
    except Exception:  # noqa: BLE001
        return raw_text, []
    n = min(len(parsed.steps), len(raw_steps))
    if n < 2:
        return raw_text, []
    if not (do_auto or do_parallel or do_safety or do_loops or do_control):
        return raw_text, []
    actions: List[Dict[str, str]] = []
    bodies = [raw_steps[i]["body"] for i in range(n)]
    nums = [parsed.steps[i].num for i in range(n)]

    if do_auto:
        for i in range(n):
            if only_auto is not None and i not in only_auto:
                continue
            body = bodies[i]
            if _JOURNAL_RE.search(body) and not re.search(r"систем\w+\s+автоматическ", body, re.I):
                bodies[i] = _automate_journal_body(body)
                actions.append(
                    {
                        "kind": "automation",
                        "detail": (
                            f"Шаг {nums[i]} «{parsed.steps[i].title or 'фиксация'}»: "
                            "scriptTask, фиксация в журнале выполняется системой."
                        ),
                    }
                )

    if do_parallel:
        i = 0
        while i < n - 1:
            a, b = parsed.steps[i], parsed.steps[i + 1]
            if only_parallel is not None and i not in only_parallel:
                i += 1
                continue
            already = _is_parallel_body(bodies[i + 1])
            if (
                not _forbid_parallel(a, b, bodies[i], bodies[i + 1])
                and _independent_steps(a, b, bodies[i], bodies[i + 1])
                and not already
            ):
                bodies[i + 1] = "Параллельно: " + bodies[i + 1]
                b.parallel = True
                actions.append(
                    {
                        "kind": "parallel",
                        "detail": f"Шаги {a.num} ({a.role}) и {b.num} ({b.role}) выполняются параллельно.",
                    }
                )
                i += 2
                continue
            i += 1

    if do_safety:
        for i in range(n):
            st = parsed.steps[i]
            if _is_prepare_ppe_ground(_ot_blob(st, bodies[i])):
                continue
            if not (_ot_sensitive(st, bodies[i]) or _is_repair_work(st, bodies[i])):
                continue
            stripped = re.sub(r"^(?:параллельно|одновременно)\s*[:,—–-]?\s*", "", bodies[i], flags=re.I)
            if stripped != bodies[i]:
                bodies[i] = stripped
                st.parallel = False
                actions.append(
                    {
                        "kind": "safety_seq",
                        "detail": (
                            f"Шаг {nums[i]} «{st.title or 'работы'}» оставлен строго последовательным: "
                            "допуск / инструктаж / установка заземлений не параллельны ремонту."
                        ),
                    }
                )

    control_at: Dict[int, str] = {}
    if do_loops:
        for i in range(n):
            st = parsed.steps[i]
            d = st.decision
            loops = bool(d and (d.no_back or (d.no_ref is not None and d.no_ref < st.num)))
            if not loops:
                continue
            if only_loops is not None and i not in only_loops:
                continue
            bodies[i] = _cut_rework_loop(bodies[i])
            prev = bodies[i - 1] if i else ""
            if _CONTROL_HINT_RE.search(prev) or _CONTROL_HINT_RE.search(bodies[i]):
                actions.append(
                    {
                        "kind": "zero_rework",
                        "detail": f"Шаг {st.num}: петля возврата снята эскалацией, входной контроль уже есть.",
                    }
                )
                continue
            inserted = False
            if do_control and i and (only_control is None or i in only_control):
                prev_st = parsed.steps[i - 1]
                prev_role = prev_st.role or ""
                ctrl_role = next(
                    (s.role for s in parsed.steps if s.role and s.role not in {st.role, prev_role}),
                    "",
                )
                ctrl_body = (
                    f"Параллельно: {ctrl_role} проводит предварительный входной контроль "
                    f"комплектности документов и исходных данных перед шагом {st.num} (5 минут)."
                )
                ctrl_step = Step(
                    idx=-1,
                    num=0,
                    role=ctrl_role,
                    title="проводит предварительный входной контроль комплектности документов",
                    parallel=True,
                    hours=5.0 / 60.0,
                )
                if (
                    ctrl_role
                    and not _forbid_parallel(prev_st, ctrl_step, bodies[i - 1], ctrl_body)
                    and not _ot_sensitive(prev_st, bodies[i - 1])
                    and not _is_repair_work(prev_st, bodies[i - 1])
                ):
                    control_at[i] = ctrl_body
                    inserted = True
            actions.append(
                {
                    "kind": "zero_rework",
                    "detail": (
                        f"Шаг {st.num} «{st.title or 'согласование'}»: цикл заменён эскалацией "
                        "на исключительной ветке"
                        + (", входной контроль параллелен независимому шагу (5 минут)." if inserted else ".")
                    ),
                }
            )

    assembled: List[Tuple[Optional[int], str]] = []
    for i in range(n):
        if i in control_at:
            assembled.append((None, control_at[i]))
        assembled.append((nums[i], bodies[i]))
    old_to_new: Dict[int, int] = {}
    new_bodies: List[str] = []
    for old_num, body in assembled:
        new_bodies.append(body)
        if old_num is not None:
            old_to_new[old_num] = len(new_bodies)

    def _fn(ref: int) -> int:
        return old_to_new.get(ref, ref)

    remapped = [_remap(b, _fn) for b in new_bodies]
    return _join_steps(header, [{"body": b} for b in remapped]), actions


def _collect_tobe_actions(old_text: str, new_text: str) -> List[Dict[str, str]]:
    actions: List[Dict[str, str]] = []
    _, old_s = _split_steps(normalize_regulation(old_text))
    _, new_s = _split_steps(normalize_regulation(new_text))
    old_par = sum(1 for s in old_s if re.match(r"^(?:параллельно|одновременно)", s["body"], re.I))
    new_par = sum(1 for s in new_s if re.match(r"^(?:параллельно|одновременно)", s["body"], re.I))
    if new_par > old_par:
        actions.append({"kind": "parallel", "detail": f"Добавлено параллельных веток: {new_par - old_par}."})
    old_ctrl = sum(1 for s in old_s if _CONTROL_HINT_RE.search(s["body"]))
    new_ctrl = sum(1 for s in new_s if _CONTROL_HINT_RE.search(s["body"]))
    if new_ctrl > old_ctrl:
        actions.append({"kind": "zero_rework", "detail": f"Введён входной контроль ({new_ctrl - old_ctrl} шаг.). Петли возврата сокращены."})
    old_auto = sum(1 for s in old_s if re.search(r"автоматическ", s["body"], re.I))
    new_auto = sum(1 for s in new_s if re.search(r"автоматическ", s["body"], re.I))
    if new_auto > old_auto:
        actions.append({"kind": "automation", "detail": f"Автоматизировано шагов фиксации: {new_auto - old_auto}."})
    if not actions:
        actions.append({"kind": "reengine", "detail": "Целевой регламент переписан с сохранением ролей и SLA."})
    return actions


def _try_llm_optimize_to_be(regulation_text: str, audit_data: dict) -> Optional[str]:
    """Облачная перепись To-Be только при BPMN_TOBE_LLM=1 — иначе 25 с лишней генерации."""
    if os.getenv("BPMN_TOBE_LLM", "").strip().lower() not in ("1", "true", "yes"):
        return None
    ctx = build_process_context(regulation_text, "", audit_data or {})
    messages = [
        {"role": "system", "content": TOBE_SYSTEM},
        {
            "role": "user",
            "content": "АУДИТ AS-IS:\n" + format_context_for_prompt(ctx, max_steps=40)
            + "\n\nРЕГЛАМЕНТ AS-IS:\n" + (regulation_text or "").strip(),
        },
    ]
    trace: List[str] = []
    got = _chat_llm(messages, trace)
    if not got:
        return None
    _, raw = got
    new_text = _strip_markdown(raw).strip()
    _, old_steps = _split_steps(normalize_regulation(regulation_text))
    _, new_steps = _split_steps(normalize_regulation(new_text))
    if len(new_steps) < 2 or abs(len(new_steps) - len(old_steps)) > 8:
        return None
    return new_text + ("\n" if not new_text.endswith("\n") else "")


def _tobe_delta(old_audit: dict, new_audit: dict, actions: List[Dict[str, str]], engine: str) -> Dict[str, Any]:
    """Эффект To-Be: экономия КП = SLA_AsIs − SLA_ToBe по Беллману-Форду (critical_path_hours)."""
    old_a, new_a = old_audit or {}, new_audit or {}
    old_sla = old_a.get("sla") or {}
    new_sla = new_a.get("sla") or {}
    before = float(old_sla.get("critical_path_hours") or 0)
    after = float(new_sla.get("critical_path_hours") or 0)
    saved = before - after
    pct = round(100.0 * max(0.0, saved) / before, 1) if before and saved > 0 else 0.0
    loops_b = len(old_a.get("rework_loops") or [])
    loops_a = len(new_a.get("rework_loops") or [])
    q_b = int((old_a.get("methodology") or {}).get("score") or 0)
    q_a = int((new_a.get("methodology") or {}).get("score") or 0)
    rw_h_b = float(old_sla.get("rework_hours") or 0)
    rw_h_a = float(new_sla.get("rework_hours") or 0)
    before_rw = float(old_sla.get("with_rework_hours") or before)
    after_rw = float(new_sla.get("with_rework_hours") or after)
    return {
        "sla_before_hours": round(before, 3),
        "sla_after_hours": round(after, 3),
        "sla_saved_hours": round(saved, 3),
        "sla_saved_pct": pct,
        "rework_before": loops_b,
        "rework_after": loops_a,
        "rework_removed": max(0, loops_b - loops_a),
        "rework_hours_before": round(rw_h_b, 3),
        "rework_hours_after": round(rw_h_a, 3),
        "with_rework_before": round(before_rw, 3),
        "with_rework_after": round(after_rw, 3),
        "quality_before": q_b,
        "quality_after": q_a,
        "quality_gain": max(0, q_a - q_b),
        "breach_before": bool(old_sla.get("breach")),
        "breach_after": bool(new_sla.get("breach")),
        "actions": actions,
        "engine": engine,
    }


def _catalog_tobe_moves(regulation_text: str) -> Tuple[List[int], List[int], List[int]]:
    """Индексы отдельных ходов: автоматизация, соседняя параллель, снятие цикла."""
    raw_text = normalize_regulation(regulation_text or "")
    try:
        parsed = parse_regulation(raw_text)
        _, raw_steps = _split_steps(raw_text)
    except Exception:  # noqa: BLE001
        return [], [], []
    n = min(len(parsed.steps), len(raw_steps))
    bodies = [raw_steps[i]["body"] for i in range(n)]
    autos: List[int] = []
    pairs: List[int] = []
    loops: List[int] = []
    for i in range(n):
        body = bodies[i]
        if _JOURNAL_RE.search(body) and not re.search(r"систем\w+\s+автоматическ", body, re.I):
            autos.append(i)
        d = parsed.steps[i].decision
        if d and (d.no_back or (d.no_ref is not None and d.no_ref < parsed.steps[i].num)):
            loops.append(i)
    for i in range(n - 1):
        a, b = parsed.steps[i], parsed.steps[i + 1]
        if _is_parallel_body(bodies[i + 1]):
            continue
        if _forbid_parallel(a, b, bodies[i], bodies[i + 1]):
            continue
        if _independent_steps(a, b, bodies[i], bodies[i + 1]):
            pairs.append(i)
    return autos, pairs, loops


_TOBE_FULL_ENUM_LIMIT = 48
_TOBE_REBUILD_BUDGET = 64


def _iter_tobe_selections(
    autos: Sequence[int],
    pairs: Sequence[int],
    loops: Sequence[int],
) -> List[Tuple[Set[int], Set[int], Set[int], Set[int]]]:
    """Сочетания ходов. У цикла три состояния: не трогать, снять, снять и добавить контроль."""
    auto_l, pair_l, loop_l = list(autos), list(pairs), list(loops)
    count = (2 ** len(auto_l)) * (2 ** len(pair_l)) * (3 ** len(loop_l))
    selections: List[Tuple[Set[int], Set[int], Set[int], Set[int]]] = []

    def add(auto_idx: Sequence[int], pair_idx: Sequence[int], loop_states: Sequence[int]) -> None:
        chosen_loops = {loop_l[i] for i, state in enumerate(loop_states) if state}
        chosen_ctrl = {loop_l[i] for i, state in enumerate(loop_states) if state == 2}
        selections.append((set(auto_idx), set(pair_idx), chosen_loops, chosen_ctrl))

    if 0 < count <= _TOBE_FULL_ENUM_LIMIT:
        for amask in range(2 ** len(auto_l)):
            auto_idx = [auto_l[i] for i in range(len(auto_l)) if amask & (1 << i)]
            for pmask in range(2 ** len(pair_l)):
                pair_idx = [pair_l[i] for i in range(len(pair_l)) if pmask & (1 << i)]
                states = [0] * len(loop_l)
                if not loop_l:
                    add(auto_idx, pair_idx, states)
                    continue
                for lmask in range(3 ** len(loop_l)):
                    value = lmask
                    for i in range(len(loop_l)):
                        states[i] = value % 3
                        value //= 3
                    add(auto_idx, pair_idx, states)
        return selections

    full_states = [1] * len(loop_l)
    ctrl_states = [2] * len(loop_l)
    add(auto_l, pair_l, full_states)
    add(auto_l, pair_l, ctrl_states)
    add([], [], full_states)
    for i in range(len(auto_l)):
        add([auto_l[i]], [], [0] * len(loop_l))
        add([a for k, a in enumerate(auto_l) if k != i], pair_l, full_states)
    for i in range(len(pair_l)):
        add([], [pair_l[i]], [0] * len(loop_l))
        add(auto_l, [p for k, p in enumerate(pair_l) if k != i], full_states)
    for i in range(len(loop_l)):
        alone = [0] * len(loop_l)
        alone[i] = 1
        add([], [], alone)
        with_ctrl = [0] * len(loop_l)
        with_ctrl[i] = 2
        add([], [], with_ctrl)
        rest = [1] * len(loop_l)
        rest[i] = 0
        add(auto_l, pair_l, rest)
    return selections


def optimize_process_to_be(regulation_text: str, audit_data: dict) -> Tuple[str, dict]:
    """Реинжиниринг As-Is → To-Be: поиск оптимума с откатом ходов, нарушающих ОТ или удлиняющих голый путь.

    Лексикография после жёстких ограничений: короче путь с возвратами, затем короче голый КП, затем меньше шагов.
    Сначала считаются прежние наборы классов ходов — это текущий To-Be. Затем перебираются сочетания отдельных ходов.
    Сочетание заменяет текущий To-Be только если оно строго лучше. Если ни один кандидат не принят — возвращается As-Is.
    Облачная LLM включается только переменной BPMN_TOBE_LLM=1 и проходит ту же проверку.
    """
    text = (regulation_text or "").strip()
    empty = {
        "sla_before_hours": 0.0, "sla_after_hours": 0.0, "sla_saved_hours": 0.0, "sla_saved_pct": 0.0,
        "rework_before": 0, "rework_after": 0, "rework_removed": 0,
        "rework_hours_before": 0.0, "rework_hours_after": 0.0,
        "with_rework_before": 0.0, "with_rework_after": 0.0,
        "quality_before": 0, "quality_after": 0, "quality_gain": 0,
        "breach_before": False, "breach_after": False, "actions": [], "engine": "semantic-optimizer",
        "tobe_xml": "", "tobe_audit": {}, "tobe_error": "",
    }
    if len(text) < 20:
        return text, empty

    asis_audit = audit_data or {}
    asis_m = _tobe_metrics(asis_audit, text)
    flagsets: List[Dict[str, bool]] = [
        {"do_auto": True, "do_parallel": True, "do_safety": True, "do_loops": True, "do_control": False},
        {"do_auto": True, "do_parallel": True, "do_safety": True, "do_loops": True, "do_control": True},
        {"do_auto": True, "do_parallel": False, "do_safety": True, "do_loops": True, "do_control": False},
        {"do_auto": False, "do_parallel": True, "do_safety": True, "do_loops": True, "do_control": False},
        {"do_auto": True, "do_parallel": True, "do_safety": True, "do_loops": False, "do_control": False},
        {"do_auto": True, "do_parallel": False, "do_safety": False, "do_loops": True, "do_control": False},
        {"do_auto": False, "do_parallel": False, "do_safety": True, "do_loops": True, "do_control": False},
    ]
    seen: set = set()
    built = 0
    best: Optional[Tuple[Tuple[float, float, float], str, List[Dict[str, str]], str, dict, str]] = None

    def _consider(cand_text: str, cand_actions: List[Dict[str, str]], cand_engine: str) -> None:
        nonlocal best, built
        if built >= _TOBE_REBUILD_BUDGET:
            return
        key = normalize_regulation(cand_text or "")
        if not key or key in seen:
            return
        seen.add(key)
        built += 1
        try:
            xml, new_audit, err = _rebuild_process(cand_text, f"to-be · {cand_engine}", [], use_llm=(cand_engine == "llm"))
        except Exception as exc:  # noqa: BLE001
            xml, new_audit, err = "", {}, f"{type(exc).__name__}: {exc}"
        if err or not new_audit:
            return
        metrics = _tobe_metrics(new_audit, cand_text)
        if not _tobe_feasible(asis_m, metrics, text, cand_text):
            return
        score = (metrics["rw"], metrics["cp"], metrics["n"])
        pack = (score, cand_text, cand_actions or [], xml, new_audit, cand_engine)
        if best is None or pack[0] < best[0]:
            best = pack

    llm_text = None
    try:
        llm_text = _try_llm_optimize_to_be(text, asis_audit)
    except Exception:  # noqa: BLE001
        llm_text = None
    if llm_text:
        _consider(llm_text, _collect_tobe_actions(text, llm_text), "llm")

    for flags in flagsets:
        opt_text, actions = _heuristic_optimize_to_be(text, **flags)
        if not actions:
            continue
        _consider(opt_text, actions, "semantic-optimizer")

    autos, pairs, loops = _catalog_tobe_moves(text)
    for auto_idx, pair_idx, loop_idx, ctrl_idx in _iter_tobe_selections(autos, pairs, loops):
        if not (auto_idx or pair_idx or loop_idx or ctrl_idx):
            continue
        opt_text, actions = _heuristic_optimize_to_be(
            text,
            do_auto=bool(auto_idx),
            do_parallel=bool(pair_idx),
            do_safety=True,
            do_loops=bool(loop_idx),
            do_control=bool(ctrl_idx),
            only_auto=auto_idx,
            only_parallel=pair_idx,
            only_loops=loop_idx,
            only_control=ctrl_idx,
        )
        if not actions:
            continue
        _consider(opt_text, actions, "semantic-optimizer")

    engine = "semantic-optimizer"
    if best is None:
        optimized = text
        actions = [{"kind": "stable", "detail": "Существенных узких мест для автоматического реинжиниринга не найдено."}]
        xml, new_audit, err = "", asis_audit, ""
        if asis_audit.get("sla"):
            pass
        else:
            try:
                xml, new_audit, err = _rebuild_process(text, "to-be · as-is", [], use_llm=False)
            except Exception as exc:  # noqa: BLE001
                err = f"{type(exc).__name__}: {exc}"
        delta = _tobe_delta(asis_audit, new_audit or asis_audit, actions, engine)
        delta["tobe_xml"] = xml or ""
        delta["tobe_audit"] = new_audit or asis_audit
        delta["tobe_error"] = err or ""
        return optimized, delta

    _score, optimized, actions, xml, new_audit, engine = best
    actions = _limit_escalation_actions(actions, asis_audit, new_audit)
    if not actions:
        actions = [{"kind": "stable", "detail": "Существенных узких мест для автоматического реинжиниринга не найдено."}]
    delta = _tobe_delta(asis_audit, new_audit or {}, actions, engine)
    delta["tobe_xml"] = xml or ""
    delta["tobe_audit"] = new_audit or {}
    delta["tobe_error"] = ""
    return optimized, delta


def _docx_shade(cell: Any, fill: str) -> None:
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}" w:val="clear"/>'))


def _docx_set_col_widths(table: Any, widths_cm: Sequence[float]) -> None:
    """Жёсткая сетка колонок: autofit выключен, ширина ячеек в сантиметрах."""
    from docx.shared import Cm

    table.autofit = False
    if hasattr(table, "allow_autofit"):
        table.allow_autofit = False
    widths = [Cm(float(w)) for w in widths_cm]
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = w


def _docx_set_cell(cell: Any, text: str, *, header: bool = False, center: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    cell.text = ""
    p = cell.paragraphs[0]
    if center or header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text if text not in (None, "") else "—"))
    run.font.size = Pt(9 if not header else 9)
    run.font.name = "Calibri"
    run.bold = header
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)
        _docx_shade(cell, "003366")


def export_docx_passport(xml_str: str, audit_data: dict, regulation_text: str) -> bytes:
    """Официальный регламент Microsoft Word: паспорт, RACI, пошаговый порядок, карта рисков."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("Для экспорта .docx установите пакет python-docx>=1.0.0") from exc

    audit = audit_data if isinstance(audit_data, dict) else {}
    text = regulation_text or ""
    try:
        parsed = parse_regulation(normalize_regulation(text)) if text.strip() else ParsedRegulation(
            title="Бизнес-процесс", sla_hours=None
        )
    except Exception:  # noqa: BLE001
        parsed = ParsedRegulation(title="Бизнес-процесс", sla_hours=None)

    sla = audit.get("sla") or {}
    stats = audit.get("stats") or {}
    bus = audit.get("bus_factor") or {}
    loops = audit.get("rework_loops") or []
    recs = audit.get("recommendations") or []
    target = sla.get("target_hours") if sla.get("target_hours") is not None else parsed.sla_hours
    owner = _process_owner(text, parsed, audit)
    title = parsed.title or "Бизнес-процесс по регламенту"
    roles = list(parsed.roles or [x.get("role") for x in (audit.get("lane_load") or []) if x.get("role")])
    raci = generate_raci_matrix(list(parsed.steps), roles)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.6)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("ПАО «Интер РАО»  ·  Дирекция бизнес-архитектуры  ·  официальный регламент процесса")
    hr.bold = True
    hr.font.size = Pt(9)
    hr.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    hr.font.name = "Calibri"
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Архитектор BPMN-диаграмм  ·  документ сформирован автоматически  ·  конфиденциально")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x45, 0x5A, 0x64)
    fr.font.name = "Calibri"

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = h.add_run("Регламент бизнес-процесса")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    r.font.name = "Calibri"
    sub = doc.add_paragraph()
    sr = sub.add_run(title)
    sr.bold = True
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)
    sr.font.name = "Calibri"

    def heading(label: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        run.font.name = "Calibri"

    heading("1. Паспорт процесса")
    pass_rows = [
        ("Наименование", title),
        ("Владелец процесса", owner),
        ("Целевой SLA", _fh(float(target)) if target else "не задан"),
        ("Критический путь", _fh(float(sla["critical_path_hours"])) if sla.get("critical_path_hours") is not None else "—"),
        ("Срок с учётом возврата", _fh(float(sla["with_rework_hours"])) if sla.get("with_rework_hours") is not None else "—"),
        ("Срыв SLA", "да" if sla.get("breach") else "нет"),
        ("Bus-factor", f"{float(bus.get('max_share') or 0):.0%} · {bus.get('top_role') or '—'}"),
        ("Узлов / ролей / подпроцессов", f"{stats.get('nodes', '—')} / {stats.get('lanes', len(roles))} / {stats.get('subprocesses', '—')}"),
        ("Quality Score", f"{int((audit.get('methodology') or {}).get('score') or 0)}%"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _docx_set_cell(table.rows[0].cells[0], "Параметр", header=True)
    _docx_set_cell(table.rows[0].cells[1], "Значение", header=True)
    for k, v in pass_rows:
        row = table.add_row().cells
        _docx_set_cell(row[0], k)
        _docx_set_cell(row[1], v)

    heading("2. Матрица ответственности (RACI)")
    legend = doc.add_paragraph()
    lg = legend.add_run("R — Responsible (исполнитель) · A — Accountable (итог) · C — Consulted (эксперт) · I — Informed (уведомление)")
    lg.font.size = Pt(8)
    lg.font.color.rgb = RGBColor(0x45, 0x5A, 0x64)
    cols = ["№", "Шаг", *roles] if roles else ["№", "Шаг"]
    rtable = doc.add_table(rows=1, cols=len(cols))
    rtable.style = "Table Grid"
    for i, name in enumerate(cols):
        _docx_set_cell(rtable.rows[0].cells[i], name, header=True, center=True)
    for item in raci:
        cells = rtable.add_row().cells
        _docx_set_cell(cells[0], str(item.get("num") or ""), center=True)
        _docx_set_cell(cells[1], str(item.get("title") or "")[:80])
        assigns = item.get("assignments") or {}
        for j, role in enumerate(roles):
            letters = "".join(assigns.get(role) or [])
            _docx_set_cell(cells[j + 2], letters or "—", center=True)
    usable = 17.4
    n_roles = len(roles)
    remain = usable - 1.0 - 6.5
    if n_roles:
        role_w = remain / n_roles
        if n_roles * 1.6 <= remain + 1e-9:
            role_w = min(1.8, max(1.6, role_w))
        raci_widths = [1.0, 6.5] + [role_w] * n_roles
    else:
        raci_widths = [1.0, 16.4]
    _docx_set_col_widths(rtable, raci_widths)

    heading("3. Пошаговый операционный регламент")
    ot = doc.add_table(rows=1, cols=6)
    ot.style = "Table Grid"
    for i, name in enumerate(["№", "Роль", "Действие", "Срок", "Системы", "Документы"]):
        _docx_set_cell(ot.rows[0].cells[i], name, header=True)
    if parsed.steps:
        for st in parsed.steps:
            row = ot.add_row().cells
            _docx_set_cell(row[0], str(st.num), center=True)
            _docx_set_cell(row[1], st.role)
            name = st.title or "(шлюз решения)"
            if st.stage:
                name = f"{name} (этап «{st.stage}»)"
            _docx_set_cell(row[2], name)
            _docx_set_cell(row[3], _fh(float(st.hours)) if st.hours else "—", center=True)
            _docx_set_cell(row[4], ", ".join(st.systems) or "—")
            _docx_set_cell(row[5], ", ".join(st.artifacts) or "—")
    else:
        row = ot.add_row().cells
        _docx_set_cell(row[0], "—")
        _docx_set_cell(row[1], "—")
        _docx_set_cell(row[2], "Шаги не распознаны")
        _docx_set_cell(row[3], "—")
        _docx_set_cell(row[4], "—")
        _docx_set_cell(row[5], "—")
    _docx_set_col_widths(ot, [1.0, 3.5, 6.0, 2.0, 2.5, 2.5])

    heading("4. Карта рисков и план мероприятий по оптимизации")
    risk_p = doc.add_paragraph()
    rp = risk_p.add_run(
        f"Циклов возврата: {len(loops)}. "
        f"Худший возврат: {_fh(float(sla.get('rework_hours') or 0))}. "
        f"Рекомендации движка аудита:"
    )
    rp.font.size = Pt(10)
    rp.font.name = "Calibri"
    if loops:
        lt = doc.add_table(rows=1, cols=4)
        lt.style = "Table Grid"
        for i, name in enumerate(["Цикл", "Откуда", "Куда", "Стоимость"]):
            _docx_set_cell(lt.rows[0].cells[i], name, header=True)
        for loop in loops:
            row = lt.add_row().cells
            _docx_set_cell(row[0], str(loop.get("label") or ""))
            _docx_set_cell(row[1], str(loop.get("from") or ""))
            _docx_set_cell(row[2], str(loop.get("to") or ""))
            _docx_set_cell(row[3], _fh(float(loop.get("cycle_hours") or 0)))
    if recs:
        for rec in recs:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(str(rec))
            run.font.size = Pt(10)
            run.font.name = "Calibri"
    else:
        p = doc.add_paragraph()
        run = p.add_run("Существенных узких мест не выявлено: процесс сбалансирован по ролям.")
        run.font.size = Pt(10)
        run.font.name = "Calibri"
    note = doc.add_paragraph()
    nr = note.add_run(
        "План To-Be: параллелить независимые шаги разных ролей, ввести входной контроль перед согласованиями, "
        "автоматизировать фиксацию в журналах (scriptTask). Целевая схема строится во вкладке «Оптимизация As-Is → To-Be»."
    )
    nr.italic = True
    nr.font.size = Pt(9)
    nr.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)
    nr.font.name = "Calibri"
    if xml_str:
        meta = doc.add_paragraph()
        mr = meta.add_run(f"Исходная BPMN-модель: {len(xml_str)} символов XML.")
        mr.font.size = Pt(8)
        mr.font.color.rgb = RGBColor(0x78, 0x90, 0x9C)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

